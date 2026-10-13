"""Resource and boundary regressions for the attachment ingestion runtime."""
import ast
import builtins
import importlib.util
import json
import symtable
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from companion_v01.attachment_ingest import AttachmentIngestService
from companion_v01 import attachment_ingest
from companion_v01.attachment_inbox import AttachmentInboxService
from companion_v01.store import MemoryStore


ROOT = Path(__file__).resolve().parents[1] / 'companion_v01'
ENTRY = ROOT / 'attachment_ingest.py'
PACKAGE = ROOT / 'attachment_ingest_runtime'


class AttachmentIngestBoundaryTests(unittest.TestCase):
    def test_modules_have_no_unbound_global_dependencies(self):
        for path in PACKAGE.glob('*.py'):
            table = symtable.symtable(path.read_text(encoding='utf-8'), str(path), 'exec')
            bound = {symbol.get_name() for symbol in table.get_symbols()
                     if symbol.is_imported() or symbol.is_assigned() or symbol.is_namespace()}
            bound.update(dir(builtins))
            bound.update({'__name__', '__package__', '__file__', '__doc__', '__builtins__'})

            def check(scope):
                for symbol in scope.get_symbols():
                    if symbol.is_referenced() and symbol.is_global():
                        self.assertIn(symbol.get_name(), bound, f'{path.name}: unbound global')
                for child in scope.get_children():
                    check(child)

            check(table)

    def test_all_implementation_modules_are_reachable_without_a_reverse_entry_import(self):
        graph = {}

        def visit(path):
            if path in graph:
                return
            graph[path] = set()
            for node in ast.walk(ast.parse(path.read_text(encoding='utf-8'))):
                if not isinstance(node, ast.ImportFrom) or not node.level:
                    continue
                base = path.parent
                for _ in range(node.level - 1):
                    base = base.parent
                dependency = base.joinpath(*(node.module or '').split('.')).with_suffix('.py')
                if path.is_relative_to(PACKAGE):
                    self.assertNotEqual(dependency, ENTRY, f'{path.name} imports the service entry')
                if dependency.is_relative_to(PACKAGE) and dependency.exists():
                    graph[path].add(dependency)
                    visit(dependency)

        visit(ENTRY)
        expected = set(PACKAGE.glob('*.py')) - {PACKAGE / '__init__.py'}
        self.assertTrue(expected)
        self.assertEqual(set(graph) - {ENTRY}, expected, 'an implementation is disconnected')
        active, complete = set(), set()

        def check(path):
            self.assertNotIn(path, active, f'cyclic ingestion import at {path.name}')
            if path in complete:
                return
            active.add(path)
            for dependency in graph[path]:
                check(dependency)
            active.remove(path)
            complete.add(path)

        check(ENTRY)

    def test_entry_contracts_keep_their_identity_and_private_urls_stay_out_of_repr(self):
        from companion_v01.attachment_ingest_runtime import definitions

        for name in ('AttachmentMaterializationError', 'RemoteMediaDescriptor', 'TEXT_SUFFIXES',
                     'DOCUMENT_SUFFIXES', 'MEDIA_SUFFIXES', 'AUDIO_MEDIA_SUFFIXES', 'IMAGE_SUFFIXES'):
            self.assertIs(getattr(attachment_ingest, name), getattr(definitions, name))
        private_url = 'https://example.com/media?token=private-fixture'
        descriptor = attachment_ingest.RemoteMediaDescriptor(
            source_url=private_url, webpage_url=private_url, thumbnail_url=private_url,
            media_urls=(private_url,), title='example', ext='mp4', mime_type='video/mp4',
            kind='file', download_mode='direct',
        )
        self.assertNotIn(private_url, repr(descriptor))
        self.assertEqual(descriptor.origin_name, 'example.mp4')

    def test_service_methods_have_one_owner(self):
        owners = {}
        for owner in AttachmentIngestService.__mro__:
            if not owner.__module__.startswith('companion_v01.attachment_ingest'):
                continue
            for name, value in vars(owner).items():
                if name.startswith('__') or not callable(value):
                    continue
                self.assertNotIn(name, owners, f'{name} has conflicting method owners')
                owners[name] = owner
        self.assertTrue(owners)


class AttachmentIngestResourceTests(unittest.TestCase):
    def test_text_byte_limit_marks_the_material_card_as_a_partial_preview(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / 'large.txt'
            source.write_bytes(b'row ' * 2048)
            runtime = object.__new__(AttachmentIngestService)
            with patch('companion_v01.attachment_ingest.config.QQ_TEXT_ATTACHMENT_MAX_READ_BYTES', 1024):
                card = runtime._build_file_card(
                    source_path=source, item={'kind': 'document', 'origin_name': source.name},
                    mime_type='text/plain',
                )
            self.assertEqual(len(card['detail']['text_preview']), 1024)
            self.assertEqual(card['detail']['file_size'], source.stat().st_size)
            self.assertTrue(card['detail']['preview_is_truncated'])

    def test_text_preview_reads_only_the_configured_byte_budget(self):
        content = b'preview-data-' * 4096
        bytes_read = []
        open_file = Path.open

        class ObservedReader:
            def __init__(self, stream):
                self.stream = stream

            def __enter__(self):
                return self

            def __exit__(self, *args):
                self.stream.close()

            def read(self, size=-1):
                data = self.stream.read(size)
                bytes_read.append(len(data))
                return data

        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / 'large.txt'
            source.write_bytes(content)

            def observe_open(path, *args, **kwargs):
                stream = open_file(path, *args, **kwargs)
                return ObservedReader(stream) if path == source else stream

            runtime = object.__new__(AttachmentIngestService)
            with patch('companion_v01.attachment_ingest.config.QQ_TEXT_ATTACHMENT_MAX_READ_BYTES', 16), patch.object(
                Path, 'open', observe_open,
            ):
                text, encoding = runtime._read_text_preview(source)
            self.assertEqual(text, content[:16].decode('utf-8'))
            self.assertEqual(encoding, 'utf-8')
            self.assertLessEqual(sum(bytes_read), 16, 'preview read beyond its configured byte budget')

    @unittest.skipIf(importlib.util.find_spec('openpyxl') is None, 'openpyxl is unavailable')
    def test_workbook_card_counts_all_sheets_and_marks_the_limited_preview(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / 'seven-sheets.xlsx'
            workbook = Workbook()
            while len(workbook.sheetnames) < 7:
                workbook.create_sheet()
            for sheet in workbook:
                sheet['A1'] = 'one cell'
            workbook.save(source)
            workbook.close()
            runtime = object.__new__(AttachmentIngestService)
            card = runtime._build_xlsx_card(
                source_path=source, title=source.name,
                mime_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                file_size=source.stat().st_size,
            )
            self.assertEqual(card['detail']['sheet_count'], 7)
            self.assertIn('7 个工作表', card['detail']['summary'])
            self.assertEqual(len(card['detail']['sheets']), 5)
            self.assertEqual(len(card['detail']['sheet_names']), 5)
            self.assertTrue(card['detail']['preview_is_truncated'])

    def test_workbook_is_closed_when_reading_a_sheet_fails(self):
        class BrokenWorkbook:
            sheetnames = ['Broken']
            closed = False

            def __getitem__(self, name):
                raise ValueError('sheet decoding failed')

            def close(self):
                self.closed = True

        workbook = BrokenWorkbook()
        provider = types.ModuleType('openpyxl')
        provider.load_workbook = lambda **kwargs: workbook
        runtime = object.__new__(AttachmentIngestService)
        with patch.dict(sys.modules, {'openpyxl': provider}), patch(
            'companion_v01.attachment_ingest.importlib.util.find_spec', return_value=object(),
        ):
            card = runtime._build_xlsx_card(
                source_path=Path('broken.xlsx'), title='broken.xlsx', mime_type='', file_size=0,
            )
        self.assertEqual(card['detail']['parser_status'], 'unsupported_or_unavailable')
        self.assertIn('sheet decoding failed', card['detail']['summary'])
        self.assertTrue(workbook.closed, 'failed spreadsheet parsing retained an open workbook')


class AttachmentIngestFailureTests(unittest.TestCase):
    def test_platform_failure_preserves_its_specific_code_in_the_saved_attachment(self):
        codes = ('remote_media_multi_segment_unsupported', 'remote_media_bilibili_stream_missing',
                 'remote_media_bilibili_api_failed', 'remote_media_bilibili_identity_missing')
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            store = MemoryStore(root / 'db')
            inbox = AttachmentInboxService(store=store)
            runtime = AttachmentIngestService(
                base_dir=root / 'attachments', store=store, attachment_service=inbox, vision_service=None,
            )
            for code in codes:
                item = inbox.create_pending(
                    profile_user_id='fixture', session_id='session', source='remote_url',
                    kind='file', origin_name='video.mp4', timestamp=100,
                )
                runtime._mark_failed(item, code, timestamp=101)
                saved = store.get_attachment_inbox_item(
                    profile_user_id='fixture', session_id='session', attachment_id=item['attachment_id'],
                )
                self.assertEqual(saved['status'], 'failed')
                self.assertEqual(saved['error_message'], code)
                self.assertEqual(saved['detail']['failure']['code'], code)
                self.assertEqual(saved['short_hint'], runtime._humanize_remote_fetch_error(code))

    def test_unclassified_failure_does_not_save_the_raw_source_locator(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            store = MemoryStore(root / 'db')
            inbox = AttachmentInboxService(store=store)
            runtime = AttachmentIngestService(
                base_dir=root / 'attachments', store=store, attachment_service=inbox, vision_service=None,
            )
            item = inbox.create_pending(
                profile_user_id='fixture', session_id='session', source='qq',
                kind='file', origin_name='fixture.bin', timestamp=100,
            )
            runtime._mark_failed(item, 'unexpected https://example.com/file?token=private-fixture', timestamp=101)
            saved = store.get_attachment_inbox_item(
                profile_user_id='fixture', session_id='session', attachment_id=item['attachment_id'],
            )
            self.assertEqual(saved['error_message'], 'attachment_materialization_failed')
            self.assertNotIn('private-fixture', json.dumps(saved))


if __name__ == '__main__':
    unittest.main()
