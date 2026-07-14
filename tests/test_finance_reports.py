from __future__ import annotations

from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

import config
from companion_v01.attachment_inbox import AttachmentInboxService
from companion_v01.finance import (
    FINANCE_REPORT_DISCLAIMER,
    FinanceAnalysisResult,
    FinanceReportRequest,
    MarketDataToolService,
    QQFinanceDeliveryAdapter,
    ComposeFinanceReportToolHandler,
    RenderMarketChartToolHandler,
)
from companion_v01.generated_files import GeneratedFileService
from companion_v01.native_tool_schema import build_openai_native_tool_specs
from companion_v01.qq_gateway import NapCatQQGateway, QQMessageContext
from companion_v01.store import MemoryStore
from companion_v01.tool_runtime import ToolExecutionContext
from services.market_data import (
    FinanceSubscription,
    MarketDataResponse,
    MarketDataValidationError,
    MarketEventStore,
    MockMarketDataProvider,
)


FIXTURE_PATH = Path(__file__).parent / "fixtures" / "choice_market_data_synthetic_v1.json"
MARKET_ZONE = ZoneInfo("Asia/Shanghai")
NOW_TS = int(datetime(2026, 7, 10, 16, 0, tzinfo=MARKET_ZONE).timestamp())


class FinanceReportRequestTests(unittest.TestCase):
    def test_request_uses_fixed_report_grammar(self) -> None:
        request = FinanceReportRequest(
            report_type="security_brief",
            codes=("000000.test",),
            output_format="pdf",
            lookback=60,
            chart_targets=("gen_001",),
            risk_notes=("等待后续披露。",),
        )

        self.assertEqual(request.codes, ("000000.TEST",))
        self.assertEqual(request.output_format, "pdf")
        with self.assertRaises(MarketDataValidationError):
            FinanceReportRequest(
                report_type="security_brief",
                codes=("000000.TEST", "000001.TEST"),
                output_format="pdf",
            )
        with self.assertRaises(MarketDataValidationError):
            FinanceReportRequest(
                report_type="security_brief",
                codes=("000000.TEST",),
                output_format="docx",
            )


class ComposeFinanceReportToolTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.memory_store = MemoryStore(self.root / "memory")
        self.generated_service = GeneratedFileService(
            base_dir=self.root / "generated_files",
            store=self.memory_store,
            attachment_service=AttachmentInboxService(store=self.memory_store),
        )
        provider = MockMarketDataProvider.from_fixture_path(FIXTURE_PATH)
        self.market_service = MarketDataToolService(
            provider=provider,
            event_store=MarketEventStore(self.root / "events.sqlite3"),
            clock=lambda: NOW_TS,
        )
        handlers = (
            RenderMarketChartToolHandler(
                service=self.market_service,
                generated_file_service=self.generated_service,
            ),
            ComposeFinanceReportToolHandler(
                service=self.market_service,
                generated_file_service=self.generated_service,
            ),
        )
        self.handlers = {handler.tool_type: handler for handler in handlers}
        self.context = ToolExecutionContext(
            profile_user_id="owner",
            session_id="session",
            now_ts=NOW_TS,
            visual_payload={},
            current_user_source_id="message::finance-report",
            request_context={"message": "请生成 000000.TEST 的正式金融报告"},
        )

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _generate_chart(self) -> dict:
        handler = self.handlers["render_market_chart"]
        call = handler.normalize_call(
            {
                "type": "render_market_chart",
                "code": "000000.TEST",
                "lookback": 20,
                "title": "合成测试量价图",
                "send_to_user": False,
            }
        )
        result = handler.execute(call=call or {}, context=self.context)
        return result.stream_events[0]["generated_file"]

    def _report_call(self, output_format: str, *, chart_handle: str = "") -> dict:
        raw = {
            "type": "compose_finance_report",
            "report_type": "security_brief",
            "codes": ["000000.TEST"],
            "output_format": output_format,
            "lookback": 20,
            "title": f"合成测试{output_format.upper()}报告",
            "analysis_summary": "程序指标显示观察期收盘价上升；这一句属于模型分析，不是原始事实字段。",
            "risk_notes": ["测试数据不代表真实市场。"],
            "watch_items": ["继续观察成交量与后续披露。"],
            "send_to_user": True,
        }
        if chart_handle:
            raw["chart_ids"] = [chart_handle]
        handler = self.handlers["compose_finance_report"]
        call = handler.normalize_call(raw)
        self.assertIsNotNone(call)
        return call or {}

    def test_native_schema_excludes_raw_data_paths_templates_and_code(self) -> None:
        specs = build_openai_native_tool_specs(self.handlers)
        report = next(item for item in specs if item["function"]["name"] == "compose_finance_report")
        parameters = report["function"]["parameters"]
        properties = parameters["properties"]

        self.assertFalse(parameters["additionalProperties"])
        self.assertEqual(parameters["required"], ["report_type", "codes", "output_format"])
        self.assertTrue(
            {"prices", "series", "raw_data", "template", "output_path", "code_snippet"}.isdisjoint(properties)
        )
        handler = self.handlers["compose_finance_report"]
        metadata = handler.tool_metadata()
        self.assertEqual((metadata.family, metadata.operation, metadata.risk), ("finance_artifact", "control", "low"))
        self.assertIsNone(
            handler.normalize_call(
                {
                    "type": "compose_finance_report",
                    "report_type": "security_brief",
                    "codes": ["000000.TEST"],
                    "output_format": "pdf",
                    "raw_data": [{"close": 999999}],
                }
            )
        )

    def test_markdown_report_refetches_evidence_and_uses_managed_chart_reference(self) -> None:
        chart = self._generate_chart()
        handler = self.handlers["compose_finance_report"]
        call = self._report_call("md", chart_handle=chart["generated_handle"])
        call["analysis_summary"] = "![任意引用](https://example.invalid/leak.png)\n正常分析。"

        result = handler.execute(
            call=call,
            context=self.context,
        )

        event = result.stream_events[0]
        self.assertEqual(event["type"], "finance_report_ready")
        generated = event["generated_file"]
        self.assertEqual(generated["output_format"], "md")
        self.assertEqual(generated["created_by_tool"], "compose_finance_report")
        self.assertEqual(generated["source_ids"], [chart["generated_id"], "message::finance-report"])
        content = Path(generated["absolute_path"]).read_text(encoding="utf-8")
        self.assertIn("## 事实与程序指标", content)
        self.assertIn("## 分析解读（模型生成，非原始事实字段）", content)
        self.assertIn("2026-07-10T15:00:00+08:00", content)
        self.assertIn(FINANCE_REPORT_DISCLAIMER, content)
        self.assertIn("![合成测试量价图]", content)
        self.assertNotIn("](https://", content)
        self.assertIn(r"\!\[任意引用\]\(https://example\.invalid/leak\.png\)", content)
        self.assertNotIn(str(self.root), content)
        self.assertNotIn(str(self.root), result.followup_context)
        self.assertIn('"evidence_sha256":', result.followup_context)

    def test_pdf_report_embeds_trusted_chart_image(self) -> None:
        chart = self._generate_chart()
        handler = self.handlers["compose_finance_report"]

        result = handler.execute(
            call=self._report_call("pdf", chart_handle=chart["generated_handle"]),
            context=self.context,
        )

        generated = result.stream_events[0]["generated_file"]
        path = Path(generated["absolute_path"])
        reader = PdfReader(str(path))
        image_count = 0
        for page in reader.pages:
            resources = page.get("/Resources")
            xobjects = resources.get("/XObject") if resources else None
            if not xobjects:
                continue
            for item in xobjects.get_object().values():
                if item.get_object().get("/Subtype") == "/Image":
                    image_count += 1
        self.assertGreater(path.stat().st_size, 0)
        self.assertGreaterEqual(len(reader.pages), 1)
        self.assertGreaterEqual(image_count, 1)
        self.assertNotIn(str(self.root).encode(), path.read_bytes())

    def test_xlsx_report_contains_metrics_raw_ohlcv_evidence_and_chart(self) -> None:
        chart = self._generate_chart()
        handler = self.handlers["compose_finance_report"]
        call = self._report_call("xlsx", chart_handle=chart["generated_handle"])
        call["analysis_summary"] = '=WEBSERVICE("https://example.invalid")'

        result = handler.execute(
            call=call,
            context=self.context,
        )

        generated = result.stream_events[0]["generated_file"]
        path = Path(generated["absolute_path"])
        workbook = load_workbook(path, read_only=False, data_only=True)
        self.assertTrue({"Summary", "Metrics", "Quotes", "Evidence", "Notes", "Charts"}.issubset(workbook.sheetnames))
        self.assertIn("000000.TEST_OHLCV", workbook.sheetnames)
        self.assertEqual(workbook["Metrics"]["A2"].value, "000000.TEST")
        self.assertEqual(workbook["Metrics"]["F2"].value, 102)
        self.assertEqual(workbook["000000.TEST_OHLCV"].max_row, 4)
        self.assertEqual(len(workbook["Charts"]._images), 1)
        self.assertEqual(workbook["Notes"]["B2"].value, '\'=WEBSERVICE("https://example.invalid")')
        workbook.close()
        self.assertNotIn(str(self.root).encode(), path.read_bytes())
        with zipfile.ZipFile(path) as archive:
            self.assertTrue(any(name.startswith("xl/media/") for name in archive.namelist()))

    def test_non_chart_generated_file_cannot_be_embedded(self) -> None:
        fake_path = self.generated_service.allocate_output_path(
            profile_user_id="owner",
            session_id="session",
            title="伪图表",
            output_format="png",
            timestamp=NOW_TS,
        )
        fake_path.parent.mkdir(parents=True, exist_ok=True)
        fake_path.write_bytes(b"not-a-trusted-chart")
        fake = self.generated_service.register_generated_artifact(
            profile_user_id="owner",
            session_id="session",
            output_path=fake_path,
            output_title="伪图表",
            output_format="png",
            mime_type="image/png",
            content_card={"kind": "other"},
            summary="not trusted",
            created_by_tool="compose_file",
            timestamp=NOW_TS,
        )
        handler = self.handlers["compose_finance_report"]

        result = handler.execute(
            call=self._report_call("pdf", chart_handle=fake["generated_handle"]),
            context=self.context,
        )

        self.assertFalse(any(event.get("type") == "finance_report_ready" for event in result.stream_events))
        self.assertEqual(result.stream_events[0]["status"], "invalid_arguments")
        self.assertEqual(
            len(self.memory_store.list_generated_files(profile_user_id="owner", session_id="session")),
            1,
        )

    def test_incomplete_upstream_series_does_not_create_partial_report(self) -> None:
        handler = self.handlers["compose_finance_report"]
        unavailable = MarketDataResponse(
            ok=False,
            status="unavailable",
            provider="synthetic_choice",
            source="Synthetic Choice Fixture",
            as_of=None,
            reason="synthetic outage",
            data=None,
        )

        with patch.object(self.market_service, "price_series_response", return_value=unavailable):
            result = handler.execute(call=self._report_call("md"), context=self.context)

        self.assertFalse(any(event.get("type") == "finance_report_ready" for event in result.stream_events))
        self.assertEqual(result.stream_events[0]["status"], "unavailable")
        self.assertEqual(
            self.memory_store.list_generated_files(profile_user_id="owner", session_id="session"),
            [],
        )


class FinanceReportQQDeliveryTests(unittest.TestCase):
    def _event(self, path: Path) -> dict:
        return {
            "type": "finance_report_ready",
            "delivery_scope": "finance_report",
            "send_to_user": True,
            "generated_file": {
                "generated_id": "generated::report",
                "absolute_path": str(path),
                "output_title": "市场简报",
                "output_format": "pdf",
                "file_ext": "pdf",
                "mime_type": "application/pdf",
                "created_by_tool": "compose_finance_report",
            },
        }

    def test_subscription_authorization_sends_report_without_current_file_intent(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = Path(temp_dir) / "report.pdf"
            report_path.write_bytes(b"non-empty")
            gateway = NapCatQQGateway()
            context = QQMessageContext(
                should_respond=True,
                reason="finance_subscription_push",
                is_group=True,
                target_id=20001,
                group_id=20001,
                session_id="qq_group_shared_20001",
                profile_user_id="qq_group_shared_20001",
                finance_mode="push",
                clean_message="",
                raw_message="",
            )

            class FakeResponse:
                def raise_for_status(self) -> None:
                    return None

                def json(self):
                    return {"status": "ok"}

            with patch("companion_v01.qq_gateway.requests.post", return_value=FakeResponse()) as post:
                result = gateway.send_finance_reports(
                    context,
                    [self._event(report_path)],
                    authorization="finance_subscription_push",
                )

            self.assertTrue(result["ok"])
            self.assertTrue(post.call_args.args[0].endswith("/upload_group_file"))
            payload = post.call_args.kwargs["json"]
            self.assertEqual(payload["group_id"], 20001)
            self.assertEqual(payload["name"], "市场简报.pdf")

    def test_report_delivery_does_not_bypass_finance_authorization(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = Path(temp_dir) / "report.pdf"
            report_path.write_bytes(b"non-empty")
            gateway = NapCatQQGateway()
            context = QQMessageContext(
                should_respond=True,
                reason="private_message",
                target_id=10001,
                user_id=10001,
                session_id="qq_10001",
                profile_user_id="qq_10001",
                finance_mode="off",
                clean_message="发文件",
                raw_message="发文件",
            )

            with patch("companion_v01.qq_gateway.requests.post") as post:
                result = gateway.send_finance_reports(
                    context,
                    [self._event(report_path)],
                    authorization="finance_tool_result",
                )

            self.assertFalse(result["ok"])
            self.assertEqual(result["status"], "blocked")
            post.assert_not_called()

    def test_proactive_adapter_passes_subscription_authorization_to_report_delivery(self) -> None:
        calls = []

        class FakeGateway:
            def send_replies(self, context, messages):
                return {"ok": True, "count": len(messages), "results": [{"ok": True}]}

            def send_finance_reports(self, context, events, *, authorization):
                calls.append((context, events, authorization))
                return {"ok": True, "status": "sent", "count": 1, "results": [{"ok": True}]}

        subscription = FinanceSubscription(
            subscription_id="sub-report",
            client="qq",
            target_id="20001",
            is_group=True,
            session_id="qq_group_shared_20001",
            profile_user_id="qq_group_shared_20001",
            character_pack_id="akane_default",
            finance_mode="push",
            enabled=True,
            filters={"codes": ["000000.TEST"]},
            delivery_policy={"level": "alert"},
        )
        analysis = FinanceAnalysisResult(
            ok=True,
            status="analyzed",
            analysis_id="market_analysis:report",
            messages=("市场快讯",),
            frame={"tool_events": [{"type": "finance_report_ready", "send_to_user": True}]},
        )

        with (
            patch.object(config, "FINANCE_ASSISTANT_ENABLED", True),
            patch.object(config, "QQ_FINANCE_PUSH_ENABLED", True),
            patch.object(config, "QQ_BRIDGE_ENABLED", True),
        ):
            result = QQFinanceDeliveryAdapter(FakeGateway()).deliver(
                subscription=subscription,
                analysis=analysis,
            )

        self.assertTrue(result.ok)
        self.assertEqual(calls[0][2], "finance_subscription_push")
        self.assertEqual(calls[0][0].reason, "finance_subscription_push")


if __name__ == "__main__":
    unittest.main()
