from __future__ import annotations

from datetime import date
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from companion_v01.document_material import MaterialError, read_document_material


class DocumentMaterialTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def docx(self, body, *, extra=None):
        path = self.root / "source.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr(
                "word/document.xml",
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                + body
                + "</w:body></w:document>",
            )
            for key, value in (extra or {}).items():
                archive.writestr(key, value)
        return path

    def test_docx_preserves_body_table_tail_order_and_whitespace(self):
        path = self.docx(
            '<w:p><w:r><w:t xml:space="preserve"> 开头 </w:t></w:r></w:p>'
            "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>0</w:t></w:r></w:p></w:tc>"
            "<w:tc><w:p><w:r><w:t>False</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
            "<w:p><w:r><w:t>结尾</w:t><w:tab/><w:t>保留</w:t><w:br/></w:r></w:p>"
        )
        before = path.read_bytes()
        result = read_document_material(path)
        self.assertTrue(result["complete"])
        self.assertEqual(
            result["blocks"],
            [
                {"kind": "paragraph", "text": " 开头 "},
                {"kind": "table", "rows": [["0", "False"]]},
                {"kind": "paragraph", "text": "结尾\t保留\n"},
            ],
        )
        self.assertEqual(before, path.read_bytes())

    def test_complex_docx_reports_partial_without_claiming_complete(self):
        path = self.docx(
            "<w:p><w:hyperlink><w:r><w:t>链接文字</w:t></w:r></w:hyperlink></w:p>",
            extra={"word/header1.xml": "header", "word/media/image1.png": "image"},
        )
        result = read_document_material(path)
        self.assertFalse(result["complete"])
        self.assertEqual(result["blocks"][0]["text"], "链接文字")
        self.assertEqual(len(result["limitations"]), 2)

    def test_xlsx_reads_every_sheet_tail_and_retains_types_formulas_dates(self):
        from openpyxl import Workbook

        path = self.root / "source.xlsx"
        workbook = Workbook()
        workbook.active.append(["数值", "布尔", "公式", "日期"])
        for index in range(120):
            workbook.active.append([index, False, "=A2+1", date(2026, 9, 6)])
        workbook.create_sheet("第二张").append(["后半段", "  空格  "])
        workbook.save(path)
        workbook.close()
        before = path.read_bytes()
        result = read_document_material(path)
        self.assertTrue(result["complete"], result["limitations"])
        self.assertEqual(len(result["blocks"]), 2)
        rows = result["blocks"][0]["rows"]
        self.assertEqual(len(rows), 121)
        self.assertIs(rows[1][1], False)
        self.assertEqual(rows[1][0], 0)
        self.assertEqual(rows[-1][0], 119)
        self.assertEqual(rows[1][2], {"type": "formula", "expression": "=A2+1"})
        self.assertEqual(rows[1][3], {"type": "datetime", "value": "2026-09-06T00:00:00"})
        self.assertEqual(result["blocks"][1]["rows"], [["后半段", "  空格  "]])
        self.assertEqual(before, path.read_bytes())

    def test_xlsx_merged_data_marked_partial(self):
        from openpyxl import Workbook

        path = self.root / "merged.xlsx"
        workbook = Workbook()
        workbook.active.append(["merged", None])
        workbook.active.merge_cells("A1:B1")
        workbook.save(path)
        workbook.close()
        result = read_document_material(path)
        self.assertFalse(result["complete"])
        self.assertIn("document_workbook_features_not_extracted", result["limitations"])

    def test_csv_preserves_embedded_newlines_blank_rows_and_formula_text(self):
        path = self.root / "source.csv"
        path.write_text('a,b\n0,False\n\n"多行\n值",=A2\n', encoding="utf-8", newline="")
        result = read_document_material(path)
        self.assertEqual(result["blocks"][0]["rows"], [["a", "b"], ["0", "False"], [], ["多行\n值", "=A2"]])
        self.assertTrue(result["complete"])

    def test_size_encoding_and_archive_failures_never_return_preview(self):
        path = self.root / "source.txt"
        for data, reason in (
            (b"\x00binary", "document_text_invalid"),
            (b"\xff", "document_encoding_unsupported"),
            (b"a" * 1_000_001, "document_content_too_large"),
        ):
            path.write_bytes(data)
            with self.assertRaisesRegex(MaterialError, reason):
                read_document_material(path)
        path = self.root / "bad.docx"
        path.write_bytes(b"not a zip; user private content")
        with self.assertRaisesRegex(MaterialError, "^document_archive_invalid$"):
            read_document_material(path)

    def test_plain_text_full_tail_and_roundtrip_json(self):
        path = self.root / "source.md"
        text = "long text\n" * 9000 + "unique tail 末尾\n"
        path.write_text(text, encoding="utf-8", newline="")
        result = read_document_material(path)
        self.assertEqual(result["blocks"][0]["text"], text)
        self.assertEqual(json.loads(json.dumps(result)), result)

    def test_pdf_unreadable_pages_and_encryption_are_not_complete_text(self):
        from pypdf import PdfWriter

        path = self.root / "blank.pdf"
        with PdfWriter() as writer:
            writer.add_blank_page(width=600, height=800)
            writer.write(path)
        result = read_document_material(path)
        self.assertFalse(result["complete"])
        self.assertEqual(result["blocks"], [{"kind": "page", "number": 1, "text": ""}])
        with PdfWriter() as writer:
            writer.add_blank_page(width=600, height=800)
            writer.encrypt("test-only-password")
            writer.write(path)
        with self.assertRaisesRegex(MaterialError, "^document_encrypted_unsupported$"):
            read_document_material(path)


if __name__ == "__main__":
    unittest.main()
