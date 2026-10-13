"""Real attachment readers and continuation behavior, rather than fake text pages."""

from __future__ import annotations

import os
import importlib.util
import types
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from companion_v01.attachment_inbox import AttachmentInboxService
from companion_v01.store import MemoryStore
from companion_v01.tool_handlers.attachments import ReadAttachmentSectionToolHandler
from companion_v01.tool_handlers.core import ToolExecutionContext


class AttachmentSectionPagingTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="akane_section_page_")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.storage = self.root / "attachments"
        self.storage.mkdir()
        self.store = MemoryStore(self.root / "store")
        self.service = AttachmentInboxService(store=self.store, base_dir=self.storage)
        self.handler = ReadAttachmentSectionToolHandler(attachment_service=self.service)

    def _register(self, text: str, *, name: str = "body.txt", encoding: str = "utf-8", preview: str = "") -> Path:
        path = self.storage / name
        path.write_bytes(text.encode(encoding))
        pending = self.service.create_pending(
            profile_user_id="p1", session_id="s1", source="desktop_pet", kind="document",
            origin_name=name, storage_relpath=name, timestamp=100,
        )
        self.service.mark_ready(
            profile_user_id="p1", session_id="s1", attachment_id=pending["attachment_id"],
            summary_title=name, detail={"text_preview": preview, "preview_is_truncated": bool(preview)}, timestamp=110,
        )
        return path

    def _section(self, section: str = "全文", *, max_chars: int = 12000):
        return self.service.read_section(
            profile_user_id="p1", session_id="s1", target="file_001", kind="document",
            section=section, max_chars=max_chars, timestamp=120,
        )

    def _execute(self, call=None, *, session_id="s1", profile_user_id="p1"):
        return self.handler.execute(
            call=call or {"type": "read_attachment_section", "target": "file_001", "section": "全文", "kind": "document"},
            context=ToolExecutionContext(profile_user_id=profile_user_id, session_id=session_id, now_ts=120, visual_payload={}),
        )

    @staticmethod
    def _body(result):
        return result.followup_context.split("本页内容如下：\n", 1)[1].split("\n请基于这段展开内容", 1)[0]

    def test_text_reading_does_not_materialize_the_entire_byte_file(self) -> None:
        self._register("内容\n" * 150_000)
        with patch.object(Path, "read_bytes", side_effect=AssertionError("whole-file read")):
            result = self._section(max_chars=1000)
        self.assertTrue(result["ok"])
        self.assertTrue(result["content"].startswith("内容"))
        self.assertLessEqual(len(result["content"]), 1000)

    def test_explicit_line_range_is_not_silently_reduced_to_81_lines(self) -> None:
        text = "\n".join(f"row-{i}" for i in range(1, 201))
        self._register(text)
        result = self._section("第1-150行")
        self.assertTrue(result["ok"])
        self.assertEqual(result["content"], "\n".join(text.splitlines()[:150]))

    def test_line_selection_can_reach_beyond_the_prefix_budget(self) -> None:
        self._register("padding\n" * 80_000 + "TARGET-LATE\nTARGET-END\n", preview="padding")
        result = self._section("第80001-80002行", max_chars=1000)
        self.assertTrue(result["ok"])
        self.assertEqual(result["content"], "TARGET-LATE\nTARGET-END")

    def test_csv_rows_after_500_are_available_by_range(self) -> None:
        self._register("\n".join(f"row-{i},value-{i}" for i in range(1, 701)), name="body.csv")
        result = self._section("第550-552行")
        self.assertTrue(result["ok"])
        self.assertEqual(result["content"], "row-550 | value-550\nrow-551 | value-551\nrow-552 | value-552")

    def test_real_pages_are_lossless_and_extract_the_section_once(self) -> None:
        text = "头🎉\r\n" + "\r\n".join(f"第{i}行 {'长' * 45}" for i in range(3000))
        self._register(text, preview="头🎉")
        pages = []
        call = None
        with patch.object(self.service, "_read_text_file_section", wraps=self.service._read_text_file_section) as extract:
            for _ in range(30):
                result = self._execute(call)
                pages.append(self._body(result))
                call = result.followup_envelope.continuation
                if call is None:
                    break
            else:
                self.fail("pagination did not finish")
        self.assertEqual("".join(pages), text)
        self.assertGreater(len(pages), 1)
        self.assertTrue(result.followup_envelope.complete)
        self.assertEqual(extract.call_count, 1)

    def test_changed_content_with_identical_file_metadata_invalidates_cursor(self) -> None:
        path = self._register("A" * 120_000)
        first = self._execute()
        stat = path.stat()
        path.write_bytes(b"B" * 120_000)
        os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns))
        result = self._execute(first.followup_envelope.continuation)
        self.assertEqual(result.followup_envelope.diagnostics["status"], "stale_cursor")
        self.assertNotIn("本页内容如下", result.followup_context)

    def test_cursor_cannot_cross_session_or_user(self) -> None:
        self._register("private-text" * 12000)
        first = self._execute()
        with patch.object(self.service, "read_section_page", wraps=self.service.read_section_page) as reader:
            for kwargs in ({"session_id": "s2"}, {"profile_user_id": "p2"}):
                result = self._execute(first.followup_envelope.continuation, **kwargs)
                self.assertEqual(result.followup_envelope.diagnostics["status"], "cursor_invalid")
                self.assertNotIn("private-text", result.followup_context)
            self.assertEqual(reader.call_count, 0)

    def test_deleted_file_cannot_resume_from_cached_text(self) -> None:
        path = self._register("original-text" * 12000, preview="old preview")
        first = self._execute()
        path.unlink()
        result = self._execute(first.followup_envelope.continuation)
        self.assertEqual(result.followup_envelope.diagnostics["status"], "stale_cursor")
        self.assertNotIn("original-text", result.followup_context)

    def test_cleared_material_cannot_resume_from_cached_text(self) -> None:
        self._register("original-text" * 12000)
        first = self._execute()
        self.service.clear_focus(profile_user_id="p1", session_id="s1", target="all", timestamp=130)
        result = self._execute(first.followup_envelope.continuation)
        self.assertEqual(result.followup_envelope.diagnostics["status"], "source_unavailable")

    def test_eviction_and_service_restart_preserve_existing_cursor(self) -> None:
        from companion_v01.attachment_inbox_runtime.section_cache import AttachmentSectionCache
        text = "content" * 20000
        self._register(text)
        self.service._section_cache = AttachmentSectionCache(max_chars=1)
        first = self._execute()
        self.service = AttachmentInboxService(store=self.store, base_dir=self.storage)
        self.handler = ReadAttachmentSectionToolHandler(attachment_service=self.service)
        resumed = self._execute(first.followup_envelope.continuation)
        self.assertEqual(self._body(first) + self._body(resumed), text[:len(self._body(first)) + len(self._body(resumed))])

    def test_4m_window_exhaustion_is_not_reported_as_full_completion(self) -> None:
        self._register("A" * 4_000_001)
        with patch("companion_v01.paged_reading.FIRST_PAGE_BUDGET_CHARS", 4_000_001):
            result = self._execute()
        self.assertEqual(len(self._body(result)), 4_000_000)
        self.assertFalse(result.followup_envelope.complete)
        self.assertIsNone(result.followup_envelope.continuation)
        self.assertTrue(result.followup_envelope.diagnostics["source_truncated"])
        self.assertIn("不能声称已读全文", result.followup_context)

    def test_exactly_4m_characters_can_be_complete(self) -> None:
        self._register("A" * 4_000_000)
        with patch("companion_v01.paged_reading.FIRST_PAGE_BUDGET_CHARS", 4_000_001):
            result = self._execute()
        self.assertTrue(result.followup_envelope.complete)
        self.assertFalse(result.followup_envelope.diagnostics["source_truncated"])

    def test_preview_fallback_is_explicitly_incomplete(self) -> None:
        path = self._register("original", preview="only-preview")
        path.unlink()
        result = self._execute()
        self.assertEqual(self._body(result), "only-preview")
        self.assertFalse(result.followup_envelope.complete)

    def test_source_change_during_extraction_does_not_publish_a_page(self) -> None:
        path = self._register("A" * 100_000)
        original = self.service._read_text_file_section
        def change_source(*args, **kwargs):
            content = original(*args, **kwargs)
            path.write_bytes(b"B" * 100_000)
            return content
        with patch.object(self.service, "_read_text_file_section", side_effect=change_source):
            result = self._execute()
        self.assertEqual(result.followup_envelope.diagnostics["status"], "stale_cursor")
        self.assertNotIn("本页内容如下", result.followup_context)

    def test_partial_utf8_budget_does_not_misdetect_another_encoding(self) -> None:
        self._register("中" * 90_000)
        result = self._section(max_chars=500)
        self.assertEqual(result["content"], "中" * 500)

    def test_deep_line_range_supports_gb18030(self) -> None:
        self._register("padding\r\n" * 40_000 + "中文末尾\r\n第二行", encoding="gb18030")
        result = self._section("第40001-40002行", max_chars=1000)
        self.assertEqual(result["content"], "中文末尾\n第二行")

    def test_csv_range_counts_records_with_embedded_newlines(self) -> None:
        self._register('head,value\n"multi\nline",2\nlast,3\n', name="body.csv")
        result = self._section("第2-3行")
        self.assertEqual(result["content"], "multi\nline | 2\nlast | 3")

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl unavailable")
    def test_excel_explicit_rows_are_not_silently_reduced(self) -> None:
        from openpyxl import Workbook
        path = self._register("", name="body.xlsx")
        workbook = Workbook()
        for index in range(1, 201):
            workbook.active.append([f"row-{index}", index])
        workbook.save(path)
        workbook.close()
        result = self._section("第1-180行")
        self.assertIn("row-180 | 180", result["content"])
        self.assertNotIn("row-181", result["content"])

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl unavailable")
    def test_named_sheet_reads_beyond_the_default_preview_rows_and_columns(self) -> None:
        from openpyxl import Workbook
        path = self._register("", name="body.xlsx")
        workbook = Workbook()
        workbook.active.title = "Sheet1"
        for index in range(1, 151):
            workbook.active.append([f"cell-{index}-{column}" for column in range(1, 46)])
        workbook.save(path)
        workbook.close()
        result = self._execute({"type": "read_attachment_section", "target": "file_001", "section": "Sheet1", "kind": "document"})
        joined = self._body(result)
        while result.followup_envelope.continuation:
            result = self._execute(result.followup_envelope.continuation)
            joined += self._body(result)
        self.assertIn("cell-150-45", joined)
        self.assertTrue(result.followup_envelope.complete)

    def test_default_pdf_preview_is_not_the_entire_document(self) -> None:
        self._register("fixture", name="body.pdf")
        pages = [types.SimpleNamespace(extract_text=lambda i=i: f"page-{i}") for i in range(1, 9)]
        module = types.ModuleType("pypdf")
        module.PdfReader = lambda _path: types.SimpleNamespace(pages=pages)
        with patch.dict("sys.modules", {"pypdf": module}), patch("importlib.util.find_spec", return_value=object()):
            result = self._execute({"type": "read_attachment_section", "target": "file_001", "section": "当前可用片段", "kind": "document"})
        self.assertIn("page-5", self._body(result))
        self.assertNotIn("page-6", self._body(result))
        self.assertFalse(result.followup_envelope.complete)

    def test_pdf_out_of_range_does_not_silently_read_the_last_page(self) -> None:
        self._register("fixture", name="body.pdf")
        pages = [types.SimpleNamespace(extract_text=lambda: "last-page")]
        module = types.ModuleType("pypdf")
        module.PdfReader = lambda _path: types.SimpleNamespace(pages=pages)
        with patch.dict("sys.modules", {"pypdf": module}), patch("importlib.util.find_spec", return_value=object()):
            result = self._section("第999页")
        self.assertFalse(result["ok"])
        self.assertNotIn("last-page", result["content"])

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl unavailable")
    def test_full_workbook_includes_each_sheet_within_the_budget(self) -> None:
        from openpyxl import Workbook
        path = self._register("", name="body.xlsx")
        workbook = Workbook()
        workbook.active.title = "First"
        workbook.active.append(["first-content"])
        workbook.create_sheet("Second").append(["second-content"])
        workbook.save(path)
        workbook.close()
        result = self._execute()
        self.assertIn("first-content", self._body(result))
        self.assertIn("second-content", self._body(result))
        self.assertTrue(result.followup_envelope.complete)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl unavailable")
    def test_sheet_name_takes_priority_over_numbers_inside_its_name(self) -> None:
        from openpyxl import Workbook
        path = self._register("", name="body.xlsx")
        workbook = Workbook()
        workbook.active.title = "Sheet2"
        workbook.active.append(["requested-sheet"])
        workbook.create_sheet("Sheet3").append(["other-sheet"])
        workbook.save(path)
        workbook.close()
        result = self._section("Sheet2")
        self.assertIn("requested-sheet", result["content"])
        self.assertNotIn("other-sheet", result["content"])

    def test_pdf_without_a_text_layer_does_not_report_a_successful_body_read(self) -> None:
        self._register("fixture", name="body.pdf")
        module = types.ModuleType("pypdf")
        module.PdfReader = lambda _path: types.SimpleNamespace(pages=[types.SimpleNamespace(extract_text=lambda: "")])
        with patch.dict("sys.modules", {"pypdf": module}), patch("importlib.util.find_spec", return_value=object()):
            result = self._section("全文")
        self.assertFalse(result["ok"])
        self.assertEqual(result["content"], "")

    def test_missing_original_sheet_does_not_substitute_its_first_preview_row_for_a_deep_range(self) -> None:
        path = self._register("fixture", name="body.xlsx")
        path.unlink()
        item = self.service.resolve_attachment(profile_user_id="p1", session_id="s1", target="file_001")
        self.store.update_attachment_inbox_item(
            profile_user_id="p1", session_id="s1", attachment_id=item["attachment_id"],
            detail={"sheets": [{"name": "Sheet1", "preview_rows": [["only-first-row"]], "row_count": 600}]},
        )
        result = self._section("Sheet1 第550-552行")
        self.assertFalse(result["ok"])
        self.assertNotIn("only-first-row", result["content"])

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx unavailable")
    def test_word_table_rows_after_120_remain_available(self) -> None:
        from docx import Document
        path = self._register("", name="body.docx")
        document = Document()
        table = document.add_table(rows=150, cols=1)
        for index, row in enumerate(table.rows, start=1):
            row.cells[0].text = f"row-{index}"
        document.save(path)
        result = self._section("第1个表")
        self.assertIn("row-150", result["content"])

    def test_pdf_explicit_page_range_is_not_silently_reduced(self) -> None:
        self._register("fixture", name="body.pdf")
        pages = [types.SimpleNamespace(extract_text=lambda i=i: f"page-{i}") for i in range(1, 21)]
        module = types.ModuleType("pypdf")
        module.PdfReader = lambda _path: types.SimpleNamespace(pages=pages)
        with patch.dict("sys.modules", {"pypdf": module}), patch("importlib.util.find_spec", return_value=object()):
            result = self._section("第1-20页")
        self.assertIn("page-20", result["content"])


class SectionCacheTests(unittest.TestCase):
    def test_cache_enforces_capacity_expiry_and_keeps_text_out_of_repr(self) -> None:
        from companion_v01.attachment_inbox_runtime.section_cache import AttachmentSectionCache, SectionSnapshot
        cache = AttachmentSectionCache(max_chars=10, max_entries=2, ttl_seconds=5)
        with patch("companion_v01.attachment_inbox_runtime.section_cache.time.monotonic", return_value=10):
            first = SectionSnapshot("secret", "f", "private-path", False, 10)
            cache.put(("p1", "s1", "a"), first)
            cache.put(("p1", "s1", "b"), SectionSnapshot("bbbbbb", "f", "stamp", False, 10))
            self.assertIsNone(cache.get(("p1", "s1", "a")))
            self.assertIsNone(cache.get(("p2", "s1", "b")))
            self.assertNotIn("secret", repr(first))
            self.assertNotIn("private-path", repr(first))
        with patch("companion_v01.attachment_inbox_runtime.section_cache.time.monotonic", return_value=16):
            self.assertIsNone(cache.get(("p1", "s1", "b")))


if __name__ == "__main__":
    unittest.main()
