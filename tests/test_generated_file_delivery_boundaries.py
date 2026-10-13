from __future__ import annotations

import sqlite3
import tempfile
import unittest
import zipfile
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import Mock, patch

from companion_v01.attachment_inbox import AttachmentInboxService
from companion_v01.generated_files import GeneratedFileService
from companion_v01.paged_reading import NEXT_PAGE_BUDGET_CHARS
from companion_v01.store import MemoryStore
from companion_v01.tool_handlers.generated_media import InspectGeneratedFileToolHandler
from companion_v01.tool_runtime import ToolExecutionContext


class GeneratedFileDeliveryBoundaryTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.store = MemoryStore(self.root / "store")
        self.attachments = AttachmentInboxService(store=self.store, base_dir=self.root / "attachments")
        self.service = GeneratedFileService(
            base_dir=self.root / "generated", store=self.store, attachment_service=self.attachments,
        )
        self.sequence = 0

    def output_path(self, kind):
        self.sequence += 1
        path = self.service.allocate_output_path(
            profile_user_id="user", session_id="session", title=f"fixture-{self.sequence}",
            output_format=kind, timestamp=100 + self.sequence,
        )
        path.parent.mkdir(parents=True, exist_ok=True)
        return path

    def register(self, path):
        return self.service.register_generated_artifact(
            profile_user_id="user", session_id="session", output_path=path,
            output_title=path.stem, output_format=path.suffix.lstrip("."),
            mime_type="application/octet-stream", content_card={}, summary="fixture",
            created_by_tool="fixture.register", timestamp=100 + self.sequence,
        )

    def text(self, value="fixture bytes"):
        path = self.output_path("txt")
        path.write_text(value, encoding="utf-8")
        return self.register(path), path

    def archive(self, members):
        path = self.output_path("zip")
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, content in members.items():
                archive.writestr(name, content)
        return self.register(path)

    def inspect(self, item, section="content", max_chars=500):
        return self.service.inspect_generated_file(
            profile_user_id="user", session_id="session", target=item["generated_handle"],
            section=section, max_chars=max_chars,
        )

    def manage(self, item, action="archive"):
        return self.service.manage_generated_files(
            profile_user_id="user", session_id="session", action=action,
            targets=[item["generated_handle"]], timestamp=200,
        )

    def stored(self, item):
        return self.store.get_generated_file(
            profile_user_id="user", session_id="session", generated_id=item["generated_id"],
        )

    def assert_read_failure(self, result):
        self.assertFalse(result["ok"])
        self.assertEqual(result["error"], "generated_file_read_failed")
        self.assertIsNone(result["inspection"])

    def document(self, kind, first, last):
        path = self.output_path(kind)
        if kind == "docx":
            from docx import Document
            document = Document()
            document.add_paragraph(first)
            document.add_paragraph(last)
            document.save(path)
        elif kind == "xlsx":
            from openpyxl import Workbook
            workbook = Workbook()
            workbook.active.title = "Data"
            workbook.active.append([first])
            workbook.active.append([last])
            workbook.save(path)
            workbook.close()
        else:
            from reportlab.pdfgen.canvas import Canvas
            canvas = Canvas(str(path))
            canvas.drawString(20, 750, first)
            canvas.showPage()
            canvas.drawString(20, 750, last)
            canvas.save()
        return self.register(path), path

    @contextmanager
    def workbook_spy(self, fail_rows=False):
        from openpyxl import load_workbook
        books = []

        def tracked(*args, **kwargs):
            book = load_workbook(*args, **kwargs)
            close = book.close
            book.close = Mock(wraps=close)
            if fail_rows:
                book.worksheets[0].iter_rows = Mock(side_effect=OSError("private reader path"))
            books.append((book, close))
            return book

        with patch("openpyxl.load_workbook", side_effect=tracked):
            try:
                yield books
            finally:
                for _book, close in books:
                    close()

    def test_corrupt_zip_is_a_failed_read_and_a_failed_tool_event(self):
        path = self.output_path("zip")
        path.write_bytes(b"not an archive")
        item = self.register(path)
        self.assert_read_failure(self.inspect(item))
        reader = InspectGeneratedFileToolHandler(generated_file_service=self.service)
        result = reader.execute(
            call={"type": reader.tool_type, "target": item["generated_handle"]},
            context=ToolExecutionContext("user", "session", 100, {}),
        )
        self.assertEqual(result.stream_events[0]["type"], "tool_execution_failed")
        self.assertEqual(result.state_updates["operation_failure"]["reason"], "generated_file_read_failed")

    def test_zip_exception_details_are_not_returned_as_content(self):
        item = self.archive({"README.md": "readme"})
        secret = "private filesystem detail"
        with patch.object(self.service, "_read_zip_member_for_inspection", side_effect=OSError(secret)):
            result = self.inspect(item, "file:README.md")
        self.assertNotIn(secret, str(result))
        self.assert_read_failure(result)

    def test_missing_zip_member_is_not_a_successful_read(self):
        item = self.archive({"README.md": "readme"})
        self.assert_read_failure(self.inspect(item, "file:missing.txt"))

    def test_ambiguous_zip_member_does_not_choose_the_first_file(self):
        item = self.archive({"first/report.txt": "first bytes", "second/report.txt": "second bytes"})
        self.assert_read_failure(self.inspect(item, "file:report.txt"))
        exact = self.inspect(item, "file:second/report.txt")
        self.assertTrue(exact["ok"])
        self.assertEqual(exact["inspection"]["content"], "second bytes")

    def test_missing_document_readers_are_failures_instead_of_document_text(self):
        for kind in ("docx", "xlsx", "pdf"):
            with self.subTest(kind=kind):
                path = self.output_path(kind)
                path.write_bytes(b"document fixture")
                item = self.register(path)
                with patch("importlib.util.find_spec", return_value=None):
                    self.assert_read_failure(self.inspect(item))

    def test_plain_text_prefix_reads_only_its_requested_window(self):
        content = "界abc\n" * 100_000
        _item, path = self.text(content)
        reads = []
        original = Path.open

        class Reader:
            def __init__(self, source):
                self.source = source

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return self.source.__exit__(*args)

            def read(self, size=-1):
                value = self.source.read(size)
                reads.append((size, len(value)))
                return value

        def tracked(candidate, *args, **kwargs):
            source = original(candidate, *args, **kwargs)
            return Reader(source) if candidate == path else source

        with patch.object(Path, "open", tracked):
            result = self.service._read_plain_text_file(path, max_chars=500)
        self.assertEqual(result, content[:501])
        self.assertTrue(reads)
        self.assertTrue(all(0 <= requested <= 501 and count <= 501 for requested, count in reads))

    def test_bounded_text_preserves_bom_legacy_encoding_and_newlines(self):
        for encoding in ("utf-8-sig", "gb18030"):
            with self.subTest(encoding=encoding):
                path = self.output_path("txt")
                source = "中文🙂\r\n" * 200
                path.write_bytes(source.encode(encoding))
                expected = source.replace("\r\n", "\n")
                self.assertEqual(self.service._read_plain_text_file(path, max_chars=500), expected[:501])

    def test_zip_preview_does_not_expand_the_entire_member(self):
        source = "界" * 400_000
        item = self.archive({"large.txt": source})
        read_bytes = []
        read = zipfile.ZipExtFile.read
        read1 = zipfile.ZipExtFile.read1

        def tracked_read(handle, *args):
            value = read(handle, *args)
            read_bytes.append(len(value))
            return value

        def tracked_read1(handle, *args):
            value = read1(handle, *args)
            read_bytes.append(len(value))
            return value

        with patch.object(zipfile.ZipExtFile, "read", tracked_read), patch.object(zipfile.ZipExtFile, "read1", tracked_read1):
            result = self.inspect(item, "file:large.txt")
        self.assertTrue(result["ok"])
        self.assertEqual(result["inspection"]["content"], source[:500])
        self.assertTrue(result["inspection"]["truncated"])
        self.assertLess(sum(read_bytes), 32_000)

    def test_plain_text_tail_reaches_the_actual_end(self):
        item, _path = self.text("prefix bytes\n" * 2000 + "actual-tail-marker")
        result = self.inspect(item, "tail")
        self.assertTrue(result["ok"])
        self.assertIn("actual-tail-marker", result["inspection"]["content"])
        self.assertTrue(result["inspection"]["truncated"])

    def test_document_tail_reaches_the_actual_end(self):
        for kind in ("docx", "xlsx", "pdf"):
            with self.subTest(kind=kind):
                item, _path = self.document(kind, "prefix" * 1000, "actual-tail-marker")
                result = self.inspect(item, "tail")
                self.assertTrue(result["ok"])
                self.assertIn("actual-tail-marker", result["inspection"]["content"])

    def test_document_paging_does_not_stop_at_an_exact_extraction_boundary(self):
        for kind in ("docx", "xlsx", "pdf"):
            with self.subTest(kind=kind):
                boundary = 500 + NEXT_PAGE_BUDGET_CHARS
                padding = {"docx": 0, "xlsx": len("## Sheet: Data\n|  |"), "pdf": len("## Page 1\n")}[kind]
                prefix = ("a" if kind == "pdf" else "界") * (boundary - padding)
                item, _path = self.document(kind, prefix, "actual-tail-marker")
                expected = {
                    "docx": prefix + "\nactual-tail-marker",
                    "xlsx": f"## Sheet: Data\n| {prefix} |\n| actual-tail-marker |",
                    "pdf": f"## Page 1\n{prefix}\n## Page 2\nactual-tail-marker",
                }[kind]
                reader = InspectGeneratedFileToolHandler(generated_file_service=self.service)
                context = ToolExecutionContext("user", "session", 100, {})
                call = {"type": reader.tool_type, "target": item["generated_handle"], "max_chars": 500}
                pieces = []
                for _ in range(5):
                    result = reader.execute(call=reader.normalize_call(call), context=context)
                    self.assertEqual(result.stream_events[0]["type"], "generated_file_inspected")
                    pieces.append(result.stream_events[0]["inspection"]["content"])
                    if result.followup_envelope.complete:
                        break
                    call = dict(result.followup_envelope.continuation)
                else:
                    self.fail("document paging did not finish")
                self.assertEqual("".join(pieces), expected)

    def test_workbook_closes_after_a_complete_read(self):
        _item, path = self.document("xlsx", "first", "last")
        with self.workbook_spy() as books:
            self.service._read_generated_xlsx(path=path, max_chars=40000)
            books[0][0].close.assert_called_once()

    def test_workbook_closes_after_a_prefix_read(self):
        _item, path = self.document("xlsx", "first" * 1000, "last")
        with self.workbook_spy() as books:
            self.service._read_generated_xlsx(path=path, max_chars=500)
            books[0][0].close.assert_called_once()

    def test_workbook_closes_when_row_reading_fails(self):
        item, _path = self.document("xlsx", "first", "last")
        with self.workbook_spy(fail_rows=True) as books:
            result = self.inspect(item)
            self.assert_read_failure(result)
            books[0][0].close.assert_called_once()

    def test_pdf_source_closes_after_prefix_and_failed_reads(self):
        item, path = self.document("pdf", "first" * 1000, "last")
        original = Path.open
        sources = []

        def tracked(candidate, *args, **kwargs):
            source = original(candidate, *args, **kwargs)
            if candidate == path:
                sources.append(source)
            return source

        with patch.object(Path, "open", tracked):
            self.assertTrue(self.inspect(item)["ok"])
            self.assertTrue(sources)
            self.assertTrue(all(source.closed for source in sources))
            sources.clear()
            with patch("pypdf._page.PageObject.extract_text", side_effect=OSError("private PDF detail")):
                result = self.inspect(item)
            self.assert_read_failure(result)
            self.assertTrue(sources)
            self.assertTrue(all(source.closed for source in sources))

    def test_png_embedded_text_paging_keeps_the_final_metadata_text(self):
        from PIL import Image, PngImagePlugin
        path = self.output_path("png")
        source = "界" * (500 + NEXT_PAGE_BUDGET_CHARS) + "\nactual-tail-marker"
        metadata = PngImagePlugin.PngInfo()
        metadata.add_text("Description", source)
        with Image.new("RGB", (1, 1), "white") as image:
            image.save(path, pnginfo=metadata)
        item = self.register(path)
        reader = InspectGeneratedFileToolHandler(generated_file_service=self.service)
        context = ToolExecutionContext("user", "session", 100, {})
        call = {"type": reader.tool_type, "target": item["generated_handle"], "max_chars": 500}
        pieces = []
        for _ in range(5):
            result = reader.execute(call=reader.normalize_call(call), context=context)
            self.assertEqual(result.stream_events[0]["type"], "generated_file_inspected")
            pieces.append(result.stream_events[0]["inspection"]["content"])
            if result.followup_envelope.complete:
                break
            call = dict(result.followup_envelope.continuation)
        else:
            self.fail("embedded text paging did not finish")
        self.assertEqual("".join(pieces), source)
        self.assertIn("actual-tail-marker", self.inspect(item, "tail")["inspection"]["content"])

    def test_empty_record_update_is_not_a_successful_management_result(self):
        for action in ("archive", "delete", "purge"):
            with self.subTest(action=action):
                item, path = self.text()
                with patch.object(self.store, "update_generated_file", return_value=None):
                    result = self.manage(item, action)
                self.assertFalse(result["ok"])
                self.assertEqual(result["status"], "failed")
                self.assertEqual(result["managed"], [])
                self.assertEqual(result["failures"][0]["code"], "record_update_failed")
                self.assertEqual(result["failures"][0]["file_deleted"], action != "archive")
                self.assertEqual(path.exists(), action == "archive")
                self.assertEqual(self.stored(item)["status"], "ready")
                retry = self.manage(item, action)
                self.assertTrue(retry["ok"])
                if action != "archive":
                    self.assertTrue(retry["managed"][0]["file_already_absent"])

    def test_management_database_error_keeps_a_truthful_retryable_result(self):
        item, path = self.text()
        with patch.object(self.store, "update_generated_file", side_effect=sqlite3.OperationalError("private database detail")):
            result = self.manage(item)
        self.assertFalse(result["ok"])
        self.assertEqual(result["failures"][0]["code"], "record_update_failed")
        self.assertEqual(result["failures"][0]["cause"], "OperationalError")
        self.assertNotIn("private database detail", str(result))
        self.assertTrue(path.exists())
        self.assertEqual(self.stored(item)["status"], "ready")

    def test_management_batch_keeps_success_and_failure_separate(self):
        first, first_path = self.text("first")
        second, _second_path = self.text("second")
        update = self.store.update_generated_file

        def selective(**kwargs):
            return None if kwargs["generated_id"] == first["generated_id"] else update(**kwargs)

        with patch.object(self.store, "update_generated_file", side_effect=selective):
            result = self.service.manage_generated_files(
                profile_user_id="user", session_id="session", action="archive",
                targets=[first["generated_handle"], second["generated_handle"]], timestamp=200,
            )
        self.assertTrue(result["ok"])
        self.assertEqual(result["status"], "partial")
        self.assertEqual([item["generated_id"] for item in result["managed"]], [second["generated_id"]])
        self.assertEqual(result["failures"][0]["generated_id"], first["generated_id"])
        self.assertTrue(first_path.exists())
        self.assertEqual(self.stored(first)["status"], "ready")

    def test_file_disappearing_before_unlink_can_still_archive_its_record(self):
        item, path = self.text()
        original = Path.unlink

        def disappears(candidate, *args, **kwargs):
            if candidate == path:
                original(candidate)
            return original(candidate, *args, **kwargs)

        with patch.object(Path, "unlink", disappears):
            result = self.manage(item, "delete")
        self.assertTrue(result["ok"])
        self.assertTrue(result["managed"][0]["file_already_absent"])
        self.assertFalse(result["managed"][0]["file_deleted"])
        self.assertEqual(self.stored(item)["status"], "removed")

    def test_file_disappearing_during_resolution_returns_missing_on_disk(self):
        item, path = self.text()
        item["file_size"] = 0
        original = Path.is_file

        def disappears(candidate):
            is_file = original(candidate)
            if is_file and candidate == path:
                path.unlink()
            return is_file

        with patch.object(Path, "is_file", disappears):
            result = self.service._generated_item_to_sendable_file(item, timestamp=200)
        self.assertEqual(result, (None, "missing_on_disk"))

    def test_attachment_disappearing_during_resolution_returns_missing_on_disk(self):
        path = self.attachments.base_dir / "user" / "session" / "input.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("original attachment bytes", encoding="utf-8")
        item = self.store.add_attachment_inbox_item(
            profile_user_id="user", session_id="session", kind="file", status="ready",
            origin_name="input.txt", storage_relpath="user/session/input.txt", file_size=0,
            timestamp=100,
        )
        original = Path.is_file

        def disappears(candidate):
            is_file = original(candidate)
            if is_file and candidate == path:
                path.unlink()
            return is_file

        with patch.object(Path, "is_file", disappears):
            result = self.service._attachment_item_to_sendable_file(item)
        self.assertEqual(result, (None, "missing_on_disk"))


if __name__ == "__main__":
    unittest.main()
