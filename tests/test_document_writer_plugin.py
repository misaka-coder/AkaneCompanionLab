"""Build/install the actual document wheel; invoke through host projection."""

import asyncio
import json
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

from capcore import InvocationContext
from companion_v01.engine import AkaneMemoryEngine
from companion_v01.background_tasks import BackgroundTaskRunner
from companion_v01.capability_registry import ExecutorBroker
from companion_v01.host_jobs import HostJobOwner, HostJobStore
from companion_v01.host_tool_jobs import HostToolJobRuntime
from companion_v01.mode_profiles import ModeProfileRegistry
from companion_v01.native_tool_schema import build_openai_native_tool_specs
from companion_v01.tool_runtime import ToolExecutionContext
from tests.dependency_supply_fixture import wheelhouse_for_plugins
from tests.image_plugin_harness import ImageHarness
from tests import test_media_convert_plugin as delivery_acceptance

PLUGIN_ID = "akane.document-writer"
CAPABILITIES = {name: PLUGIN_ID + f".{name}.v1" for name in ("compose", "revise", "style")}
DEPENDENCY_WHEELHOUSE = wheelhouse_for_plugins("akane_document_writer")


class DocumentHarness(ImageHarness):
    plugin_id = PLUGIN_ID
    capability_id = CAPABILITIES["compose"]

    def __init__(self, root):
        super().__init__(root, None, dependency_wheelhouse=DEPENDENCY_WHEELHOUSE)

    async def document(self, operation="compose", **options):
        capability = CAPABILITIES[operation]
        handler = self.handlers()[capability]
        return await asyncio.to_thread(
            handler.execute,
            call=handler.normalize_call({"type": capability, **options}),
            context=ToolExecutionContext("owner", "session", 1, {}, client_mode="desktop_pet"),
        )

    def record(self, handle):
        return self.files.resolve_generated_artifact(profile_user_id="owner", session_id="session", target=handle)


class DocumentPluginTests(unittest.IsolatedAsyncioTestCase):
    async def test_installed_job_duplicate_cancel_disable_and_delivery(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            hooks, scratch = root / "hooks", root / "scratch"
            hooks.mkdir()
            scratch.mkdir()
            marker = root / "parser-child-pid.txt"
            # Explicit test-only startup barrier in the real owned child. It
            # does not replace rendering or fabricate a successful operation.
            (hooks / "sitecustomize.py").write_text(
                "import json,os,sys,time\nfrom pathlib import Path\n"
                "if len(sys.argv)==2 and Path(sys.argv[0]).name=='worker.py' and Path(sys.argv[0]).parent.name=='akane_document_writer':\n"
                " p=Path(sys.argv[1])\n"
                " if p.is_file() and json.loads(p.read_text(encoding='utf-8')).get('output_title')=='取消边界':\n"
                "  Path(os.environ['DOCUMENT_TEST_BARRIER']).write_text(str(os.getpid()))\n"
                "  time.sleep(120)\n",
                encoding="utf-8",
            )
            environment = {
                "PYTHONPATH": str(hooks) + os.pathsep + os.environ.get("PYTHONPATH", ""),
                "DOCUMENT_TEST_BARRIER": str(marker),
                "TEMP": str(scratch),
                "TMP": str(scratch),
                "TMPDIR": str(scratch),
            }
            with patch.dict(os.environ, environment):
                harness = await DocumentHarness(root).start()
                await harness.install()
                engine = harness.engine
                engine.executor_broker = ExecutorBroker(None)
                engine._tool_hook_result_status = AkaneMemoryEngine._tool_hook_result_status
                engine._resolve_client_protocol_context = ModeProfileRegistry().resolve_from_payload
                jobs, runner = HostJobStore(root / "jobs.db"), BackgroundTaskRunner({"host-jobs": 1})
                completed = []
                runtime = HostToolJobRuntime(
                    engine=engine,
                    store=jobs,
                    background_tasks=runner,
                    conversation_ref_issuer=lambda _: "document-test-conversation",
                    terminal_callback=lambda job: completed.append(job) or True,
                )
                owner = HostJobOwner("owner", "session")

                def submit(invocation, **options):
                    result = runtime.submit(
                        capability_id=CAPABILITIES["compose"],
                        invocation_id=invocation,
                        call={
                            "type": CAPABILITIES["compose"],
                            "arguments": {"output_format": "docx", "content_markdown": "真实文档正文与结尾", **options},
                        },
                        context=ToolExecutionContext("owner", "session", 1, {}, client_mode="qq_text"),
                    )
                    return result.stream_events[0]["job_id"]

                async def settle():
                    self.assertTrue(await asyncio.to_thread(runner.wait_idle, lane="host-jobs", timeout=30))
                    self.assertTrue(await asyncio.to_thread(runner.wait_idle, lane="host-job-completions", timeout=5))

                try:
                    started = time.monotonic()
                    success = submit("document-success")
                    self.assertLess(time.monotonic() - started, 1)
                    await settle()
                    self.assertEqual(jobs.get(success, owner=owner).status, "succeeded")
                    self.assertEqual(len(jobs.get(success, owner=owner).artifacts), 1)
                    self.assertEqual(submit("document-success"), success)
                    cancelled = submit("document-cancel", output_title="取消边界")
                    for _ in range(500):
                        if marker.is_file():
                            break
                        await asyncio.sleep(0.01)
                    self.assertTrue(marker.is_file(), jobs.get(cancelled, owner=owner))
                    child_pid = int(marker.read_text())
                    self.assertTrue(self.child_running(child_pid))
                    disabling = asyncio.create_task(harness.service.set_enabled(plugin_id=PLUGIN_ID, enabled=False))
                    self.assertTrue(jobs.request_cancel(cancelled, owner=owner)["ok"])
                    self.assertTrue(jobs.request_cancel(cancelled, owner=owner)["ok"])
                    self.assertTrue((await disabling)["ok"])
                    await settle()
                    self.assertEqual(jobs.get(cancelled, owner=owner).status, "cancelled")
                    self.assertFalse(jobs.get(cancelled, owner=owner).artifacts)
                    self.assertFalse(self.child_running(child_pid))
                    self.assertFalse(list(scratch.glob("akane-plugin-work-*")))
                    self.discovery(harness, False)
                    self.assertTrue((await harness.service.set_enabled(plugin_id=PLUGIN_ID, enabled=True))["ok"])
                    failed = submit("document-invalid-json", output_format="json", content_markdown="invalid json")
                    await settle()
                    self.assertEqual(jobs.get(failed, owner=owner).status, "failed")
                    self.assertEqual(jobs.get(failed, owner=owner).last_error, "document_json_invalid")
                    self.assertEqual([job.job_id for job in completed], [success, cancelled, failed])
                    self.assertEqual(runtime.recover(), 0)
                    self.assertEqual(
                        len(
                            harness.files.store.list_generated_files(
                                profile_user_id="owner", session_id="session", limit=100
                            )
                        ),
                        1,
                    )
                    delivery_acceptance.MediaPluginInstallationTests.verify_job_memory_and_delivery(
                        self, root=root, files=harness.files, completed=completed, expected_extension="docx"
                    )
                finally:
                    await asyncio.to_thread(runner.close, timeout=10)
                    await harness.close()

    @staticmethod
    def child_running(pid):
        if os.name != "nt":
            return Path(f"/proc/{pid}").exists()
        import ctypes
        from ctypes import wintypes

        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.OpenProcess.argtypes = (wintypes.DWORD, wintypes.BOOL, wintypes.DWORD)
        kernel.OpenProcess.restype = wintypes.HANDLE
        kernel.GetExitCodeProcess.argtypes = (wintypes.HANDLE, ctypes.POINTER(wintypes.DWORD))
        kernel.CloseHandle.argtypes = (wintypes.HANDLE,)
        handle = kernel.OpenProcess(0x1000, False, pid)
        if not handle:
            return False
        try:
            code = wintypes.DWORD()
            return bool(kernel.GetExitCodeProcess(handle, ctypes.byref(code))) and code.value == 259
        finally:
            kernel.CloseHandle(handle)

    def discovery(self, harness, installed):
        result = []
        for mode in ("desktop_pet", "qq_text"):
            client = ModeProfileRegistry().resolve_from_payload({"client_mode": mode})
            handlers = harness.engine._resolve_tool_handlers(
                client_context=client, profile_user_id="owner", session_id="session"
            )
            prompt = AkaneMemoryEngine._build_tool_prompt_context(
                harness.engine,
                allow_tool_call=True,
                client_context=client,
                profile_user_id="owner",
                session_id="session",
            )
            tools = build_openai_native_tool_specs(handlers)
            schema = json.dumps(tools, sort_keys=True)
            if installed:
                for operation in ("compose", "revise"):
                    tool = next(item for item in tools if item["_akane_capability_id"] == CAPABILITIES[operation])
                    row = tool["function"]["parameters"]["properties"]["table_rows"]["items"]
                    self.assertEqual(row["type"], "array")
                    self.assertEqual(row["minItems"], 1)
                    self.assertEqual(row["maxItems"], 50)
                    self.assertEqual(set(row["items"]["type"]), {"string", "number", "boolean", "null"})
            for retired in ("compose_file", "revise_generated_file", "apply_style_to_existing_file"):
                self.assertNotIn(retired, handlers)
                self.assertNotIn(retired, prompt)
                self.assertNotIn(retired, schema)
            for capability in CAPABILITIES.values():
                self.assertEqual(capability in handlers, installed)
                self.assertEqual(capability in prompt, installed)
                self.assertEqual(capability in schema, installed)
            result.append(schema)
        return result

    def output(self, harness, result, *, sent=False):
        self.assertEqual(result.state_updates["adapter_capability_status"], "ok", result)
        self.assertEqual(result.state_updates["plugin_result_experience"], "projected")
        events = [event for event in result.stream_events if event["type"] == "generated_file_ready"]
        self.assertEqual(len(events), 1)
        self.assertIs(events[0]["send_to_user"], sent)
        record = harness.record(events[0]["generated_file"]["generated_handle"])
        self.assertTrue(Path(record["absolute_path"]).is_file())
        self.assertNotIn(str(harness.root), result.followup_context)
        self.assertFalse(list((harness.root / "workers").glob("**/outbox/**/*.*")))
        if (harness.root / "copies").exists():
            self.assertFalse(list((harness.root / "copies").iterdir()))
        return record

    async def test_installed_formats_sources_revisions_styles_and_lifecycle(self):
        from docx import Document
        from openpyxl import Workbook, load_workbook
        from pypdf import PdfReader

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            harness = await DocumentHarness(root).start()
            try:
                self.discovery(harness, False)
                await harness.install()
                baseline = self.discovery(harness, True)
                text_samples = {
                    "txt": "完整正文\n结尾",
                    "md": "# 标题\n\n原文末尾",
                    "html": "<p>完整页面</p>",
                    "json": '{"value":0,"ready":false}',
                    "srt": "1\n00:00:01,000 --> 00:00:02,000\n正文\n",
                    "lrc": "[00:01.00]歌词\n",
                    "vtt": "WEBVTT\n\n00:01.000 --> 00:02.000\n正文\n",
                }
                records = {}
                for fmt, content in text_samples.items():
                    result = await harness.document(
                        output_format=fmt, output_title="文档验收", content_markdown=content
                    )
                    record = self.output(harness, result)
                    self.assertEqual(Path(record["absolute_path"]).read_bytes(), content.encode("utf-8"))
                    inspection = harness.files.inspect_generated_file(
                        profile_user_id="owner",
                        session_id="session",
                        target=record["generated_handle"],
                        section="content",
                        max_chars=2000,
                    )
                    self.assertTrue(inspection["ok"], inspection)
                    self.assertEqual(inspection["inspection"]["source_kind"], fmt)
                    self.assertEqual(inspection["inspection"]["content"].strip(), content.strip())
                    records[fmt] = record
                rows = [["项目", "数值", "状态"], ["中文长项目", 0, False], ["**原样符号**", -2.5, "=1+1"]]
                for fmt in ("csv", "xlsx", "docx", "pdf"):
                    result = await harness.document(output_format=fmt, output_title="数据验收", table_rows=rows)
                    record = self.output(harness, result)
                    records[fmt] = record
                    if fmt == "docx":
                        self.assertEqual(Document(record["absolute_path"]).tables[0].cell(2, 0).text, "**原样符号**")
                    elif fmt == "xlsx":
                        book = load_workbook(record["absolute_path"])
                        try:
                            self.assertEqual(book.active["B2"].value, 0)
                            self.assertIs(book.active["C2"].value, False)
                            self.assertEqual(book.active["C3"].data_type, "s")
                        finally:
                            book.close()
                    elif fmt == "pdf":
                        with PdfReader(record["absolute_path"]) as reader:
                            self.assertIn("中文长项目", reader.pages[0].extract_text())
                original = records["docx"]
                before = Path(original["absolute_path"]).read_bytes()
                revised = self.output(
                    harness,
                    await harness.document(
                        "revise",
                        target=original["generated_handle"],
                        content_markdown="# 修订说明\n完整替换正文和结尾",
                        send_to_user=True,
                    ),
                    sent=True,
                )
                self.assertEqual(revised["version_no"], 2)
                self.assertEqual(revised["version_of_generated_id"], original["generated_id"])
                self.assertIn(original["generated_id"], revised["source_ids"])
                self.assertEqual(Path(original["absolute_path"]).read_bytes(), before)
                self.assertEqual(harness.record(original["generated_handle"])["delivery_status"], "not_requested")
                styled = self.output(
                    harness,
                    await harness.document(
                        "style",
                        target=revised["generated_handle"],
                        formatting={"paragraphs": [{"contains": "完整替换", "font_color": "FF0000"}]},
                    ),
                )
                self.assertEqual(styled["version_no"], 3)
                self.assertEqual(styled["version_of_generated_id"], revised["generated_id"])
                self.assertEqual(
                    [p.text for p in Document(styled["absolute_path"]).paragraphs],
                    [p.text for p in Document(revised["absolute_path"]).paragraphs],
                )
                self.assertEqual(str(Document(styled["absolute_path"]).paragraphs[-1].runs[0].font.color.rgb), "FF0000")
                # Uploaded originals use the same resource contract and keep their bytes.
                (root / "attachments").mkdir(exist_ok=True)
                source = root / "attachments" / "source.txt"
                content = "完整原文\n" * 9000 + "唯一末尾\n"
                source.write_bytes(content.encode("utf-8"))
                handle = harness.register_image(source, kind="file", mime="text/plain")
                full = self.output(
                    harness, await harness.document(source_ids=[handle], source_mode="full_text", output_format="txt")
                )
                self.assertEqual(Path(full["absolute_path"]).read_bytes(), source.read_bytes())
                regenerated = self.output(
                    harness,
                    await harness.document(
                        source_ids=[full["generated_handle"]], source_mode="full_text", output_format="txt"
                    ),
                )
                self.assertEqual(Path(regenerated["absolute_path"]).read_bytes(), source.read_bytes())
                self.assertIn(full["generated_id"], regenerated["source_ids"])
                workbook_path = root / "attachments" / "source.xlsx"
                workbook = Workbook()
                workbook.active.title = "第一张"
                workbook.active.append(["值", "公式"])
                workbook.active.append([0, "=A2+1"])
                workbook.create_sheet("第二张").append(["不应变更"])
                workbook.save(workbook_path)
                workbook.close()
                workbook_bytes = workbook_path.read_bytes()
                workbook_handle = harness.register_image(
                    workbook_path, kind="file", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                workbook_styled = self.output(
                    harness,
                    await harness.document(
                        "style",
                        target=workbook_handle,
                        formatting={
                            "sheet_name": "第一张",
                            "cells": [{"row": 2, "column_index": 1, "font_color": "FF0000"}],
                        },
                    ),
                )
                book = load_workbook(workbook_styled["absolute_path"])
                try:
                    self.assertEqual(book["第一张"]["B2"].value, "=A2+1")
                    self.assertEqual(book["第一张"]["B2"].data_type, "f")
                    self.assertEqual(str(book["第一张"]["A2"].font.color.rgb), "00FF0000")
                    self.assertEqual(book["第二张"]["A1"].value, "不应变更")
                finally:
                    book.close()
                self.assertEqual(workbook_path.read_bytes(), workbook_bytes)
                complex_path = root / "attachments" / "complex.docx"
                complex_document = Document()
                complex_document.add_paragraph("正文")
                complex_document.sections[0].header.paragraphs[0].text = "页眉不可遗漏"
                complex_document.save(complex_path)
                complex_handle = harness.register_image(
                    complex_path,
                    kind="file",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                failures = (
                    ("compose", {"output_format": "md"}),
                    ("compose", {"output_format": "json", "content_markdown": "not JSON"}),
                    ("compose", {"output_format": "xlsx", "content_markdown": "would be dropped", "table_rows": [[1]]}),
                    (
                        "compose",
                        {
                            "output_format": "txt",
                            "source_ids": [handle, "file_missing"],
                            "content_markdown": "not silently accepted",
                        },
                    ),
                    ("compose", {"output_format": "txt", "source_ids": [complex_handle], "source_mode": "full_text"}),
                    ("compose", {"output_format": "xlsx", "source_ids": [workbook_handle], "source_mode": "full_text"}),
                    ("revise", {"target": original["generated_handle"]}),
                    ("revise", {"target": handle, "content_markdown": "not a generated file"}),
                    ("style", {"target": revised["generated_handle"]}),
                    ("style", {"target": handle, "formatting": {"header": {"bold": True}}}),
                    ("style", {"target": workbook_handle, "formatting": {"header": {"bold": True}}}),
                    (
                        "style",
                        {
                            "target": revised["generated_handle"],
                            "formatting": {"paragraphs": [{"paragraph_index": 999, "bold": True}]},
                        },
                    ),
                )
                before_failures = harness.files.store.list_generated_files(
                    profile_user_id="owner", session_id="session", limit=100
                )
                for operation, options in failures:
                    with self.subTest(operation=operation, options=options):
                        failure = await harness.document(operation, **options)
                        self.assertNotEqual(failure.state_updates["adapter_capability_status"], "ok", failure)
                        self.assertFalse(any(e["type"] == "generated_file_ready" for e in failure.stream_events))
                        self.assertNotIn(str(root), failure.followup_context)
                self.assertEqual(
                    harness.files.store.list_generated_files(profile_user_id="owner", session_id="session", limit=100),
                    before_failures,
                )
                for context in (
                    InvocationContext("other", "session", "desktop_pet"),
                    InvocationContext("owner", "other", "desktop_pet"),
                ):
                    denied = await harness.runtime.invoke(
                        CAPABILITIES["compose"],
                        {"source_ids": [handle], "source_mode": "full_text", "output_format": "txt"},
                        context=context,
                    )
                    self.assertEqual(denied.reason, "resource_not_found")
                self.assertEqual(self.discovery(harness, True), baseline)
                self.assertTrue((await harness.service.set_enabled(plugin_id=PLUGIN_ID, enabled=False))["ok"])
                self.discovery(harness, False)
                self.assertTrue((await harness.service.set_enabled(plugin_id=PLUGIN_ID, enabled=True))["ok"])
                self.discovery(harness, True)
                self.assertTrue((await harness.service.uninstall(plugin_id=PLUGIN_ID))["ok"])
                self.discovery(harness, False)
            finally:
                await harness.close()


if __name__ == "__main__":
    unittest.main()
