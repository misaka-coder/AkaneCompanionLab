from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
ASSET_DIR = OUT_DIR / "assets"
APPARATUS = ASSET_DIR / "apparatus_crop.jpg"
RESULT = ASSET_DIR / "result_observation.jpg"
SHEET = ASSET_DIR / "principle_sheet_rotated.jpg"
PPTX = OUT_DIR / "磁混沌摆_五分钟讲解PPT_初稿.pptx"

BG = RGBColor(248, 250, 252)
DARK = RGBColor(26, 32, 44)
MUTED = RGBColor(82, 95, 112)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(13, 148, 136)
ORANGE = RGBColor(234, 88, 12)
PURPLE = RGBColor(124, 58, 237)
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = RGBColor(219, 234, 254)
LIGHT_ORANGE = RGBColor(255, 237, 213)
LIGHT_TEAL = RGBColor(204, 251, 241)
FONT = "Microsoft YaHei"


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=24,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, 0.7, 0.45, 9.5, 0.65, 30, DARK, True)
    if subtitle:
        add_text(slide, subtitle, 0.72, 1.08, 11.2, 0.35, 13, MUTED)


def add_card(slide, x, y, w, h, title, body, accent=BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(226, 232, 240)
    shape.line.width = Pt(1)
    add_text(slide, title, x + 0.22, y + 0.18, w - 0.44, 0.32, 15, accent, True)
    add_text(slide, body, x + 0.22, y + 0.58, w - 0.44, h - 0.68, 14, DARK)
    return shape


def add_chip(slide, text, x, y, w=4.3, color=BLUE, fill_color=LIGHT_BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.4)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = color
    return shape


def add_image_fit(slide, path, x, y, w, h, border=True):
    img = Image.open(path)
    iw, ih = img.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        pw = w
        ph = w / img_ratio
        px = x
        py = y + (h - ph) / 2
    else:
        ph = h
        pw = h * img_ratio
        px = x + (w - pw) / 2
        py = y
    if border:
        rect = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = WHITE
        rect.line.color.rgb = RGBColor(203, 213, 225)
    return slide.shapes.add_picture(
        str(path), Inches(px), Inches(py), width=Inches(pw), height=Inches(ph)
    )


def add_footer(slide, n):
    add_text(
        slide,
        f"{n}/7  磁混沌摆五分钟讲解",
        10.55,
        7.08,
        2.1,
        0.25,
        9,
        RGBColor(100, 116, 139),
        align=PP_ALIGN.RIGHT,
    )


def add_connector(slide, x1, y1, x2, y2, color, width=2):
    con = slide.shapes.add_connector(
        1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    con.line.color.rgb = color
    con.line.width = Pt(width)
    return con


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(15, 23, 42))
    add_image_fit(slide, APPARATUS, 6.7, 0.8, 5.8, 4.2, border=False)
    add_text(slide, "实验 27\n磁混沌摆", 0.75, 1.15, 5.3, 1.6, 42, WHITE, True)
    add_text(
        slide,
        "一个小球为什么会走出复杂轨迹，最后停在某一个磁铁上？",
        0.8,
        3.05,
        5.4,
        0.6,
        20,
        RGBColor(226, 232, 240),
    )
    add_chip(slide, "5分钟讲解 + 实验演示", 0.8, 4.0, 2.5, TEAL, LIGHT_TEAL)
    add_text(
        slide,
        "结构：现象观察 -> 实验装置 -> 原理解释 -> 结果与结论",
        0.8,
        4.6,
        5.6,
        0.7,
        15,
        RGBColor(203, 213, 225),
    )
    add_text(
        slide,
        "1/7",
        12.1,
        7.05,
        0.5,
        0.25,
        9,
        RGBColor(148, 163, 184),
        align=PP_ALIGN.RIGHT,
    )

    # 2
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(
        slide,
        "先看现象：运动很复杂",
        "把摆锤从偏离平衡的位置释放，它不会像普通单摆那样规则往复。",
    )
    add_image_fit(slide, APPARATUS, 0.75, 1.55, 5.7, 3.4)
    add_card(
        slide,
        6.8,
        1.55,
        5.55,
        1.05,
        "观察到什么？",
        "摆锤会被三个磁铁共同影响，轨迹不断弯曲、折返、改变方向。",
        BLUE,
    )
    add_card(
        slide,
        6.8,
        2.8,
        5.55,
        1.05,
        "最后会怎样？",
        "由于空气阻力和摩擦消耗能量，运动逐渐变慢，最终停在某个磁极上方。",
        TEAL,
    )
    add_card(
        slide,
        6.8,
        4.05,
        5.55,
        1.05,
        "这一点很关键",
        "过程看起来“不规则”，但不是没有规律，而是规律非常敏感、很难长期预测。",
        ORANGE,
    )
    add_chip(
        slide,
        "插入演示视频 1：释放到开始复杂运动，约25秒",
        1.0,
        5.35,
        5.0,
        ORANGE,
        LIGHT_ORANGE,
    )
    add_footer(slide, 2)

    # 3
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(
        slide,
        "实验装置：三个磁铁 + 一个磁性摆锤",
        "装置本身很简单，但叠加出来的运动并不简单。",
    )
    base = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(0.95), Inches(1.55), Inches(4.9), Inches(4.9)
    )
    base.fill.solid()
    base.fill.fore_color.rgb = RGBColor(241, 245, 249)
    base.line.color.rgb = RGBColor(148, 163, 184)
    for x, y, label, col, fill in [
        (3.2, 2.3, "N", BLUE, LIGHT_BLUE),
        (2.15, 4.25, "N", BLUE, LIGHT_BLUE),
        (4.25, 4.25, "S", ORANGE, LIGHT_ORANGE),
    ]:
        m = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x - 0.23), Inches(y - 0.23), Inches(0.46), Inches(0.46)
        )
        m.fill.solid()
        m.fill.fore_color.rgb = fill
        m.line.color.rgb = col
        add_text(
            slide,
            label,
            x - 0.11,
            y - 0.13,
            0.22,
            0.22,
            12,
            col,
            True,
            PP_ALIGN.CENTER,
            MSO_ANCHOR.MIDDLE,
        )
    add_connector(slide, 3.2, 1.0, 3.55, 3.35, RGBColor(71, 85, 105), 1.6)
    bob = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(3.4), Inches(3.2), Inches(0.35), Inches(0.35)
    )
    bob.fill.solid()
    bob.fill.fore_color.rgb = RGBColor(250, 204, 21)
    bob.line.color.rgb = RGBColor(202, 138, 4)
    for x1, y1, x2, y2, col in [
        (2.4, 2.8, 4.4, 3.2, TEAL),
        (4.4, 3.2, 2.2, 4.4, ORANGE),
        (2.2, 4.4, 4.7, 4.65, PURPLE),
        (4.7, 4.65, 3.2, 2.3, BLUE),
    ]:
        add_connector(slide, x1, y1, x2, y2, col, 2)
    add_card(slide, 6.55, 1.65, 5.8, 0.9, "摆锤", "小球或摆锤带有磁性，会受到底板磁铁的吸引或排斥。", BLUE)
    add_card(slide, 6.55, 2.75, 5.8, 0.9, "三个磁铁", "磁铁位置不在同一直线上，摆锤受到的磁力方向不断变化。", TEAL)
    add_card(slide, 6.55, 3.85, 5.8, 0.9, "释放方式", "从静止、非平衡位置释放，多试几次，比较轨迹和最终停留位置。", ORANGE)
    add_chip(slide, "拍摄建议：先给装置一个5秒特写，再放手", 6.75, 5.25, 4.3, TEAL, LIGHT_TEAL)
    add_footer(slide, 3)

    # 4
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "受力原理：多种力叠加", "磁混沌摆不是“乱动”，而是在几个确定因素共同作用下运动。")
    bob = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(6.25), Inches(3.15), Inches(0.55), Inches(0.55)
    )
    bob.fill.solid()
    bob.fill.fore_color.rgb = RGBColor(250, 204, 21)
    bob.line.color.rgb = RGBColor(202, 138, 4)
    add_text(slide, "摆锤", 6.18, 3.78, 0.7, 0.25, 10, MUTED, True, PP_ALIGN.CENTER)
    for x, y, t, b, c, tx, ty in [
        (0.75, 1.75, "重力", "让摆锤倾向于回到最低位置，像普通单摆一样往复。", BLUE, 6.25, 3.35),
        (0.75, 4.2, "阻力/摩擦", "空气阻力和转轴摩擦会消耗能量，所以摆动越来越慢。", TEAL, 6.25, 3.45),
        (8.15, 1.75, "磁力", "三个磁铁从不同方向影响摆锤，距离越近作用越明显。", ORANGE, 6.8, 3.35),
        (8.15, 4.2, "非线性", "磁力随位置变化很快，微小位置差别会逐渐被放大。", PURPLE, 6.8, 3.45),
    ]:
        add_card(slide, x, y, 4.2, 1.1, t, b, c)
        sx = x + 4.2 if x < 6 else x
        add_connector(slide, sx, y + 0.55, tx, ty, c, 1.8)
    add_text(
        slide,
        "普通单摆：主要受重力，轨迹较规则\n磁混沌摆：重力 + 多个磁力 + 阻力，轨迹更复杂",
        3.45,
        5.85,
        6.45,
        0.65,
        17,
        DARK,
        True,
        PP_ALIGN.CENTER,
    )
    add_footer(slide, 4)

    # 5
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "为什么叫“混沌”？", "混沌不是完全随机，而是“确定系统中的难预测运动”。")
    add_card(slide, 0.8, 1.55, 3.85, 1.25, "不是纯随机", "摆锤每一刻都受物理规律控制，力的来源是确定的。", BLUE)
    add_card(slide, 0.8, 3.05, 3.85, 1.25, "但难以长期预测", "只要释放位置、角度或速度有一点点差别，后面的轨迹就可能明显不同。", ORANGE)
    add_card(slide, 0.8, 4.55, 3.85, 1.25, "最终状态不唯一", "同样是释放摆锤，最后可能停在三个磁铁中的不同一个上方。", TEAL)
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.45), Inches(1.55), Inches(6.7), Inches(4.55)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = WHITE
    panel.line.color.rgb = RGBColor(226, 232, 240)
    for x, y, c in [(8.0, 2.3, BLUE), (6.9, 4.75, TEAL), (9.35, 4.65, ORANGE)]:
        s = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x - 0.22), Inches(y - 0.22), Inches(0.44), Inches(0.44)
        )
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor(241, 245, 249)
        s.line.color.rgb = c
        s.line.width = Pt(1.5)
    for x, y, c, label in [(7.65, 3.25, BLUE, "A"), (7.82, 3.28, ORANGE, "B")]:
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.13), Inches(0.13)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = c
        dot.line.color.rgb = c
        add_text(slide, label, x - 0.08, y - 0.33, 0.3, 0.18, 10, c, True, PP_ALIGN.CENTER)
    for pts, col in [
        ([(7.72, 3.31), (8.35, 2.9), (7.95, 2.25), (8.85, 2.6), (9.4, 4.45)], BLUE),
        ([(7.89, 3.34), (7.15, 3.75), (6.82, 4.6), (7.6, 4.25), (8.1, 2.45)], ORANGE),
    ]:
        for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
            add_connector(slide, x1, y1, x2, y2, col, 2)
    add_text(slide, "两个几乎相同的起点\n可能走向不同结局", 9.55, 1.85, 2.05, 0.75, 17, DARK, True, PP_ALIGN.CENTER)
    add_chip(slide, "插入演示视频 2：连续两次释放，约35秒", 6.7, 5.45, 4.4, ORANGE, LIGHT_ORANGE)
    add_footer(slide, 5)

    # 6
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "我们的观察记录", "实验说明中提到：摆锤会作复杂运动，约2分钟后停在某个磁极上方。")
    add_image_fit(slide, RESULT, 0.75, 1.45, 5.15, 4.85)
    add_image_fit(slide, SHEET, 6.25, 1.45, 5.95, 2.25)
    add_card(slide, 6.25, 4.0, 5.95, 0.95, "记录角度 1", "从不同位置释放，比较最终停在哪个磁铁上方。", BLUE)
    add_card(slide, 6.25, 5.1, 5.95, 0.95, "记录角度 2", "同一位置附近重复释放，观察轨迹是否仍然完全一样。", TEAL)
    add_chip(slide, "插入演示视频 3：最后变慢并停下，约30秒", 1.0, 6.25, 4.8, TEAL, LIGHT_TEAL)
    add_footer(slide, 6)

    # 7
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(241, 245, 249))
    add_title(slide, "结论：简单装置，也能产生复杂运动", "磁混沌摆展示了“确定性规律”和“长期不可预测性”可以同时存在。")
    add_card(slide, 0.85, 1.65, 3.75, 1.25, "一句话总结", "三个磁铁改变了摆锤的受力方向，使轨迹对初始条件非常敏感。", BLUE)
    add_card(slide, 4.82, 1.65, 3.75, 1.25, "实验价值", "它把抽象的“混沌”变成可以直接看见的运动轨迹。", ORANGE)
    add_card(slide, 8.8, 1.65, 3.75, 1.25, "现实联系", "天气、湍流、生态系统、交通流等复杂系统中，也常见类似的敏感性。", TEAL)
    add_text(
        slide,
        "所以，这个实验最想告诉我们的不是“小球乱动”，而是：\n复杂现象背后仍有规律，只是预测它需要非常精确的初始信息。",
        1.25,
        3.55,
        10.8,
        1.0,
        26,
        DARK,
        True,
        PP_ALIGN.CENTER,
    )
    add_chip(slide, "结尾可用演示视频定格：小球停在某一磁铁上方", 4.7, 5.2, 4.2, PURPLE, RGBColor(237, 233, 254))
    add_text(slide, "谢谢观看", 5.3, 6.1, 2.7, 0.4, 24, BLUE, True, PP_ALIGN.CENTER)
    add_footer(slide, 7)

    prs.save(PPTX)
    print(PPTX)


if __name__ == "__main__":
    main()
