from pathlib import Path
import tempfile
import unittest

from akane_document_writer import DocumentError, render_document, render_sources


def material(blocks, fmt="docx", complete=True):
    return {
        "schema": "akane.document-material.v1",
        "source_format": fmt,
        "complete": complete,
        "limitations": [] if complete else ["image_not_extracted"],
        "blocks": blocks,
    }


class SourceTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def test_literal_document_content_and_interleaved_table_are_not_reparsed_as_markdown(self):
        from docx import Document

        source = material(
            [
                {"kind": "paragraph", "text": "# **原样符号**\n保留后续"},
                {"kind": "table", "rows": [["**不丢星号**", "数值"], [False, 0]]},
                {"kind": "paragraph", "text": "后文尾部"},
            ]
        )
        output = self.root / "source.docx"
        result = render_sources(materials=[source], output_path=output, output_format="docx")
        document = Document(output)
        self.assertEqual(document.paragraphs[0].text, "# **原样符号**\n保留后续")
        self.assertEqual(document.tables[0].cell(0, 0).text, "**不丢星号**")
        self.assertEqual(document.tables[0].cell(1, 0).text, "False")
        self.assertEqual(document.tables[0].cell(1, 1).text, "0")
        self.assertEqual(document.paragraphs[-1].text, "后文尾部")
        tags = [node.tag.rsplit("}", 1)[-1] for node in document.element.body]
        self.assertEqual(tags[:2], ["p", "tbl"])
        self.assertIn("source_text_export_not_original_layout", result["notices"])

    def test_explicit_rows_also_preserve_literal_emphasis_characters(self):
        from docx import Document

        output = self.root / "rows.docx"
        render_document(output_path=output, output_format="docx", rows=[["value"], ["**literal**"]])
        self.assertEqual(Document(output).tables[0].cell(1, 0).text, "**literal**")

    def test_full_text_preserves_tail_and_original_markdown_for_text_output(self):
        output = self.root / "source.md"
        content = "# head\n" + "line\n" * 12000 + "exact tail\n"
        render_sources(
            materials=[material([{"kind": "text", "text": content}], fmt="md")], output_path=output, output_format="md"
        )
        self.assertEqual(output.read_bytes(), content.encode("utf-8"))

    def test_incomplete_source_or_ambiguous_sheet_export_never_creates_file(self):
        output = self.root / "source.xlsx"
        with self.assertRaisesRegex(DocumentError, "document_source_incomplete"):
            render_sources(
                materials=[material([{"kind": "text", "text": "short preview"}], complete=False)],
                output_path=output,
                output_format="xlsx",
            )
        sheets = [{"kind": "sheet", "name": name, "rows": [[0]]} for name in ("one", "two")]
        with self.assertRaisesRegex(DocumentError, "document_source_table_selection_required"):
            render_sources(materials=[material(sheets, fmt="xlsx")], output_path=output, output_format="xlsx")
        self.assertFalse(output.exists())

    def test_formula_and_date_export_are_explicitly_literal_and_typed_values_remain(self):
        from openpyxl import load_workbook

        output = self.root / "source.xlsx"
        source = material(
            [
                {
                    "kind": "sheet",
                    "name": "Values",
                    "rows": [
                        ["number", "bool", "formula", "date"],
                        [
                            0,
                            False,
                            {"type": "formula", "expression": "=A2+1"},
                            {"type": "datetime", "value": "2026-09-06"},
                        ],
                    ],
                }
            ],
            fmt="xlsx",
        )
        result = render_sources(materials=[source], output_path=output, output_format="xlsx")
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.active["A2"].value, 0)
        self.assertIs(book.active["B2"].value, False)
        self.assertEqual(book.active["C2"].value, "=A2+1")
        self.assertEqual(book.active["C2"].data_type, "s")
        self.assertEqual(book.active["D2"].value, "2026-09-06")
        self.assertIn("source_formulas_exported_as_text_not_evaluated", result["notices"])
        self.assertIn("source_dates_exported_as_iso_text", result["notices"])

    def test_source_export_does_not_invent_subtitle_times(self):
        output = self.root / "source.srt"
        with self.assertRaisesRegex(DocumentError, "document_source_timing_required"):
            render_sources(
                materials=[material([{"kind": "text", "text": "untimed speech"}], fmt="txt")],
                output_path=output,
                output_format="srt",
            )
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
