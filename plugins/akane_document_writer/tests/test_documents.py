from __future__ import annotations

import csv
import hashlib
from pathlib import Path
import tempfile
import unittest
import zipfile

from akane_document_writer import DocumentError, render_document, style_document
from akane_document_writer.markdown import blocks
from akane_document_writer.validation import column, formatting, row_range, table_rows


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def render(self, fmt="xlsx", **kwargs):
        path = self.root / f"output.{fmt}"
        result = render_document(output_path=path, output_format=fmt, **kwargs)
        self.assertEqual(result["status"], "ok")
        self.assertNotIn(str(self.root), str(result))
        return path

    def test_typed_values_and_formula_strings_survive_xlsx(self):
        from openpyxl import load_workbook

        path = self.render(rows=[["项目", "值", "备注"], [0, False, "  空格  "], [None, -2.5, "=1+1"], ["", "", ""]])
        book = load_workbook(path)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet["A2"].value, 0)
        self.assertEqual(sheet["A2"].data_type, "n")
        self.assertIs(sheet["B2"].value, False)
        self.assertEqual(sheet["B2"].data_type, "b")
        self.assertEqual(sheet["C2"].value, "  空格  ")
        self.assertEqual(sheet["B3"].value, -2.5)
        self.assertEqual(sheet["C3"].value, "=1+1")
        self.assertEqual(sheet["C3"].data_type, "s")
        self.assertEqual(sheet.max_row, 4)
        self.assertTrue(sheet["C2"].alignment.wrap_text)

    def test_cjk_width_is_applied_without_explicit_formatting(self):
        from openpyxl import load_workbook

        path = self.render(rows=[["需要完整显示的中文标题", "数量"], ["甲", 0]])
        book = load_workbook(path)
        self.addCleanup(book.close)
        self.assertGreaterEqual(book.active.column_dimensions["A"].width, 24)

    def test_csv_preserves_empty_rows_zero_boolean_and_whitespace(self):
        path = self.render("csv", rows=[["A", "B"], [0, False], ["", ""], [" space ", "=1+1"]])
        with path.open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(list(csv.reader(stream)), [["A", "B"], ["0", "False"], ["", ""], [" space ", "=1+1"]])

    def test_over_limit_cells_rows_and_bad_types_fail_without_truncation(self):
        for rows in ([["x" * 32768]], [[1]] * 1001, [[1] * 51], [[{}]], [[float("nan")]], [[]]):
            with self.subTest(rows_type=type(rows)), self.assertRaises(DocumentError):
                table_rows(rows)
        self.assertEqual(table_rows([[0, False, None, " "]]), [[0, False, None, " "]])

    def test_markdown_heading_does_not_consume_following_body(self):
        parsed = blocks("# 标题\n正文 **加粗**\n- 第一项\n- 第二项\n\n| 姓名 | 值 |\n| --- | --- |\n| A\\|B | 0 |")
        self.assertEqual([block.kind for block in parsed], ["heading", "paragraph", "bullet", "bullet", "table"])
        self.assertEqual(parsed[-1].rows[-1], ["A|B", "0"])

    def test_docx_real_structure_typed_table_and_indexed_paragraph(self):
        from docx import Document
        from docx.oxml.ns import qn

        path = self.render(
            "docx",
            title="文档验收",
            content="# 第一节\n完整正文 **强调**\n- 内容一\n- 内容二",
            rows=[["数量", "状态"], [0, False]],
            styles={"paragraphs": [{"paragraph_index": 3, "font_color": "red"}]},
        )
        document = Document(path)
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        for name in ("Normal", "Title", "Heading 1", "Heading 2"):
            self.assertIsNone(document.styles[name]._element.find(".//" + qn("w:pBdr")))
            self.assertFalse(document.styles[name].font.underline)
        self.assertNotIn(qn("w:eastAsiaTheme"), document.styles["Title"]._element.rPr.rFonts.attrib)
        self.assertEqual(document.paragraphs[1].text, "第一节")
        self.assertEqual(document.paragraphs[2].text, "完整正文 强调")
        self.assertTrue(document.paragraphs[2].runs[-1].bold)
        self.assertEqual(str(document.paragraphs[2].runs[0].font.color.rgb), "FF0000")
        self.assertEqual(document.tables[0].cell(1, 0).text, "0")
        self.assertEqual(document.tables[0].cell(1, 1).text, "False")
        self.assertIsNotNone(document.tables[0]._tbl.tblPr.find(qn("w:tblBorders")))
        self.assertIsNotNone(document.tables[0].rows[0]._tr.trPr.find(qn("w:tblHeader")))

    def test_docx_style_preserves_body_and_source_bytes_exactly(self):
        from docx import Document
        from docx.oxml.ns import qn

        source = self.render("docx", content="第一段\n\n第二段", rows=[["项目"], ["值"]])
        original = source.read_bytes()
        output = self.root / "styled.docx"
        style_document(
            source_path=source,
            output_path=output,
            styles={
                "paragraphs": [{"index": 2, "bold": True, "highlight_color": "123456"}],
                "header": {"fill_color": "ABCDEF"},
            },
        )
        self.assertEqual(source.read_bytes(), original)
        doc = Document(output)
        self.assertEqual([p.text for p in doc.paragraphs], [p.text for p in Document(source).paragraphs])
        self.assertTrue(doc.paragraphs[1].runs[0].bold)
        self.assertFalse(bool(doc.paragraphs[0].runs[0].bold))
        self.assertEqual(doc.paragraphs[1].runs[0]._r.rPr.find(qn("w:shd")).get(qn("w:fill")), "123456")
        self.assertEqual(len(doc.tables[0].cell(0, 0)._tc.tcPr.findall(qn("w:shd"))), 1)
        self.assertEqual(doc.tables[0].cell(0, 0)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")), "ABCDEF")

    def test_negative_out_of_range_and_ambiguous_selectors_rejected(self):
        for rule in (
            {"row": -2},
            {"row": "row2"},
            {"row": 0},
            {"row_start": 5},
            {"row_start": 2, "row_end": 9},
            {"row": 2, "start": 1},
        ):
            with self.subTest(rule=rule), self.assertRaises(DocumentError):
                row_range(rule, 3)
        self.assertEqual(list(row_range({"row_start": 2, "row_end": 3}, 3)), [2, 3])
        for rule in ({"match_header": "姓名"}, {"column_index": -1}, {"column": "Z"}, {"index": 1, "column": "A"}):
            with self.subTest(rule=rule), self.assertRaises(DocumentError):
                column(rule, ["", "姓名", "姓名"])
        self.assertEqual(column({"match_header": "姓名"}, ["", "姓名"]), 2)
        with self.assertRaises(DocumentError):
            column({"match_header": "姓名"}, ["", "姓名后缀"])

    def test_style_rule_validation_never_coerces_false_string(self):
        for rules in (
            {"header": {"bold": "false"}},
            {"auto_width": "false"},
            {"header": {"font_color": "not-a-color"}},
            {"rows": [{"row": 2, "unknown": 1}]},
            {"unknown": 1},
            {"rows": [{"row": 2, "bold": True}] * 121},
        ):
            with self.subTest(rules=rules), self.assertRaises(DocumentError):
                formatting(rules, "xlsx")
        with self.assertRaisesRegex(DocumentError, "unsupported"):
            formatting({"header": {"bold": True}}, "pdf")

    def test_xlsx_styles_apply_to_exact_columns_rows_and_conditions(self):
        from openpyxl import load_workbook

        path = self.render(
            rows=[["项目", "分数", "是否"], ["甲", 0, False], ["乙", 80, True]],
            styles={
                "columns": [{"column_index": 2, "italic": True}],
                "rows": [{"row_start": 3, "row_end": 3, "font_color": "blue"}],
                "row_rules": [{"where": {"match_header": "分数", "lt": 60}, "fill_color": "pink"}],
                "cells": [{"row": 2, "column": "C", "bold": True}],
            },
        )
        book = load_workbook(path)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertTrue(sheet["B2"].font.italic)
        self.assertTrue(sheet["C2"].font.bold)
        self.assertEqual(sheet["A3"].font.color.rgb[-6:], "0000FF")
        self.assertEqual(sheet["C2"].fill.fgColor.rgb[-6:], "F4CCCC")
        self.assertNotEqual(sheet["C3"].fill.fgColor.rgb[-6:], "F4CCCC")

    def test_style_multisheet_requires_selector_and_keeps_formulas(self):
        from openpyxl import Workbook, load_workbook

        book = Workbook()
        first = book.active
        first.title = "原表"
        first.append(["数值", "公式"])
        first.append([4, "=A2*2"])
        first.column_dimensions["A"].width = 22
        second = book.create_sheet("备注")
        second["A1"] = "保留"
        source = self.root / "source.xlsx"
        book.save(source)
        book.close()
        original = hashlib.sha256(source.read_bytes()).digest()
        output = self.root / "styled.xlsx"
        with self.assertRaisesRegex(DocumentError, "sheet_selector_required"):
            style_document(source_path=source, output_path=output, styles={"header": {"bold": True}})
        self.assertFalse(output.exists())
        style_document(
            source_path=source, output_path=output, styles={"sheet_name": "原表", "header": {"font_color": "red"}}
        )
        self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), original)
        actual = load_workbook(output)
        self.addCleanup(actual.close)
        self.assertEqual(actual["原表"]["B2"].value, "=A2*2")
        self.assertEqual(actual["备注"]["A1"].value, "保留")
        self.assertEqual(actual["原表"].column_dimensions["A"].width, 22)

    def test_no_match_is_failure_and_leaves_no_output_or_temporary(self):
        source = self.render("docx", content="文本")
        output = self.root / "styled.docx"
        with self.assertRaisesRegex(DocumentError, "no_match"):
            style_document(
                source_path=source, output_path=output, styles={"paragraphs": [{"contains": "不存在", "bold": True}]}
            )
        self.assertFalse(output.exists())
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), ["output.docx"])

    def test_atomic_publication_does_not_overwrite(self):
        output = self.render("txt", content="原文")
        with self.assertRaisesRegex(DocumentError, "output_exists"):
            render_document(output_path=output, output_format="txt", content="替换")
        self.assertEqual(output.read_text(encoding="utf-8"), "原文")

    def test_text_and_subtitle_formats_preserve_raw_content(self):
        for fmt in ("txt", "md", "html", "srt", "lrc", "vtt", "json"):
            with self.subTest(fmt=fmt):
                content = '{"zero":0,"false":false}' if fmt == "json" else "\n保留\r\n换行\n"
                path = self.render(fmt, content=content)
                self.assertEqual(path.read_bytes(), content.encode("utf-8"))

    def test_json_invalid_is_not_wrapped_as_fake_success(self):
        for content in ("invalid", "NaN", '{"x":Infinity}'):
            with self.subTest(content=content), self.assertRaisesRegex(DocumentError, "json_invalid"):
                self.render("json", content=content)

    def test_real_pdf_contains_cjk_table_and_zero(self):
        from pypdf import PdfReader

        path = self.render(
            "pdf",
            title="文档输出验收",
            content="# 中文内容\n正文内容应完整显示。",
            rows=[["项目", "数量"], ["样本", 0]],
        )
        reader = PdfReader(path)
        self.assertGreaterEqual(len(reader.pages), 1)
        extracted = "\n".join(page.extract_text() for page in reader.pages)
        for expected in ("文档输出验收", "正文内容应完整显示", "样本", "0"):
            self.assertIn(expected, extracted)
        fonts = reader.pages[0]["/Resources"]["/Font"]
        self.assertTrue(any("/FontFile2" in font.get_object().get("/FontDescriptor", {}) for font in fonts.values()))

    def test_pdf_missing_font_does_not_fall_back_to_broken_glyphs(self):
        from unittest.mock import patch

        with patch.dict("os.environ", {"AKANE_DOCUMENT_FONT": str(self.root / "missing.ttf")}):
            with self.assertRaisesRegex(DocumentError, "font_unavailable_or_missing_glyphs"):
                self.render("pdf", content="中文")
        self.assertFalse((self.root / "output.pdf").exists())

    def test_document_noop_style_is_rejected(self):
        source = self.render(rows=[["A"], [0]])
        with self.assertRaisesRegex(DocumentError, "style_required"):
            style_document(
                source_path=source,
                output_path=self.root / "noop.xlsx",
                styles={"sheet_name": "生成内容", "auto_width": False},
            )

    def test_hyperlink_run_gets_paragraph_style_without_losing_link(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE

        document = Document()
        paragraph = document.add_paragraph()
        link = OxmlElement("w:hyperlink")
        rel = paragraph.part.relate_to("https://example.com", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link.set(qn("r:id"), rel)
        run = OxmlElement("w:r")
        value = OxmlElement("w:t")
        value.text = "链接文本"
        run.append(value)
        link.append(run)
        paragraph._p.append(link)
        source = self.root / "link.docx"
        document.save(source)
        output = self.root / "styled.docx"
        style_document(source_path=source, output_path=output, styles={"paragraphs": [{"index": 1, "bold": True}]})
        result = Document(output)
        self.assertEqual(result.paragraphs[0].text, "链接文本")
        self.assertIsNotNone(result.paragraphs[0]._p.find(".//" + qn("w:b")))
        self.assertEqual(result.part.rels[rel].target_ref, "https://example.com")

    def test_pdf_unsupported_glyphs_and_styles_fail_structurally(self):
        with self.assertRaisesRegex(DocumentError, "glyph_unsupported"):
            self.render("pdf", content="测试😀")
        self.assertFalse((self.root / "output.pdf").exists())
        with self.assertRaisesRegex(DocumentError, "formatting_unsupported"):
            self.render("pdf", content="正文", styles={"header": {"bold": True}})

    def test_duplicate_archive_rejected_before_edit(self):
        import warnings

        source = self.root / "duplicate.docx"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(source, "w") as archive:
                archive.writestr("word/document.xml", "first")
                archive.writestr("word/document.xml", "second")
        with self.assertRaisesRegex(DocumentError, "duplicate_member"):
            style_document(source_path=source, output_path=self.root / "styled.docx", styles={"header": {"bold": True}})

    def test_ambiguous_content_is_not_silently_dropped(self):
        with self.assertRaisesRegex(DocumentError, "table_content_ambiguous"):
            self.render(rows=[["A"]], content="other content")
        with self.assertRaisesRegex(DocumentError, "rows_unsupported"):
            self.render("txt", rows=[["A"]])


if __name__ == "__main__":
    unittest.main()
