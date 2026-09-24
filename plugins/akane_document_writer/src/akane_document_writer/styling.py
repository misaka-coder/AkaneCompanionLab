"""Shared selectors with format-specific mutation; no host dependencies."""

from __future__ import annotations

from copy import copy
from typing import Any
import unicodedata

from .validation import DocumentError, column, one, positive_int, row_matches, row_range, text


def table_styles(values: list[list[Any]], rules: dict) -> list[tuple[int, int, dict]]:
    if not values or not values[0]:
        raise DocumentError("document_table_empty")
    count, width = len(values), len(values[0])
    result = []
    if "header" in rules:
        result.extend((1, col, rules["header"]) for col in range(1, width + 1))
    for rule in rules.get("columns", []):
        col = column(rule, values[0])
        result.extend((row, col, rule) for row in range(1, count + 1))
    for rule in rules.get("rows", []):
        result.extend((row, col, rule) for row in row_range(rule, count) for col in range(1, width + 1))
    for rule in rules.get("cells", []):
        row = positive_int(one(rule, ("row", "row_index")))
        if row > count:
            raise DocumentError("document_row_out_of_range")
        col = column({key: value for key, value in rule.items() if key not in {"row", "row_index"}}, values[0])
        result.append((row, col, rule))
    for rule in rules.get("row_rules", []):
        matches = [row for row in range(1, count + 1) if row_matches(rule, values[row - 1], values[0])]
        if not matches:
            raise DocumentError("document_style_no_match")
        result.extend((row, col, rule) for row in matches for col in range(1, width + 1))
    for rule in rules.get("highlights", []):
        needle = one(rule, ("text", "contains"))
        if not isinstance(needle, str) or not needle:
            raise DocumentError("document_selector_text_invalid")
        matches = [
            (row, col, rule)
            for row in range(1, count + 1)
            for col in range(1, width + 1)
            if needle in text(values[row - 1][col - 1])
        ]
        if not matches:
            raise DocumentError("document_style_no_match")
        result.extend(matches)
    return result


def xlsx_style(cell: Any, rule: dict) -> None:
    from openpyxl.styles import PatternFill

    font = copy(cell.font)
    for key in ("bold", "italic"):
        if key in rule:
            setattr(font, key, rule[key])
    if "font_color" in rule:
        font.color = rule["font_color"]
    cell.font = font
    fill = rule.get("fill_color", rule.get("highlight_color"))
    if fill:
        cell.fill = PatternFill(fill_type="solid", fgColor=fill)


def display_width(value: Any) -> int:
    return max(
        (
            sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in line)
            for line in text(value).splitlines()
        ),
        default=0,
    )


def xlsx_format(sheet: Any, rules: dict, *, new: bool = False) -> None:
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    if sheet.max_row > 10000 or sheet.max_column > 200:
        raise DocumentError("document_style_sheet_too_large")
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    operations = table_styles(values, rules)
    for row, col, rule in operations:
        xlsx_style(sheet.cell(row, col), rule)
    if rules.get("auto_width", new):
        for col in range(1, sheet.max_column + 1):
            width = max(10, min(60, max(display_width(row[col - 1]) for row in values) + 2))
            sheet.column_dimensions[get_column_letter(col)].width = width
            for row in range(1, sheet.max_row + 1):
                cell = sheet.cell(row, col)
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                alignment.vertical = "center"
                cell.alignment = alignment
    elif new:
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)


def shade(properties: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for element in list(properties.findall(qn("w:shd"))):
        properties.remove(element)
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def docx_paragraph_style(paragraph: Any, rule: dict) -> None:
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    from docx.text.run import Run

    # Include hyperlink runs; paragraph.runs alone silently misses them.
    for element in paragraph._p.iter(qn("w:r")):
        run = Run(element, paragraph)
        for key in ("bold", "italic"):
            if key in rule:
                setattr(run, key, rule[key])
        if "font_color" in rule:
            run.font.color.rgb = RGBColor.from_string(rule["font_color"])
        if "highlight_color" in rule or "fill_color" in rule:
            shade(run._r.get_or_add_rPr(), rule.get("highlight_color", rule.get("fill_color")))


def docx_format(document: Any, rules: dict) -> None:
    # Indices deliberately address top-level body paragraphs only, including title.
    paragraphs = document.paragraphs
    for rule in rules.get("paragraphs", []):
        index = one(rule, ("index", "paragraph_index"), required=False)
        needle = one(rule, ("text", "contains"), required=False)
        if index is not None and needle is not None:
            raise DocumentError("document_selector_ambiguous")
        if index is not None:
            index = positive_int(index)
            if index > len(paragraphs):
                raise DocumentError("document_paragraph_out_of_range")
            matches = [paragraphs[index - 1]]
        elif needle is not None:
            if not isinstance(needle, str) or not needle:
                raise DocumentError("document_selector_text_invalid")
            matches = [paragraph for paragraph in paragraphs if needle in paragraph.text]
        else:
            raise DocumentError("document_selector_missing")
        if not matches:
            raise DocumentError("document_style_no_match")
        for paragraph in matches:
            docx_paragraph_style(paragraph, rule)

    tables = document.tables
    table_keys = {"header", "columns", "rows", "cells", "row_rules"}
    if table_keys & rules.keys():
        if not tables:
            raise DocumentError("document_table_missing")
        if "table_index" in rules:
            index = rules["table_index"]
            if index > len(tables):
                raise DocumentError("document_table_out_of_range")
            tables = [tables[index - 1]]
        elif len(tables) > 1:
            raise DocumentError("document_table_selector_required")
        for table in tables:
            values = [[cell.text for cell in row.cells] for row in table.rows]
            for row, col, rule in table_styles(values, {key: val for key, val in rules.items() if key != "highlights"}):
                cell = table.cell(row - 1, col - 1)
                if "fill_color" in rule:
                    shade(cell._tc.get_or_add_tcPr(), rule["fill_color"])
                for paragraph in cell.paragraphs:
                    docx_paragraph_style(paragraph, {key: val for key, val in rule.items() if key != "fill_color"})

    # A text highlight styles the entire matching paragraph/cell, explicitly
    # documented in the descriptor; never pretends to perform substring edits.
    all_paragraphs = list(paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)
    for rule in rules.get("highlights", []):
        needle = one(rule, ("text", "contains"))
        if not isinstance(needle, str) or not needle:
            raise DocumentError("document_selector_text_invalid")
        matches = [paragraph for paragraph in all_paragraphs if needle in paragraph.text]
        if not matches:
            raise DocumentError("document_style_no_match")
        for paragraph in matches:
            docx_paragraph_style(paragraph, rule)
