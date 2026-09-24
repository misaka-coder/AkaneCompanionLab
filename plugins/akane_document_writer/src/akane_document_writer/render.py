"""Bounded real document I/O. Each operation creates a new file atomically."""

from __future__ import annotations

from copy import copy
import csv
import json
from pathlib import Path
import re
from typing import Any
from uuid import uuid4
from xml.sax.saxutils import escape
import zipfile
import warnings

from .markdown import Block, blocks, table_from_content
from .fonts import pdf_font
from .styling import display_width, docx_format, shade, xlsx_format
from .validation import DocumentError, FORMATS, formatting, table_rows, text


def _check_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 10000 or sum(member.file_size for member in members) > 100_000_000:
                raise DocumentError("document_archive_too_large")
            if len({member.filename for member in members}) != len(members):
                raise DocumentError("document_archive_duplicate_member")
            if any(member.flag_bits & 1 for member in members):
                raise DocumentError("document_encrypted_unsupported")
            if any(
                any(
                    part in member.filename.lower()
                    for part in ("vbaproject", "_xmlsignatures", "/embeddings/", "/activex/")
                )
                for member in members
            ):
                raise DocumentError("document_complex_package_unsupported")
    except zipfile.BadZipFile:
        raise DocumentError("document_archive_invalid") from None


def _atomic(output: Path, writer: Any) -> None:
    if output.exists():
        raise DocumentError("document_output_exists")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.stem}.{uuid4().hex}{output.suffix}")
    try:
        writer(temporary)
        if not temporary.is_file() or temporary.stat().st_size == 0:
            raise DocumentError("document_output_empty")
        # Producer-owned path; hard linking makes publication fail if a second
        # invocation unexpectedly claimed the target, rather than overwrite it.
        output.hardlink_to(temporary)
    finally:
        temporary.unlink(missing_ok=True)


def _rectangular(rows: list[list[Any]]) -> list[list[Any]]:
    width = max((len(row) for row in rows), default=0)
    return [row + [None] * (width - len(row)) for row in rows]


def render_document(
    *,
    output_path: Path,
    output_format: str,
    title: str = "",
    content: str = "",
    rows: Any = None,
    styles: Any = None,
    source_blocks: list[Block] | None = None,
) -> dict:
    if output_format not in FORMATS or output_path.suffix.lower() != f".{output_format}":
        raise DocumentError("document_format_invalid")
    if not isinstance(content, str) or len(content) > 80000 or not isinstance(title, str) or len(title) > 80:
        raise DocumentError("document_content_invalid")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", content + title):
        raise DocumentError("document_control_character")
    validated_rows = table_rows(rows)
    if source_blocks is not None:
        if output_format not in {"docx", "pdf"} or content or rows or not isinstance(source_blocks, list):
            raise DocumentError("document_source_blocks_invalid")
        total = 0
        for block in source_blocks:
            if not isinstance(block, Block) or block.kind not in {
                "paragraph",
                "heading",
                "bullet",
                "number",
                "code",
                "table",
            }:
                raise DocumentError("document_source_blocks_invalid")
            if not isinstance(block.text, str) or type(block.literal) is not bool:
                raise DocumentError("document_source_blocks_invalid")
            if block.kind == "heading" and not 1 <= block.level <= 6:
                raise DocumentError("document_source_blocks_invalid")
            if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", block.text):
                raise DocumentError("document_control_character")
            validated = table_rows(block.rows)
            if block.kind == "table" and not validated:
                raise DocumentError("document_table_empty")
            total += len(block.text) + sum(len(text(cell)) for row in validated for cell in row)
        if len(source_blocks) > 10000 or total > 1_000_000:
            raise DocumentError("document_content_invalid")
    if not content and not validated_rows and not source_blocks:
        raise DocumentError("document_content_required")
    if validated_rows and output_format not in {"csv", "xlsx", "docx", "pdf"}:
        raise DocumentError("document_rows_unsupported")
    if validated_rows and content and output_format in {"csv", "xlsx"}:
        raise DocumentError("document_table_content_ambiguous")
    rules = formatting(styles, output_format)
    if output_format == "json":
        try:
            json.loads(content, parse_constant=lambda _: (_ for _ in ()).throw(ValueError()))
        except (ValueError, RecursionError):
            raise DocumentError("document_json_invalid") from None

    def write(path: Path) -> None:
        if output_format in {"txt", "md", "html", "json", "srt", "lrc", "vtt"}:
            path.write_text(content, encoding="utf-8", newline="")
        elif output_format == "csv":
            data = validated_rows or table_rows(table_from_content(content))
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream).writerows(data)
        elif output_format == "xlsx":
            _xlsx(path, title, validated_rows or table_rows(table_from_content(content)), rules)
        elif output_format == "docx":
            _docx(path, title, content, validated_rows, rules, source_blocks=source_blocks)
        elif output_format == "pdf":
            _pdf(path, title, content, validated_rows, source_blocks=source_blocks)

    _atomic(output_path, write)
    notices = []
    if output_format == "csv" and any(
        isinstance(cell, str) and cell.startswith(("=", "+", "-", "@"))
        for row in (validated_rows or table_from_content(content))
        for cell in row
    ):
        notices.append("csv_formula_like_text_preserved_use_text_import")
    return {"status": "ok", "output_format": output_format, "file_size": output_path.stat().st_size, "notices": notices}


def style_document(*, source_path: Path, output_path: Path, styles: Any) -> dict:
    fmt = source_path.suffix.lower().removeprefix(".")
    if fmt not in {"docx", "xlsx"} or output_path.suffix.lower() != source_path.suffix.lower():
        raise DocumentError("document_style_format_invalid")
    rules = formatting(styles, fmt)
    if not set(rules) - {"sheet_name", "table_index", "auto_width"} and not rules.get("auto_width"):
        raise DocumentError("document_style_required")
    if not source_path.is_file() or source_path.stat().st_size > 50_000_000:
        raise DocumentError("document_source_invalid")
    _check_zip(source_path)

    def write(path: Path) -> None:
        if fmt == "docx":
            from docx import Document

            document = Document(source_path)
            docx_format(document, rules)
            document.save(path)
        else:
            from openpyxl import load_workbook

            with warnings.catch_warnings(record=True) as import_warnings:
                warnings.simplefilter("always")
                workbook = load_workbook(source_path, data_only=False, keep_links=True, rich_text=True)
            try:
                if import_warnings:
                    raise DocumentError("document_workbook_feature_unsupported")
                if "sheet_name" in rules:
                    if rules["sheet_name"] not in workbook.sheetnames:
                        raise DocumentError("document_sheet_missing")
                    sheet = workbook[rules["sheet_name"]]
                elif len(workbook.worksheets) == 1:
                    sheet = workbook.worksheets[0]
                else:
                    raise DocumentError("document_sheet_selector_required")
                xlsx_format(sheet, rules)
                with warnings.catch_warnings(record=True) as export_warnings:
                    warnings.simplefilter("always")
                    workbook.save(path)
                if export_warnings:
                    raise DocumentError("document_workbook_feature_unsupported")
            finally:
                workbook.close()

    _atomic(output_path, write)
    return {"status": "ok", "output_format": fmt, "file_size": output_path.stat().st_size, "notices": []}


def _xlsx(path: Path, title: str, rows: list[list[Any]], rules: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = re.sub(r"[\[\]:*?/\\]", "_", title or "生成内容").strip("'")[:31] or "Sheet"
        if "sheet_name" in rules and rules["sheet_name"] != sheet.title:
            raise DocumentError("document_sheet_missing")
        for row in _rectangular(rows):
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    # Model/user cell strings are literal text, never formulas.
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="000000")
            cell.fill = PatternFill("solid", fgColor="E7EDF3")
        xlsx_format(sheet, rules, new=True)
        workbook.save(path)
    finally:
        workbook.close()


def _inline(paragraph: Any, value: str) -> None:
    for part in re.split(r"(\*\*[^*\n]+\*\*|`[^`\n]+`|\*[^*\n]+\*)", value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        else:
            paragraph.add_run(part)


def _word_table(document: Any, rows: list[list[Any]], *, literal: bool = False) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    data = _rectangular(rows)
    table = document.add_table(rows=len(data), cols=len(data[0]))
    table.autofit = False
    weights = [max(6, min(35, max(display_width(row[col]) for row in data))) for col in range(len(data[0]))]
    for col, weight in enumerate(weights):
        table.columns[col].width = Inches(6.5 * weight / sum(weights))
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, val in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn(f"w:{key}"), val)
        borders.append(border)
    properties.append(borders)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, row in enumerate(data):
        for col, value in enumerate(row):
            cell = table.cell(row_index, col)
            cell.width = table.columns[col].width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if type(value) in (int, float, bool) else WD_ALIGN_PARAGRAPH.LEFT
            )
            if literal:
                paragraph.add_run(text(value))
            else:
                _inline(paragraph, text(value))
            if row_index == 0:
                shade(cell._tc.get_or_add_tcPr(), "E7EDF3")
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def _docx(
    path: Path,
    title: str,
    content: str,
    rows: list[list[Any]],
    rules: dict,
    *,
    source_blocks: list[Block] | None = None,
) -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(0.8)
    for name in ("Normal", "Title", *[f"Heading {n}" for n in range(1, 7)]):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        properties = style._element.get_or_add_pPr()
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("asciiTheme", "eastAsiaTheme", "hAnsiTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{key}"), None)
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    if title:
        document.add_paragraph(title, style="Title")
    for block in source_blocks if source_blocks is not None else blocks(content):
        if block.kind == "table":
            _word_table(document, table_rows(block.rows), literal=block.literal)
            continue
        if block.kind == "heading":
            paragraph = document.add_paragraph(style=f"Heading {block.level}")
        elif block.kind in {"bullet", "number"}:
            paragraph = document.add_paragraph(style="List Bullet" if block.kind == "bullet" else None)
            paragraph.paragraph_format.left_indent = Inches(0.25 + min(block.level, 8) * 0.2)
            if block.kind == "number":
                paragraph.add_run(block.marker + " ")
        else:
            paragraph = document.add_paragraph()
        if block.kind == "code":
            run = paragraph.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif block.literal:
            paragraph.add_run(block.text)
        else:
            _inline(paragraph, block.text)
    if rows:
        _word_table(document, rows, literal=True)
    docx_format(document, rules)
    document.save(path)


def _pdf(
    path: Path, title: str, content: str, rows: list[list[Any]], *, source_blocks: list[Block] | None = None
) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    all_text = content + title + "".join(text(cell) for row in rows for cell in row)
    if source_blocks is not None:
        all_text += "".join(
            block.text + "".join(text(cell) for row in block.rows for cell in row) for block in source_blocks
        )
    if any(ord(char) > 0xFFFF for char in all_text):
        raise DocumentError("document_pdf_glyph_unsupported")
    font_name = pdf_font(all_text)
    base = ParagraphStyle("Body", fontName=font_name, fontSize=11, leading=17, wordWrap="CJK", spaceAfter=8)
    heading = ParagraphStyle("Heading", parent=base, fontSize=15, leading=21, spaceBefore=10, keepWithNext=True)
    title_style = ParagraphStyle("Title", parent=base, fontSize=22, leading=28, spaceAfter=16)
    story = [Paragraph(escape(title), title_style)] if title else []

    def add_table(data: list[list[Any]]) -> None:
        rectangular = _rectangular(data)
        weights = [
            max(6, min(35, max(display_width(row[col]) for row in rectangular))) for col in range(len(rectangular[0]))
        ]
        width = A4[0] - 108
        cells = [[Paragraph(escape(text(value)).replace("\n", "<br/>"), base) for value in row] for row in rectangular]
        table = LongTable(
            cells, colWidths=[width * weight / sum(weights) for weight in weights], repeatRows=1, hAlign="LEFT"
        )
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7EDF3")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.extend([table, Spacer(1, 10)])

    for block in source_blocks if source_blocks is not None else blocks(content):
        if block.kind == "table":
            add_table(table_rows(block.rows))
        else:
            value = escape(block.text).replace("\n", "<br/>")
            if block.kind != "code" and not block.literal:
                value = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", value)
                value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", value)
                value = re.sub(r"`([^`]+)`", r"\1", value)
            if block.kind in {"bullet", "number"}:
                value = (block.marker if block.kind == "number" else "-") + " " + value
            style = copy(heading if block.kind == "heading" else base)
            if block.kind in {"bullet", "number"}:
                style.leftIndent = 12 + min(block.level, 8) * 12
            story.append(Paragraph(value, style))
    if rows:
        add_table(rows)
    SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=54, rightMargin=54, topMargin=48, bottomMargin=48, title=title
    ).build(story)
