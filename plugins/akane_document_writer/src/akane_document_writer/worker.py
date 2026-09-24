"""Single owned document operation; no host imports or child processes."""

from __future__ import annotations

import json
from pathlib import Path
import sys
import tempfile

if not __package__:
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from akane_document_writer import DocumentError, render_document, render_sources, style_document


def deny_child(event, args):
    if event in {"subprocess.Popen", "os.system", "os.posix_spawn", "os.spawn"}:
        raise PermissionError("document_descendant_process_forbidden")


def health():
    if sys.version_info < (3, 11):
        raise DocumentError("document_python_version_unsupported")
    from docx import Document
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory(prefix="akane-document-health-") as temporary:
        root = Path(temporary)
        for fmt in ("docx", "xlsx", "pdf"):
            path = root / f"sample.{fmt}"
            render_document(output_path=path, output_format=fmt, rows=[["值", "状态"], [0, False]])
            if fmt != "pdf":
                style_document(
                    source_path=path,
                    output_path=root / f"styled.{fmt}",
                    styles={"header": {"font_color": "000000", "bold": True}},
                )
        if Document(root / "styled.docx").tables[0].cell(1, 0).text != "0":
            raise DocumentError("document_runtime_probe_failed")
        book = load_workbook(root / "styled.xlsx")
        try:
            if book.active["A2"].value != 0 or book.active["B2"].value is not False:
                raise DocumentError("document_runtime_probe_failed")
        finally:
            book.close()
        if not (root / "sample.pdf").read_bytes().startswith(b"%PDF-"):
            raise DocumentError("document_runtime_probe_failed")
    return {"status": "ready"}


def execute(request_path):
    path = Path(request_path)
    if path.stat().st_size > 4_000_000:
        raise DocumentError("document_request_too_large")
    request = json.loads(path.read_text(encoding="utf-8"))
    operation = request["operation"]
    common = {"output_path": Path(request["output_path"]), "styles": request.get("formatting")}
    if operation == "style":
        return style_document(source_path=Path(request["source_path"]), **common)
    common.update(output_format=request["output_format"], title=request.get("output_title", ""))
    if operation == "sources":
        paths = request["material_paths"]
        if not isinstance(paths, list) or not 1 <= len(paths) <= 20:
            raise DocumentError("document_material_invalid")
        materials = []
        for value in paths:
            source = Path(value)
            if source.stat().st_size > 8_000_000:
                raise DocumentError("document_material_invalid")
            materials.append(json.loads(source.read_text(encoding="utf-8")))
        return render_sources(materials=materials, **common)
    if operation != "render":
        raise DocumentError("document_operation_invalid")
    return render_document(content=request.get("content_markdown", ""), rows=request.get("table_rows"), **common)


def main():
    sys.addaudithook(deny_child)
    try:
        result = health() if sys.argv[1:] == ["health"] else execute(sys.argv[1])
        print(json.dumps({"ok": True, "result": result}, ensure_ascii=True))
        return 0
    except DocumentError as error:
        reason = str(error)
    except ImportError:
        reason = "document_dependency_missing"
    except Exception:
        reason = "document_execution_failed"
    print(json.dumps({"ok": False, "reason": reason}))
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
