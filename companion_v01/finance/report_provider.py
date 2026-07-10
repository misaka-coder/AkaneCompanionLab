from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from html import escape
import hashlib
import importlib.util
import json
import math
import os
from pathlib import Path
import re
from typing import Any
import uuid
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from services.market_data import (
    MarketDataValidationError,
    MarketQuoteSnapshot,
    MarketSeries,
    timestamp_to_iso,
)

from .market_service import compute_quote_metrics, compute_series_metrics


SUPPORTED_FINANCE_REPORT_TYPES = ("security_brief", "market_comparison")
SUPPORTED_FINANCE_REPORT_FORMATS = ("md", "pdf", "xlsx")
FINANCE_REPORT_DISCLAIMER = (
    "本报告仅用于信息整理与研究辅助，不构成投资建议、收益承诺或交易指令。"
    "行情、公告和统计指标可能存在延迟、修订或口径差异，请以标注的数据来源和 as_of 时间为准。"
)


@dataclass(frozen=True)
class FinanceReportRequest:
    codes: tuple[str, ...]
    output_format: str
    report_type: str = "security_brief"
    interval: str = "1d"
    adjusted: str = "none"
    lookback: int = 120
    chart_targets: tuple[str, ...] = ()
    title: str = ""
    analysis_summary: str = ""
    risk_notes: tuple[str, ...] = ()
    watch_items: tuple[str, ...] = ()

    def __post_init__(self) -> None:
        codes: list[str] = []
        for raw_code in tuple(self.codes or ()):
            code = str(raw_code or "").strip().upper()
            if not code or len(code) > 40 or not re.fullmatch(r"[A-Z0-9_.:-]+", code):
                raise _invalid("codes", "codes must contain explicit provider security codes")
            if code not in codes:
                codes.append(code)
        if not 1 <= len(codes) <= 5:
            raise _invalid("codes", "one to five security codes are required")

        output_format = str(self.output_format or "").strip().lower().lstrip(".")
        if output_format not in SUPPORTED_FINANCE_REPORT_FORMATS:
            raise _invalid("output_format", "output_format must be md, pdf, or xlsx")
        report_type = str(self.report_type or "security_brief").strip().lower()
        if report_type not in SUPPORTED_FINANCE_REPORT_TYPES:
            raise _invalid("report_type", "report_type must be security_brief or market_comparison")
        if report_type == "security_brief" and len(codes) != 1:
            raise _invalid("codes", "security_brief requires exactly one security code")

        interval = str(self.interval or "1d").strip().lower()
        if interval != "1d":
            raise _invalid("interval", "finance reports currently support only 1d series")
        adjusted = str(self.adjusted or "none").strip().lower()
        if adjusted not in {"none", "forward", "backward"}:
            raise _invalid("adjusted", "adjusted must be none, forward, or backward")
        if isinstance(self.lookback, bool):
            raise _invalid("lookback", "lookback must be an integer from 20 to 250")
        try:
            lookback = int(self.lookback)
        except (TypeError, ValueError) as exc:
            raise _invalid("lookback", "lookback must be an integer from 20 to 250") from exc
        if not 20 <= lookback <= 250:
            raise _invalid("lookback", "lookback must be between 20 and 250")

        chart_targets: list[str] = []
        for raw_target in tuple(self.chart_targets or ()):
            target = str(raw_target or "").strip()
            if not target or len(target) > 120 or not re.fullmatch(r"[A-Za-z0-9_.:-]+", target):
                raise _invalid("chart_ids", "chart_ids must contain generated ids or handles")
            if target not in chart_targets:
                chart_targets.append(target)
        if len(chart_targets) > 5:
            raise _invalid("chart_ids", "at most five chart references are supported")

        title = _clean_text(self.title, max_chars=100, multiline=False)
        analysis_summary = _clean_text(self.analysis_summary, max_chars=6000, multiline=True)
        risk_notes = _clean_text_list(self.risk_notes, field="risk_notes", limit=10, max_chars=500)
        watch_items = _clean_text_list(self.watch_items, field="watch_items", limit=10, max_chars=500)

        object.__setattr__(self, "codes", tuple(codes))
        object.__setattr__(self, "output_format", output_format)
        object.__setattr__(self, "report_type", report_type)
        object.__setattr__(self, "interval", interval)
        object.__setattr__(self, "adjusted", adjusted)
        object.__setattr__(self, "lookback", lookback)
        object.__setattr__(self, "chart_targets", tuple(chart_targets))
        object.__setattr__(self, "title", title)
        object.__setattr__(self, "analysis_summary", analysis_summary)
        object.__setattr__(self, "risk_notes", risk_notes)
        object.__setattr__(self, "watch_items", watch_items)


@dataclass(frozen=True)
class FinanceSeriesEvidence:
    source: str
    series: MarketSeries

    def public_summary(self) -> dict[str, Any]:
        points = tuple(self.series.points)
        latest = points[-1]
        metrics = compute_series_metrics(self.series)
        return {
            "provider": self.series.provider,
            "source": self.source,
            "code": self.series.code,
            "interval": self.series.interval,
            "adjusted": self.series.adjusted,
            "timezone": self.series.timezone,
            "date_from": timestamp_to_iso(points[0].timestamp, self.series.timezone),
            "date_to": timestamp_to_iso(points[-1].timestamp, self.series.timezone),
            "as_of": timestamp_to_iso(self.series.as_of, self.series.timezone),
            "point_count": len(points),
            "latest_bar": {
                **latest.to_public_dict(),
                "datetime": timestamp_to_iso(latest.timestamp, self.series.timezone),
            },
            "program_metrics": metrics,
            "series_sha256": _series_fingerprint(self.series),
        }


@dataclass(frozen=True)
class FinanceQuoteEvidence:
    source: str
    quotes: tuple[MarketQuoteSnapshot, ...]
    status: str = "ok"
    reason: str = ""

    def public_summary(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "source": self.source,
            "reason": self.reason,
            "quotes": [
                {
                    **quote.to_public_dict(),
                    "as_of": timestamp_to_iso(quote.as_of, quote.timezone),
                    "program_metrics": compute_quote_metrics(quote),
                }
                for quote in self.quotes
            ],
        }


@dataclass(frozen=True)
class FinanceChartReference:
    generated_id: str
    generated_handle: str
    title: str
    path: Path
    code: str
    as_of: str
    source: str
    series_sha256: str
    width: int
    height: int

    def public_summary(self) -> dict[str, Any]:
        return {
            "generated_id": self.generated_id,
            "generated_handle": self.generated_handle,
            "title": self.title,
            "code": self.code,
            "as_of": self.as_of,
            "source": self.source,
            "series_sha256": self.series_sha256,
            "width": self.width,
            "height": self.height,
        }


@dataclass(frozen=True)
class FinanceReportArtifact:
    output_path: Path
    output_format: str
    title: str
    report_type: str
    codes: tuple[str, ...]
    created_at: str
    as_of: str
    sources: tuple[str, ...]
    evidence_sha256: str
    series_summaries: tuple[dict[str, Any], ...]
    quote_summary: dict[str, Any]
    chart_summaries: tuple[dict[str, Any], ...]
    markdown: str

    def public_metadata(self) -> dict[str, Any]:
        return {
            "report_type": self.report_type,
            "output_format": self.output_format,
            "title": self.title,
            "codes": list(self.codes),
            "created_at": self.created_at,
            "as_of": self.as_of,
            "sources": list(self.sources),
            "evidence_sha256": self.evidence_sha256,
            "series": list(self.series_summaries),
            "quotes": self.quote_summary,
            "charts": list(self.chart_summaries),
        }


class FinanceReportProvider:
    provider_id = "local_finance_report_v1"

    def render(
        self,
        *,
        request: FinanceReportRequest,
        series_evidence: tuple[FinanceSeriesEvidence, ...],
        quote_evidence: FinanceQuoteEvidence,
        chart_references: tuple[FinanceChartReference, ...],
        output_path: Path,
        created_at: int,
    ) -> FinanceReportArtifact:
        if not series_evidence:
            raise _invalid_data("series", "at least one normalized market series is required")
        series_codes = [item.series.code for item in series_evidence]
        if len(series_codes) != len(request.codes) or len(set(series_codes)) != len(series_codes):
            raise _invalid_data("series", "series evidence must contain one unique entry per requested code")
        if set(series_codes) != set(request.codes):
            raise _invalid_data("series", "series evidence must cover every requested code exactly once")
        for item in series_evidence:
            self._validate_series_evidence(item, request=request)
        self._validate_quote_evidence(quote_evidence, request=request)
        for chart in chart_references:
            if chart.code not in request.codes:
                raise _invalid("chart_ids", "referenced chart code is outside the report code set")
            if not chart.path.is_file() or chart.path.stat().st_size <= 0:
                raise _invalid_data("charts", "referenced chart file is unavailable")

        target = Path(output_path)
        if target.suffix.lower() != f".{request.output_format}":
            raise _invalid("output_format", "report output path extension does not match output_format")
        title = request.title or self._default_title(request)
        created_at_iso = timestamp_to_iso(int(created_at), "Asia/Shanghai")
        series_summaries = tuple(item.public_summary() for item in series_evidence)
        quote_summary = quote_evidence.public_summary()
        chart_summaries = tuple(item.public_summary() for item in chart_references)
        as_of_timestamps = [item.series.as_of for item in series_evidence]
        as_of_timestamps.extend(quote.as_of for quote in quote_evidence.quotes)
        report_as_of_timestamp = max(as_of_timestamps)
        for chart in chart_references:
            try:
                chart_as_of_timestamp = int(datetime.fromisoformat(chart.as_of).timestamp())
            except (TypeError, ValueError) as exc:
                raise _invalid_data("charts.as_of", "chart as_of must be an ISO timestamp") from exc
            if chart_as_of_timestamp > report_as_of_timestamp:
                raise _invalid_data("charts.as_of", "chart as_of cannot be later than the report evidence")
        as_of = timestamp_to_iso(report_as_of_timestamp, "Asia/Shanghai")
        sources = tuple(
            dict.fromkeys(
                [item.source for item in series_evidence]
                + ([quote_evidence.source] if quote_evidence.source else [])
                + [item.source for item in chart_references if item.source]
            )
        )
        evidence_payload = {
            "request": {
                "report_type": request.report_type,
                "codes": list(request.codes),
                "interval": request.interval,
                "adjusted": request.adjusted,
                "lookback": request.lookback,
            },
            "series": series_summaries,
            "quotes": quote_summary,
            "charts": chart_summaries,
        }
        evidence_sha256 = hashlib.sha256(
            json.dumps(evidence_payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode("utf-8")
        ).hexdigest()
        markdown = self._build_markdown(
            request=request,
            title=title,
            created_at=created_at_iso,
            as_of=as_of,
            series_summaries=series_summaries,
            quote_summary=quote_summary,
            chart_references=chart_references,
            output_path=target,
            sources=sources,
            evidence_sha256=evidence_sha256,
        )

        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name(f".{target.stem}.{uuid.uuid4().hex}.tmp{target.suffix}")
        try:
            if request.output_format == "md":
                temporary.write_text(markdown, encoding="utf-8")
            elif request.output_format == "pdf":
                self._write_pdf(
                    output_path=temporary,
                    request=request,
                    title=title,
                    created_at=created_at_iso,
                    as_of=as_of,
                    series_summaries=series_summaries,
                    quote_summary=quote_summary,
                    chart_references=chart_references,
                    sources=sources,
                    evidence_sha256=evidence_sha256,
                )
            elif request.output_format == "xlsx":
                self._write_xlsx(
                    output_path=temporary,
                    request=request,
                    title=title,
                    created_at=created_at_iso,
                    as_of=as_of,
                    series_evidence=series_evidence,
                    series_summaries=series_summaries,
                    quote_evidence=quote_evidence,
                    chart_references=chart_references,
                    sources=sources,
                    evidence_sha256=evidence_sha256,
                )
            if not temporary.is_file() or temporary.stat().st_size <= 0:
                raise RuntimeError("finance report renderer produced an empty file")
            temporary.replace(target)
        finally:
            if temporary.exists():
                temporary.unlink(missing_ok=True)

        return FinanceReportArtifact(
            output_path=target,
            output_format=request.output_format,
            title=title,
            report_type=request.report_type,
            codes=request.codes,
            created_at=created_at_iso,
            as_of=as_of,
            sources=sources,
            evidence_sha256=evidence_sha256,
            series_summaries=series_summaries,
            quote_summary=quote_summary,
            chart_summaries=chart_summaries,
            markdown=markdown,
        )

    def _validate_series_evidence(self, evidence: FinanceSeriesEvidence, *, request: FinanceReportRequest) -> None:
        series = evidence.series
        if not str(evidence.source or "").strip() or not str(series.provider or "").strip():
            raise _invalid_data("series.source", "series provider and source are required")
        if series.code not in request.codes:
            raise _invalid_data("series.code", "series code is outside the report request")
        if series.interval != request.interval or series.adjusted != request.adjusted:
            raise _invalid_data("series", "series interval or adjustment does not match the report request")
        if len(series.points) < 2 or len(series.points) > request.lookback:
            raise _invalid_data("series.points", "series point count is outside the requested report range")
        previous_timestamp = 0
        for index, point in enumerate(series.points):
            if int(point.timestamp) <= previous_timestamp:
                raise _invalid_data(
                    f"series.points[{index}].timestamp",
                    "series timestamps must be unique and strictly increasing",
                )
            previous_timestamp = int(point.timestamp)
            values = (point.open, point.high, point.low, point.close)
            if any(isinstance(value, bool) or not math.isfinite(float(value)) or float(value) <= 0 for value in values):
                raise _invalid_data(f"series.points[{index}]", "OHLC values must be finite and positive")
            if point.low > min(point.open, point.close) or point.high < max(point.open, point.close):
                raise _invalid_data(
                    f"series.points[{index}]",
                    "OHLC values must satisfy low <= open/close <= high",
                )
            if point.volume is not None and (
                isinstance(point.volume, bool) or not math.isfinite(float(point.volume)) or float(point.volume) < 0
            ):
                raise _invalid_data(f"series.points[{index}].volume", "volume must be finite and non-negative")
        if int(series.as_of) < int(series.points[-1].timestamp):
            raise _invalid_data("series.as_of", "series as_of cannot be earlier than its latest bar")

    def _validate_quote_evidence(self, evidence: FinanceQuoteEvidence, *, request: FinanceReportRequest) -> None:
        if evidence.quotes and not str(evidence.source or "").strip():
            raise _invalid_data("quotes.source", "quote source is required when quotes are present")
        seen: set[str] = set()
        for index, quote in enumerate(evidence.quotes):
            if quote.code not in request.codes or quote.code in seen:
                raise _invalid_data("quotes", "quote codes must be unique and inside the report code set")
            seen.add(quote.code)
            if isinstance(quote.as_of, bool) or int(quote.as_of) <= 0:
                raise _invalid_data(f"quotes[{index}].as_of", "quote as_of must be positive")
            try:
                ZoneInfo(str(quote.timezone or "Asia/Shanghai"))
            except ZoneInfoNotFoundError as exc:
                raise _invalid_data(f"quotes[{index}].timezone", "quote timezone is unknown") from exc
            for field_name in (
                "previous_close",
                "open",
                "high",
                "low",
                "last",
                "volume",
                "amount",
                "change",
                "change_pct",
            ):
                value = getattr(quote, field_name)
                if value is not None and (isinstance(value, bool) or not math.isfinite(float(value))):
                    raise _invalid_data(f"quotes[{index}].{field_name}", "quote values must be finite")
            if quote.volume is not None and quote.volume < 0:
                raise _invalid_data(f"quotes[{index}].volume", "quote volume cannot be negative")
            if quote.amount is not None and quote.amount < 0:
                raise _invalid_data(f"quotes[{index}].amount", "quote amount cannot be negative")
            if quote.high is not None and quote.low is not None and quote.high < quote.low:
                raise _invalid_data(f"quotes[{index}]", "quote high cannot be lower than low")

    def _default_title(self, request: FinanceReportRequest) -> str:
        code_text = "、".join(request.codes)
        label = "证券简报" if request.report_type == "security_brief" else "市场对比简报"
        return f"{code_text} {label}"

    def _build_markdown(
        self,
        *,
        request: FinanceReportRequest,
        title: str,
        created_at: str,
        as_of: str,
        series_summaries: tuple[dict[str, Any], ...],
        quote_summary: dict[str, Any],
        chart_references: tuple[FinanceChartReference, ...],
        output_path: Path,
        sources: tuple[str, ...],
        evidence_sha256: str,
    ) -> str:
        lines = [
            f"# {_md_plain(title)}",
            "",
            f"> 报告类型：{request.report_type}  ",
            f"> 生成时间：{created_at}  ",
            f"> 数据截至：{as_of}  ",
            f"> 标的：{'、'.join(request.codes)}  ",
            f"> 日线口径：{request.adjusted}；观察窗口：最多 {request.lookback} 个交易日",
            "",
            "## 事实与程序指标",
            "",
            "| 代码 | 区间 | 最新收盘 | 区间收益 | 最大回撤 | 年化波动率 | MA5 | MA20 |",
            "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: |",
        ]
        for item in series_summaries:
            metrics = item["program_metrics"]
            moving = metrics.get("moving_averages") or {}
            latest = item["latest_bar"]
            lines.append(
                "| "
                + " | ".join(
                    [
                        _md_cell(item["code"]),
                        _md_cell(f"{str(item['date_from'])[:10]} ~ {str(item['date_to'])[:10]}"),
                        _fmt_number(latest.get("close")),
                        _fmt_percent(metrics.get("interval_return_pct")),
                        _fmt_percent(metrics.get("max_drawdown_pct")),
                        _fmt_percent(metrics.get("annualized_volatility_pct")),
                        _fmt_number(moving.get("ma5")),
                        _fmt_number(moving.get("ma20")),
                    ]
                )
                + " |"
            )

        lines.extend(["", "## 最新行情快照", ""])
        quotes = list(quote_summary.get("quotes") or [])
        if quotes:
            lines.extend(
                [
                    "| 代码 | as_of | 最新价 | 涨跌幅 | 最高 | 最低 | 成交量 | 状态 |",
                    "| --- | --- | ---: | ---: | ---: | ---: | ---: | --- |",
                ]
            )
            for quote in quotes:
                metrics = quote.get("program_metrics") or {}
                lines.append(
                    "| "
                    + " | ".join(
                        [
                            _md_cell(quote.get("code")),
                            _md_cell(quote.get("as_of")),
                            _fmt_number(quote.get("last")),
                            _fmt_percent(metrics.get("change_pct")),
                            _fmt_number(quote.get("high")),
                            _fmt_number(quote.get("low")),
                            _fmt_number(quote.get("volume"), digits=0),
                            _md_cell(quote.get("status")),
                        ]
                    )
                    + " |"
                )
            if str(quote_summary.get("status") or "") != "ok":
                lines.extend(["", f"- 快照降级说明：{_md_plain(quote_summary.get('reason') or '部分代码缺少快照')}"])
        else:
            reason = str(quote_summary.get("reason") or "当前快照不可用").strip()
            lines.append(f"- 当前行情快照未纳入：{reason}")

        lines.extend(["", "## 图表引用", ""])
        if chart_references:
            for chart in chart_references:
                relative_path = os.path.relpath(chart.path, output_path.parent).replace("\\", "/")
                lines.extend(
                    [
                        f"### {_md_plain(chart.title)}",
                        "",
                        f"![{_md_plain(chart.title)}](<{relative_path}>)",
                        "",
                        f"- 生成物：{chart.generated_handle or chart.generated_id}",
                        f"- 代码：{chart.code}；数据截至：{chart.as_of}",
                        f"- 来源：{_md_plain(chart.source)}；序列指纹：`{chart.series_sha256}`",
                        "",
                    ]
                )
        else:
            lines.append("- 本报告未附加图表；所有数值仍来自上方程序读取与计算的结构化证据。")

        lines.extend(["", "## 分析解读（模型生成，非原始事实字段）", ""])
        lines.append(_md_plain(request.analysis_summary or "本次未附加模型分析，只保留程序核验的数据与指标。"))
        lines.extend(["", "## 风险与待验证项", ""])
        risks = list(request.risk_notes) or ["需要结合后续公告、成交结构和新的带时间戳行情继续核验。"]
        lines.extend(f"- {_md_plain(item)}" for item in risks)
        lines.extend(["", "## 后续观察", ""])
        watches = list(request.watch_items) or ["观察后续价格、成交量与公开披露是否强化或削弱当前判断。"]
        lines.extend(f"- {_md_plain(item)}" for item in watches)
        lines.extend(["", "## 数据来源与证据口径", ""])
        lines.extend(f"- {_md_plain(source)}" for source in sources)
        for item in series_summaries:
            lines.append(
                f"- {item['code']}：{item['interval']} / {item['adjusted']} / {item['point_count']} 条；"
                f"as_of {item['as_of']}；序列指纹 `{item['series_sha256']}`"
            )
        lines.extend(
            [
                f"- 报告证据指纹：`{evidence_sha256}`",
                "",
                "## 免责声明",
                "",
                FINANCE_REPORT_DISCLAIMER,
                "",
            ]
        )
        return "\n".join(lines)

    def _write_pdf(
        self,
        *,
        output_path: Path,
        request: FinanceReportRequest,
        title: str,
        created_at: str,
        as_of: str,
        series_summaries: tuple[dict[str, Any], ...],
        quote_summary: dict[str, Any],
        chart_references: tuple[FinanceChartReference, ...],
        sources: tuple[str, ...],
        evidence_sha256: str,
    ) -> None:
        if importlib.util.find_spec("reportlab") is None:
            raise _unavailable("report_renderer", "reportlab is required for PDF finance reports")
        from reportlab.lib import colors  # type: ignore
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore
        from reportlab.lib.units import mm  # type: ignore
        from reportlab.pdfbase import pdfmetrics  # type: ignore
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # type: ignore
        from reportlab.platypus import (  # type: ignore
            Image as ReportImage,
            KeepTogether,
            LongTable,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
        )

        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            font_name = "STSong-Light"
        except Exception:
            font_name = "Helvetica"
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "FinanceTitle",
            parent=styles["Title"],
            fontName=font_name,
            fontSize=20,
            leading=26,
            textColor=colors.HexColor("#1D2733"),
            spaceAfter=10,
        )
        heading_style = ParagraphStyle(
            "FinanceHeading",
            parent=styles["Heading2"],
            fontName=font_name,
            fontSize=13,
            leading=18,
            textColor=colors.HexColor("#1D2733"),
            spaceBefore=10,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "FinanceBody",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=9.5,
            leading=14,
            textColor=colors.HexColor("#344252"),
        )
        small_style = ParagraphStyle(
            "FinanceSmall",
            parent=body_style,
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#647386"),
        )
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=16 * mm,
            rightMargin=16 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
            title=title,
            author="Akane Finance Assistant",
        )
        story: list[Any] = [Paragraph(escape(title), title_style)]
        summary_table = Table(
            [
                ["报告类型", request.report_type, "标的", "、".join(request.codes)],
                ["生成时间", created_at, "数据截至", as_of],
                ["日线口径", request.adjusted, "观察窗口", f"最多 {request.lookback} 个交易日"],
            ],
            colWidths=[23 * mm, 60 * mm, 23 * mm, 60 * mm],
        )
        summary_table.setStyle(_pdf_table_style(font_name, colors=colors, header=False))
        story.extend([summary_table, Spacer(1, 8), Paragraph("事实与程序指标", heading_style)])
        metrics_rows = [["代码", "区间", "最新收盘", "区间收益", "最大回撤", "年化波动率"]]
        for item in series_summaries:
            metrics = item["program_metrics"]
            metrics_rows.append(
                [
                    item["code"],
                    f"{str(item['date_from'])[:10]}~{str(item['date_to'])[:10]}",
                    _fmt_number(item["latest_bar"].get("close")),
                    _fmt_percent(metrics.get("interval_return_pct")),
                    _fmt_percent(metrics.get("max_drawdown_pct")),
                    _fmt_percent(metrics.get("annualized_volatility_pct")),
                ]
            )
        metrics_table = LongTable(
            metrics_rows,
            repeatRows=1,
            colWidths=[25 * mm, 45 * mm, 28 * mm, 28 * mm, 28 * mm, 30 * mm],
        )
        metrics_table.setStyle(_pdf_table_style(font_name, colors=colors, header=True))
        story.append(metrics_table)

        quotes = list(quote_summary.get("quotes") or [])
        story.append(Paragraph("最新行情快照", heading_style))
        if quotes:
            quote_rows = [["代码", "as_of", "最新价", "涨跌幅", "最高/最低", "状态"]]
            for quote in quotes:
                quote_rows.append(
                    [
                        quote.get("code"),
                        str(quote.get("as_of") or "")[:19],
                        _fmt_number(quote.get("last")),
                        _fmt_percent((quote.get("program_metrics") or {}).get("change_pct")),
                        f"{_fmt_number(quote.get('high'))}/{_fmt_number(quote.get('low'))}",
                        quote.get("status"),
                    ]
                )
            quote_table = LongTable(
                quote_rows,
                repeatRows=1,
                colWidths=[25 * mm, 45 * mm, 25 * mm, 25 * mm, 35 * mm, 27 * mm],
            )
            quote_table.setStyle(_pdf_table_style(font_name, colors=colors, header=True))
            story.append(quote_table)
            if str(quote_summary.get("status") or "") != "ok":
                story.append(
                    Paragraph(
                        escape(str(quote_summary.get("reason") or "部分代码缺少快照")),
                        small_style,
                    )
                )
        else:
            story.append(Paragraph(escape(str(quote_summary.get("reason") or "当前快照不可用")), body_style))

        if chart_references:
            story.append(Paragraph("图表引用", heading_style))
            for index, chart in enumerate(chart_references):
                width = 176 * mm
                ratio = chart.height / chart.width if chart.width > 0 and chart.height > 0 else 0.5625
                image = ReportImage(str(chart.path), width=width, height=width * ratio)
                caption = Paragraph(
                    escape(
                        f"{chart.title}｜{chart.generated_handle or chart.generated_id}｜"
                        f"{chart.code}｜as_of {chart.as_of}｜{chart.source}"
                    ),
                    small_style,
                )
                story.append(KeepTogether([image, Spacer(1, 3), caption]))
                if index < len(chart_references) - 1:
                    story.append(PageBreak())

        story.append(Paragraph("分析解读（模型生成，非原始事实字段）", heading_style))
        _append_pdf_text(story, request.analysis_summary or "本次未附加模型分析。", body_style, Paragraph, Spacer)
        risk_flow = [Paragraph("风险与待验证项", heading_style)]
        for item in request.risk_notes or ("需要结合后续公告、成交结构和新的带时间戳行情继续核验。",):
            risk_flow.append(Paragraph(f"• {escape(item)}", body_style))
        story.append(KeepTogether(risk_flow))
        watch_flow = [Paragraph("后续观察", heading_style)]
        for item in request.watch_items or ("观察后续价格、成交量与公开披露是否强化或削弱当前判断。",):
            watch_flow.append(Paragraph(f"• {escape(item)}", body_style))
        story.append(KeepTogether(watch_flow))
        source_flow = [Paragraph("数据来源与证据口径", heading_style)]
        if sources:
            source_flow.append(Paragraph(f"• {escape(sources[0])}", small_style))
        story.append(KeepTogether(source_flow))
        for source in sources[1:]:
            story.append(Paragraph(f"• {escape(source)}", small_style))
        for item in series_summaries:
            story.append(
                Paragraph(
                    escape(
                        f"{item['code']}｜{item['interval']}｜{item['adjusted']}｜{item['point_count']} 条｜"
                        f"as_of {item['as_of']}｜series_sha256 {item['series_sha256']}"
                    ),
                    small_style,
                )
            )
        story.append(Paragraph(escape(f"evidence_sha256 {evidence_sha256}"), small_style))
        story.append(Paragraph("免责声明", heading_style))
        story.append(Paragraph(escape(FINANCE_REPORT_DISCLAIMER), body_style))

        def footer(canvas, document) -> None:
            canvas.saveState()
            canvas.setFont(font_name, 7)
            canvas.setFillColor(colors.HexColor("#7B8794"))
            canvas.drawString(16 * mm, 8 * mm, f"as_of {as_of}")
            canvas.drawRightString(194 * mm, 8 * mm, f"Page {document.page}")
            canvas.restoreState()

        doc.build(story, onFirstPage=footer, onLaterPages=footer)

    def _write_xlsx(
        self,
        *,
        output_path: Path,
        request: FinanceReportRequest,
        title: str,
        created_at: str,
        as_of: str,
        series_evidence: tuple[FinanceSeriesEvidence, ...],
        series_summaries: tuple[dict[str, Any], ...],
        quote_evidence: FinanceQuoteEvidence,
        chart_references: tuple[FinanceChartReference, ...],
        sources: tuple[str, ...],
        evidence_sha256: str,
    ) -> None:
        if importlib.util.find_spec("openpyxl") is None:
            raise _unavailable("report_renderer", "openpyxl is required for XLSX finance reports")
        from openpyxl import Workbook  # type: ignore
        from openpyxl.drawing.image import Image as ExcelImage  # type: ignore
        from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore

        workbook = Workbook()
        workbook.properties.title = title
        workbook.properties.subject = f"{request.report_type}; as_of {as_of}"
        workbook.properties.creator = "Akane Finance Assistant"
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        summary_rows = [
            ["字段", "值"],
            ["标题", title],
            ["报告类型", request.report_type],
            ["标的", "、".join(request.codes)],
            ["生成时间", created_at],
            ["数据截至", as_of],
            ["周期", request.interval],
            ["复权", request.adjusted],
            ["观察窗口", request.lookback],
            ["证据指纹", evidence_sha256],
        ]
        for row in summary_rows:
            summary_sheet.append(row)
        _style_xlsx_sheet(summary_sheet, Font, PatternFill, Alignment, freeze="A2", auto_filter="A1:B1")

        metrics_sheet = workbook.create_sheet("Metrics")
        metrics_headers = [
            "code",
            "date_from",
            "date_to",
            "as_of",
            "point_count",
            "latest_close",
            "interval_return_pct",
            "max_drawdown_pct",
            "annualized_volatility_pct",
            "relative_volume_20",
            "ma5",
            "ma10",
            "ma20",
            "ma60",
            "series_sha256",
        ]
        metrics_sheet.append(metrics_headers)
        for item in series_summaries:
            metrics = item["program_metrics"]
            moving = metrics.get("moving_averages") or {}
            metrics_sheet.append(
                [
                    item["code"],
                    item["date_from"],
                    item["date_to"],
                    item["as_of"],
                    item["point_count"],
                    item["latest_bar"].get("close"),
                    metrics.get("interval_return_pct"),
                    metrics.get("max_drawdown_pct"),
                    metrics.get("annualized_volatility_pct"),
                    metrics.get("relative_volume_20"),
                    moving.get("ma5"),
                    moving.get("ma10"),
                    moving.get("ma20"),
                    moving.get("ma60"),
                    item["series_sha256"],
                ]
            )
        _style_xlsx_sheet(
            metrics_sheet,
            Font,
            PatternFill,
            Alignment,
            freeze="A2",
            auto_filter=f"A1:O{metrics_sheet.max_row}",
        )

        quotes_sheet = workbook.create_sheet("Quotes")
        quote_headers = [
            "code",
            "as_of",
            "previous_close",
            "open",
            "high",
            "low",
            "last",
            "volume",
            "amount",
            "change",
            "change_pct",
            "status",
            "source",
        ]
        quotes_sheet.append(quote_headers)
        for quote in quote_evidence.quotes:
            quotes_sheet.append(
                [
                    quote.code,
                    timestamp_to_iso(quote.as_of, quote.timezone),
                    quote.previous_close,
                    quote.open,
                    quote.high,
                    quote.low,
                    quote.last,
                    quote.volume,
                    quote.amount,
                    quote.change,
                    quote.change_pct,
                    quote.status,
                    quote_evidence.source,
                ]
            )
        if not quote_evidence.quotes:
            quotes_sheet.append(["unavailable", quote_evidence.reason])
        _style_xlsx_sheet(
            quotes_sheet,
            Font,
            PatternFill,
            Alignment,
            freeze="A2",
            auto_filter=f"A1:M{quotes_sheet.max_row}",
        )

        for evidence in series_evidence:
            sheet = workbook.create_sheet(_safe_sheet_name(f"{evidence.series.code}_OHLCV"))
            sheet.append(["datetime", "open", "high", "low", "close", "volume", "amount"])
            for point in evidence.series.points:
                sheet.append(
                    [
                        timestamp_to_iso(point.timestamp, evidence.series.timezone),
                        point.open,
                        point.high,
                        point.low,
                        point.close,
                        point.volume,
                        point.amount,
                    ]
                )
            _style_xlsx_sheet(
                sheet,
                Font,
                PatternFill,
                Alignment,
                freeze="A2",
                auto_filter=f"A1:G{sheet.max_row}",
            )

        evidence_sheet = workbook.create_sheet("Evidence")
        evidence_sheet.append(["kind", "code_or_id", "source", "as_of", "sha256"])
        for item in series_summaries:
            evidence_sheet.append(["market_series", item["code"], item["source"], item["as_of"], item["series_sha256"]])
        for chart in chart_references:
            evidence_sheet.append(
                [
                    "market_chart",
                    chart.generated_handle or chart.generated_id,
                    chart.source,
                    chart.as_of,
                    chart.series_sha256,
                ]
            )
        for source in sources:
            evidence_sheet.append(["source", "", source, as_of, ""])
        evidence_sheet.append(["report_evidence", "", self.provider_id, as_of, evidence_sha256])
        _style_xlsx_sheet(
            evidence_sheet,
            Font,
            PatternFill,
            Alignment,
            freeze="A2",
            auto_filter=f"A1:E{evidence_sheet.max_row}",
        )

        notes_sheet = workbook.create_sheet("Notes")
        notes_sheet.append(["section", "content"])
        notes_sheet.append(["analysis_model_generated", request.analysis_summary or "本次未附加模型分析。"])
        notes_sheet.append(["quote_status", f"{quote_evidence.status}: {quote_evidence.reason}".strip()])
        for item in request.risk_notes or ("需要结合后续公告和新行情继续核验。",):
            notes_sheet.append(["risk", item])
        for item in request.watch_items or ("观察后续价格、成交量与公开披露。",):
            notes_sheet.append(["watch", item])
        notes_sheet.append(["disclaimer", FINANCE_REPORT_DISCLAIMER])
        _style_xlsx_sheet(notes_sheet, Font, PatternFill, Alignment, freeze="A2", auto_filter="A1:B1")
        notes_sheet.column_dimensions["B"].width = 90
        for cell in notes_sheet["B"]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        if chart_references:
            chart_sheet = workbook.create_sheet("Charts")
            row = 1
            for chart in chart_references:
                chart_sheet.cell(row=row, column=1, value=chart.title)
                chart_sheet.cell(
                    row=row + 1,
                    column=1,
                    value=f"{chart.generated_handle or chart.generated_id} | {chart.code} | as_of {chart.as_of}",
                )
                image = ExcelImage(str(chart.path))
                image.width = 960
                image.height = int(960 * (chart.height / chart.width)) if chart.width and chart.height else 540
                chart_sheet.add_image(image, f"A{row + 3}")
                row += 34
            chart_sheet.column_dimensions["A"].width = 120

        workbook.save(str(output_path))


def _pdf_table_style(font_name: str, *, colors: Any, header: bool):
    from reportlab.platypus import TableStyle  # type: ignore

    commands = [
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D7DEE8")),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#F6F8FB")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF3")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1D2733")),
            ]
        )
    return TableStyle(commands)


def _append_pdf_text(story: list[Any], text: str, style: Any, paragraph_type: Any, spacer_type: Any) -> None:
    for block in re.split(r"\n\s*\n", str(text or "").strip()):
        clean = "<br/>".join(escape(line) for line in block.splitlines() if line.strip())
        if clean:
            story.extend([paragraph_type(clean, style), spacer_type(1, 4)])


def _style_xlsx_sheet(
    sheet: Any, font_type: Any, fill_type: Any, alignment_type: Any, *, freeze: str, auto_filter: str
) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.value = f"'{cell.value}"
    header_fill = fill_type("solid", fgColor="E8EDF3")
    for cell in sheet[1]:
        cell.font = font_type(bold=True, color="1D2733")
        cell.fill = header_fill
        cell.alignment = alignment_type(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = alignment_type(vertical="top", wrap_text=True)
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = auto_filter
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(48, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width


def _safe_sheet_name(value: str) -> str:
    text = re.sub(r"[\[\]:*?/\\]+", "_", str(value or "Sheet").strip())
    return (text or "Sheet")[:31]


def _series_fingerprint(series: MarketSeries) -> str:
    payload = {
        "provider": series.provider,
        "code": series.code,
        "interval": series.interval,
        "adjusted": series.adjusted,
        "timezone": series.timezone,
        "as_of": series.as_of,
        "points": [point.to_public_dict() for point in series.points],
    }
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode("utf-8")
    ).hexdigest()


def _fmt_number(value: Any, *, digits: int = 2) -> str:
    if value is None:
        return "—"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "—"
    if not math.isfinite(number):
        return "—"
    return f"{number:,.{digits}f}"


def _fmt_percent(value: Any) -> str:
    rendered = _fmt_number(value)
    return f"{rendered}%" if rendered != "—" else rendered


def _md_cell(value: Any) -> str:
    return str(value if value is not None else "—").replace("|", "\\|").replace("\n", " ").strip() or "—"


def _md_plain(value: Any) -> str:
    text = str(value or "").replace("\\", "\\\\").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"([`*_{}\[\]()#+.!|])", r"\\\1", text)


def _clean_text(value: Any, *, max_chars: int, multiline: bool) -> str:
    text = str(value or "").replace("\x00", "").strip()
    text = "\n".join(line.rstrip() for line in text.splitlines()) if multiline else re.sub(r"\s+", " ", text)
    if any(ord(char) < 32 and char not in {"\n", "\t"} for char in text):
        raise _invalid("text", "text contains unsupported control characters")
    return text[:max_chars].strip()


def _clean_text_list(value: Any, *, field: str, limit: int, max_chars: int) -> tuple[str, ...]:
    if not isinstance(value, (list, tuple)):
        raise _invalid(field, f"{field} must be an array")
    if len(value) > limit:
        raise _invalid(field, f"{field} contains too many items")
    result: list[str] = []
    for item in value:
        text = _clean_text(item, max_chars=max_chars, multiline=False)
        if not text:
            raise _invalid(field, f"{field} items cannot be empty")
        if text not in result:
            result.append(text)
    return tuple(result)


def _invalid(field: str, reason: str) -> MarketDataValidationError:
    return MarketDataValidationError(
        field=field,
        reason=reason,
        code="invalid_arguments",
        status="invalid_arguments",
        provider="local_finance_report_v1",
    )


def _invalid_data(field: str, reason: str) -> MarketDataValidationError:
    return MarketDataValidationError(
        field=field,
        reason=reason,
        code="invalid_finance_evidence",
        status="invalid_data",
        provider="local_finance_report_v1",
    )


def _unavailable(field: str, reason: str) -> MarketDataValidationError:
    return MarketDataValidationError(
        field=field,
        reason=reason,
        code="dependency_unavailable",
        status="unavailable",
        provider="local_finance_report_v1",
    )


__all__ = [
    "FINANCE_REPORT_DISCLAIMER",
    "FinanceChartReference",
    "FinanceQuoteEvidence",
    "FinanceReportArtifact",
    "FinanceReportProvider",
    "FinanceReportRequest",
    "FinanceSeriesEvidence",
    "SUPPORTED_FINANCE_REPORT_FORMATS",
    "SUPPORTED_FINANCE_REPORT_TYPES",
]
