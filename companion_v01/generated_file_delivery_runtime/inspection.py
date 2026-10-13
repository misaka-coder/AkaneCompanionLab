"""Inspect generated content and render bounded file previews."""
from __future__ import annotations

import importlib.util
import io
import json
import zipfile
from collections import deque
from collections.abc import Iterable, Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any



class GeneratedFileReadError(RuntimeError):
    """A failed extraction, never exception prose masquerading as file text."""


def inspect_generated_file(
    service: Any,
    *,
    profile_user_id: str,
    session_id: str,
    target: str = "latest",
    section: str = "content",
    max_chars: int = 12000,
) -> dict[str, Any]:
    normalized_section = service._normalize_generated_inspection_section(section)
    char_budget = service._normalize_inspection_max_chars(max_chars)
    generated = service._resolve_generated_file(
        profile_user_id=profile_user_id,
        session_id=session_id,
        target=target,
    )
    if generated is None:
        return {
            "ok": False,
            "generated": None,
            "inspection": None,
            "error": "generated_file_not_found",
            "followup_context": (
                f"你刚刚想查看 {target or 'latest'}，但没有找到对应的生成文件。"
                "请自然向用户确认要查看哪一份，不要重复调用 inspect_generated_file。"
            ),
        }

    absolute_path = service.absolute_path(generated)
    generated["absolute_path"] = str(absolute_path)
    if normalized_section != "summary" and (not absolute_path.exists() or not absolute_path.is_file()):
        return {
            "ok": False,
            "generated": generated,
            "inspection": None,
            "error": "generated_file_missing_on_disk",
            "followup_context": (
                f"你刚刚想查看 {service._generated_display_name(generated)}，但本地文件本体已经找不到了。"
                "你仍然可以参考生成文件工作台里的摘要；请自然告诉用户文件记录还在，但文件本体缺失。"
            ),
        }

    try:
        inspection = service._inspect_generated_file_content(
            generated=generated, path=absolute_path, section=normalized_section, max_chars=char_budget,
        )
    except GeneratedFileReadError:
        return {
            "ok": False, "generated": generated, "inspection": None,
            "error": "generated_file_read_failed",
            "followup_context": "生成文件仍有记录，但本次无法读取正文（generated_file_read_failed）。不能把读取失败当成已读全文。",
        }
    return {
        "ok": True,
        "generated": generated,
        "inspection": inspection,
        "followup_context": service._build_generated_inspection_followup(
            generated=generated,
            inspection=inspection,
        ),
    }


def normalize_generated_inspection_section(service: Any, value: Any) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    if lowered.startswith(("file:", "member:")):
        prefix, _, rest = text.partition(":")
        return f"file:{rest.strip()[:240]}"
    aliases = {
        "": "content",
        "full": "content",
        "all": "content",
        "content": "content",
        "body": "content",
        "正文": "content",
        "全文": "content",
        "全部": "content",
        "head": "head",
        "start": "head",
        "开头": "head",
        "前面": "head",
        "tail": "tail",
        "end": "tail",
        "结尾": "tail",
        "末尾": "tail",
        "summary": "summary",
        "摘要": "summary",
        "概览": "summary",
        "manifest": "manifest",
        "manifest.json": "file:manifest.json",
        "readme": "file:README.md",
        "readme.md": "file:README.md",
        "file_list": "file_list",
        "list": "file_list",
        "members": "file_list",
        "文件列表": "file_list",
        "列表": "file_list",
    }
    return aliases.get(lowered, lowered if lowered in {"content", "head", "tail", "summary", "manifest", "file_list"} else "content")


def normalize_inspection_max_chars(service: Any, value: Any) -> int:
    try:
        parsed = int(float(str(value).strip()))
    except Exception:
        parsed = 12000
    return max(500, min(40000, parsed))


def inspect_generated_file_content(
    service: Any,
    *,
    generated: dict[str, Any],
    path: Path,
    section: str,
    max_chars: int,
) -> dict[str, Any]:
    output_format = str(generated.get("output_format") or path.suffix.lstrip(".")).strip().lower()
    if section == "summary":
        content = service._render_generated_summary_inspection(generated=generated, path=path)
        return {
            "section": "summary",
            "content": service._clip_inspection_text(content, max_chars=max_chars),
            "truncated": len(content) > max_chars,
            "source_kind": "metadata",
        }

    if output_format == "zip" or path.suffix.lower() == ".zip":
        return service._inspect_generated_zip(
            generated=generated,
            path=path,
            section=section,
            max_chars=max_chars,
        )

    try:
        if section == "tail" and output_format in {"docx", "xlsx", "pdf", "png"}:
            text = _read_tail_text_material(path, output_format=output_format, max_chars=max_chars)
        else:
            text = service._read_generated_text_material(path=path, output_format=output_format, max_chars=max_chars * 3)
            if text and section == "tail" and len(text) > max_chars:
                text = _read_tail_text_material(path, output_format=output_format, max_chars=max_chars)
    except Exception:
        raise GeneratedFileReadError("generated_file_read_failed") from None
    if not text:
        content = service._render_generated_binary_inspection(generated=generated, path=path, output_format=output_format)
        return {
            "section": section,
            "content": service._clip_inspection_text(content, max_chars=max_chars),
            "truncated": False,
            "source_kind": "binary_or_unsupported",
        }

    sliced = service._slice_inspection_text(text, section=section, max_chars=max_chars)
    return {
        "section": section,
        "content": sliced["content"],
        "truncated": sliced["truncated"],
        "source_kind": output_format or "text",
    }


def read_generated_text_material(
    service: Any,
    *,
    path: Path,
    output_format: str,
    max_chars: int,
    text_inspect_formats: set[str],
) -> str:
    suffix = (output_format or path.suffix.lstrip(".")).lower()
    try:
        if suffix in text_inspect_formats:
            return service._read_plain_text_file(path, max_chars=max_chars)
        if suffix == "docx":
            return service._read_generated_docx(path=path, max_chars=max_chars)
        if suffix == "xlsx":
            return service._read_generated_xlsx(path=path, max_chars=max_chars)
        if suffix == "pdf":
            return service._read_generated_pdf(path=path, max_chars=max_chars)
        if suffix == "png":
            # Locally typeset cards carry their exact text in PNG metadata.
            # This is embedded text, not an OCR claim about arbitrary pixels.
            from PIL import Image

            with Image.open(path) as image:
                return str(image.info.get("Description") or "")[:max_chars + 1]
    except Exception:
        raise GeneratedFileReadError("generated_file_read_failed") from None
    return ""


def read_plain_text_file(service: Any, path: Path, *, max_chars: int) -> str:
    return _read_encoded_prefix(lambda encoding: path.open("r", encoding=encoding), max_chars=max_chars)


def read_generated_docx(service: Any, *, path: Path, max_chars: int) -> str:
    with _document_lines(path, "docx") as lines:
        return _render_text_window(lines, max_chars=max_chars)


def read_generated_xlsx(service: Any, *, path: Path, max_chars: int) -> str:
    with _document_lines(path, "xlsx") as lines:
        return _render_text_window(lines, max_chars=max_chars)


def read_generated_pdf(service: Any, *, path: Path, max_chars: int) -> str:
    with _document_lines(path, "pdf") as lines:
        return _render_text_window(lines, max_chars=max_chars)


def inspect_generated_zip(
    service: Any,
    *,
    generated: dict[str, Any],
    path: Path,
    section: str,
    max_chars: int,
) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(path, "r") as archive:
            if section == "file_list":
                content = service._render_zip_file_list(archive)
                sliced = service._slice_inspection_text(content, section="content", max_chars=max_chars)
                return {
                    "section": "file_list",
                    "content": sliced["content"],
                    "truncated": sliced["truncated"],
                    "source_kind": "zip",
                }
            if section == "manifest":
                section = "file:manifest.json"
            if section.startswith("file:"):
                member_name = service._resolve_zip_member_name(archive, section.partition(":")[2].strip())
                content = service._read_zip_member_for_inspection(archive, member_name=member_name, max_chars=max_chars)
                sliced = service._slice_inspection_text(content, section="content", max_chars=max_chars)
                return {
                    "section": f"file:{member_name}" if member_name else section,
                    "content": sliced["content"],
                    "truncated": sliced["truncated"],
                    "source_kind": "zip_member",
                }

            member = service._resolve_zip_member_name(archive, "manifest.json") or service._resolve_zip_member_name(archive, "README.md")
            if member:
                content = service._read_zip_member_for_inspection(archive, member_name=member, max_chars=max_chars)
            else:
                content = service._render_zip_file_list(archive)
            sliced = service._slice_inspection_text(content, section="content", max_chars=max_chars)
            return {
                "section": f"file:{member}" if member else "file_list",
                "content": sliced["content"],
                "truncated": sliced["truncated"],
                "source_kind": "zip",
            }
    except Exception:
        raise GeneratedFileReadError("generated_file_read_failed") from None


def render_zip_file_list(service: Any, archive: zipfile.ZipFile) -> str:
    lines = ["zip 文件列表："]
    for info in archive.infolist()[:400]:
        if info.is_dir():
            continue
        lines.append(f"- {info.filename}（{service._format_file_size(int(info.file_size))}）")
    if len(archive.infolist()) > 400:
        lines.append(f"... 还有 {len(archive.infolist()) - 400} 个条目未显示。")
    return "\n".join(lines)


def resolve_zip_member_name(service: Any, archive: zipfile.ZipFile, target: str) -> str:
    normalized = str(target or "").replace("\\", "/").strip().lstrip("/")
    if not normalized or ".." in normalized.split("/"):
        return ""
    names = [info.filename for info in archive.infolist() if not info.is_dir()]
    for name in names:
        if name == normalized:
            return name
    lowered = normalized.lower()
    for matches in (
        [name for name in names if name.lower() == lowered],
        [name for name in names if name.lower().endswith("/" + lowered)],
        [name for name in names if lowered in name.lower()],
    ):
        if matches:
            return matches[0] if len(matches) == 1 else ""
    return ""


def read_zip_member_for_inspection(
    service: Any,
    archive: zipfile.ZipFile,
    *,
    member_name: str,
    max_chars: int,
    text_inspect_formats: set[str],
) -> str:
    if not member_name:
        raise GeneratedFileReadError("generated_file_read_failed")
    info = archive.getinfo(member_name)
    suffix = Path(member_name).suffix.lower().lstrip(".")
    basename = Path(member_name).name.lower()
    if suffix not in text_inspect_formats and basename not in {"readme", "license"}:
        return (
            f"{member_name} 是 zip 内的二进制或不适合直接展开的文件，"
            f"大小 {service._format_file_size(int(info.file_size))}。"
        )
    text = _read_encoded_prefix(
        lambda encoding: io.TextIOWrapper(archive.open(member_name), encoding=encoding, newline=""),
        max_chars=max_chars,
    )
    if suffix == "json":
        try:
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except Exception:
            pass
    return text[: max_chars + 1]


def slice_inspection_text(service: Any, text: str, *, section: str, max_chars: int) -> dict[str, Any]:
    raw = str(text or "")
    if section == "tail" and len(raw) > max_chars:
        return {
            "content": "（以下是文件末尾片段）\n" + raw[-max_chars:],
            "truncated": True,
        }
    if section == "head" and len(raw) > max_chars:
        return {
            "content": raw[:max_chars] + "\n\n（内容未显示完；如需继续查看，可指定 tail 或更大的 max_chars。）",
            "truncated": True,
        }
    return {
        # Keep source text separate from display notices: the continuation
        # offset is the number of source characters, not the notice length.
        "content": raw[:max_chars],
        "truncated": len(raw) > max_chars,
    }


def clip_inspection_text(service: Any, text: str, *, max_chars: int) -> str:
    raw = str(text or "")
    if len(raw) <= max_chars:
        return raw
    return raw[:max_chars] + "\n\n（内容已按 max_chars 截断；如需更多，请指定更大的 max_chars 或查看 tail。）"


def _read_encoded_prefix(open_source, *, max_chars: int) -> str:
    """Decode a bounded prefix, retaining one source character to prove more text."""
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp936", "latin-1"):
        try:
            with open_source(encoding) as source:
                return source.read(max(0, int(max_chars)) + 1)
        except UnicodeDecodeError:
            continue
    raise GeneratedFileReadError("generated_file_read_failed")


def _render_text_window(
    lines: Iterable[str], *, max_chars: int, tail: bool = False, separator: str = "\n",
) -> str:
    """Count once per piece and retain at most one character beyond the budget."""
    limit = max(0, int(max_chars)) + 1
    parts: deque[str] = deque()
    used = 0
    first = True
    for line in lines:
        if not line:
            continue
        piece = ("" if first else separator) + line
        first = False
        piece = piece[-limit:] if tail else piece[:limit - used]
        parts.append(piece)
        used += len(piece)
        if tail:
            while used > limit:
                extra = used - limit
                if len(parts[0]) <= extra:
                    used -= len(parts.popleft())
                else:
                    parts[0] = parts[0][extra:]
                    used = limit
        elif used >= limit:
            break
    return "".join(parts)


@contextmanager
def _document_lines(path: Path, kind: str) -> Iterator[Iterable[str]]:
    module = {"docx": "docx", "xlsx": "openpyxl", "pdf": "pypdf"}[kind]
    if importlib.util.find_spec(module) is None:
        raise GeneratedFileReadError("generated_file_read_failed")
    if kind == "docx":
        from docx import Document  # type: ignore
        document = Document(str(path))

        def lines():
            for paragraph in document.paragraphs:
                text = str(paragraph.text or "").strip()
                if text:
                    yield text
            for table in document.tables:
                for row in table.rows:
                    cells = [str(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        yield "| " + " | ".join(cells) + " |"

        yield lines()
    elif kind == "xlsx":
        from openpyxl import load_workbook  # type: ignore
        workbook = load_workbook(str(path), read_only=True, data_only=True)

        def lines():
            for sheet in workbook.worksheets:
                yield f"## Sheet: {sheet.title}"
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if cell is None else str(cell) for cell in row]
                    if any(cell.strip() for cell in cells):
                        yield "| " + " | ".join(cells) + " |"

        rows = lines()
        try:
            yield rows
        finally:
            rows.close()
            workbook.close()
    else:
        from pypdf import PdfReader  # type: ignore
        with path.open("rb") as source:
            reader = PdfReader(source)

            def lines():
                for index, page in enumerate(reader.pages, start=1):
                    text = str(page.extract_text() or "").strip()
                    if text:
                        yield f"## Page {index}"
                        yield text

            yield lines()


def _read_tail_text_material(path: Path, *, output_format: str, max_chars: int) -> str:
    """Scan readable content to EOF while retaining only the actual tail window."""
    if output_format in {"docx", "xlsx", "pdf"}:
        with _document_lines(path, output_format) as lines:
            return _render_text_window(lines, max_chars=max_chars, tail=True)
    if output_format == "png":
        from PIL import Image
        with Image.open(path) as image:
            return str(image.info.get("Description") or "")[-(max_chars + 1):]
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp936", "latin-1"):
        try:
            with path.open("r", encoding=encoding) as source:
                chunks = iter(lambda: source.read(64 * 1024), "")
                return _render_text_window(chunks, max_chars=max_chars, tail=True, separator="")
        except UnicodeDecodeError:
            continue
    raise GeneratedFileReadError("generated_file_read_failed")
