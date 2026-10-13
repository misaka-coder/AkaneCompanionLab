"""Document reading for AttachmentInboxService."""

from __future__ import annotations

import importlib.util
from pathlib import Path

from .section_content import SectionContent


class AttachmentDocumentReadingMixin:
    def _read_docx_section(self, path: Path, *, section: str, max_chars: int) -> str:
        if importlib.util.find_spec("docx") is None:
            return ""
        from docx import Document  # type: ignore

        document = Document(str(path))
        if "表" in section or "table" in section.lower():
            index = self._parse_section_index(section) or 1
            if not (1 <= index <= len(document.tables)):
                return ""
            table = document.tables[index - 1]
            lines = [f"[表格 {index}] 行数约 {len(table.rows)}，列数约 {len(table.columns)}"]
            used = len(lines[0])
            for row in table.rows:
                lines.append(
                    " | ".join(str(cell.text or "").strip() for cell in row.cells if str(cell.text or "").strip())
                )
                used += len(lines[-1]) + 1
                if used > max_chars:
                    break
            return "\n".join(lines).strip()[:max_chars]
        paragraphs = [str(paragraph.text or "").strip() for paragraph in document.paragraphs]
        paragraphs = [text for text in paragraphs if text]
        return self._slice_plain_text_section("\n".join(paragraphs), section=section, max_chars=max_chars)

    def _read_xlsx_section(self, path: Path, *, section: str, max_chars: int) -> str:
        return self._read_xlsx_result(path, section=section, max_chars=max_chars).text

    def _read_xlsx_result(self, path: Path, *, section: str, max_chars: int) -> SectionContent:
        if importlib.util.find_spec("openpyxl") is None:
            return SectionContent("", reason="parser_unavailable")
        from openpyxl import load_workbook  # type: ignore

        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        try:
            normalized = str(section or "").strip().lower()
            full = normalized in {"全文", "全部", "完整", "all", "full"}
            sheet_names = workbook.sheetnames if full else [self._select_sheet_name(workbook.sheetnames, section)]
            named = any(name.lower() in normalized for name in workbook.sheetnames)
            line_range = self._parse_line_range(section)
            truncated = normalized in {"", "预览", "当前可用片段"} and len(sheet_names) < len(workbook.sheetnames)
            lines: list[str] = []
            used = 0
            for sheet_name in sheet_names:
                sheet = workbook[sheet_name]
                if line_range is None:
                    start_row, end_row = 1, int(sheet.max_row or 1) if full or named else min(int(sheet.max_row or 1), 80)
                    truncated = truncated or end_row < int(sheet.max_row or end_row)
                else:
                    start_row, end_row = line_range
                    if start_row > int(sheet.max_row or 0):
                        return SectionContent("", reason="section_out_of_range")
                    end_row = min(end_row, int(sheet.max_row or end_row))
                header = f"[{sheet_name}] 行 {start_row}-{end_row} / 共 {sheet.max_row or '?'} 行"
                used += len(header) + (1 if lines else 0)
                lines.append(header)
                for row in sheet.iter_rows(min_row=start_row, max_row=end_row, values_only=True):
                    values = [self._cell_to_text(value) for value in row]
                    if any(values):
                        lines.append(" | ".join(values))
                        used += len(lines[-1]) + 1
                    if used > max_chars:
                        truncated = True
                        break
                if used > max_chars:
                    truncated = True
                    break
            return SectionContent("\n".join(lines).strip()[:max_chars], truncated=truncated)
        finally:
            workbook.close()

    def _read_pdf_section(self, path: Path, *, section: str, max_chars: int) -> str:
        return self._read_pdf_result(path, section=section, max_chars=max_chars).text

    def _read_pdf_result(self, path: Path, *, section: str, max_chars: int) -> SectionContent:
        if importlib.util.find_spec("pypdf") is None:
            return SectionContent("", reason="parser_unavailable")
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        page_count = len(reader.pages)
        if page_count <= 0:
            return SectionContent("", reason="section_empty")
        truncated = False
        normalized_section = str(section or "").strip().lower()
        if normalized_section in {"全文", "全部", "完整", "all", "full"}:
            start_page, end_page = 1, page_count
        else:
            page_range = self._parse_page_range(section)
            if page_range is None:
                start_page, end_page = 1, min(page_count, 5)
                truncated = end_page < page_count
            else:
                start_page, end_page = page_range
                if start_page > page_count:
                    return SectionContent("", reason="section_out_of_range")
                start_page = max(1, min(page_count, start_page))
                end_page = max(start_page, min(page_count, end_page))
        chunks = [f"[PDF 页 {start_page}-{end_page} / 共 {page_count} 页]"]
        used = len(chunks[0])
        for page_no in range(start_page, end_page + 1):
            text = str(reader.pages[page_no - 1].extract_text() or "").strip()
            if text:
                chunks.append(f"--- 第 {page_no} 页 ---\n{text}")
                used += len(chunks[-1]) + 2
            if used > max_chars:
                truncated = True
                break
        if len(chunks) == 1:
            return SectionContent("", reason="source_text_unavailable")
        return SectionContent("\n\n".join(chunks).strip()[:max_chars], truncated=truncated)
