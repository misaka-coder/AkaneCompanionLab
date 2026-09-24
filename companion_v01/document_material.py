"""Host-owned, bounded content extraction, not a renderer or preview fallback.

Used in a disposable process by the resource port. The JSON representation is
private invocation data; parser errors never contain source text or paths.
"""

from __future__ import annotations

import csv
from datetime import date, datetime, time
import io
import json
from pathlib import Path
import sys
import warnings
import xml.etree.ElementTree as ET
import zipfile


MATERIAL_REASONS = frozenset(
    {
        "document_source_too_large",
        "document_content_too_large",
        "document_archive_invalid",
        "document_encrypted_unsupported",
        "document_dependency_missing",
        "document_read_failed",
        "document_format_unsupported",
        "document_encoding_unsupported",
        "document_text_invalid",
    }
)
TEXT_FORMATS = frozenset(
    {
        "txt",
        "md",
        "markdown",
        "html",
        "htm",
        "json",
        "xml",
        "csv",
        "tsv",
        "srt",
        "lrc",
        "vtt",
        "log",
        "yaml",
        "yml",
        "toml",
        "ini",
        "cfg",
        "conf",
        "py",
        "js",
        "ts",
        "css",
        "sql",
    }
)
MAX_CHARACTERS = 1_000_000
MAX_CELLS = 100_000
MAX_JSON_BYTES = 8_000_000
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class MaterialError(ValueError):
    pass


def _archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 100_000_000:
            raise MaterialError("document_source_too_large")
        if len({entry.filename for entry in entries}) != len(entries):
            raise MaterialError("document_archive_invalid")
        if any(entry.flag_bits & 1 for entry in entries):
            raise MaterialError("document_encrypted_unsupported")


def _text(path: Path) -> str:
    data = path.read_bytes()
    encodings = ("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
    for encoding in encodings:
        try:
            value = data.decode(encoding)
            break
        except UnicodeError:
            continue
    else:
        raise MaterialError("document_encoding_unsupported")
    if any(ord(char) < 32 and char not in "\t\r\n\f" for char in value):
        raise MaterialError("document_text_invalid")
    if len(value) > MAX_CHARACTERS:
        raise MaterialError("document_content_too_large")
    return value


def _paragraph(element: ET.Element) -> str:
    return "".join(
        node.text or "" if node.tag == W + "t" else "\t" if node.tag == W + "tab" else "\n"
        for node in element.iter()
        if node.tag in {W + "t", W + "tab", W + "br", W + "cr"}
    )


def _docx(path: Path, limitations: set[str]) -> list[dict]:
    _archive(path)
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        for name in names:
            lowered = name.lower()
            if lowered.startswith(
                (
                    "word/header",
                    "word/footer",
                    "word/footnotes",
                    "word/endnotes",
                    "word/comments",
                    "word/media/",
                    "word/embeddings/",
                )
            ):
                limitations.add("document_additional_parts_not_extracted")
        raw = archive.read("word/document.xml")
    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise MaterialError("document_archive_invalid")
    document = ET.fromstring(raw)
    body = document.find(W + "body")
    if body is None:
        raise MaterialError("document_archive_invalid")
    unsupported = {
        "drawing",
        "pict",
        "object",
        "sdt",
        "altChunk",
        "ins",
        "del",
        "moveFrom",
        "moveTo",
        "fldSimple",
        "fldChar",
        "sym",
        "footnoteReference",
        "endnoteReference",
        "numPr",
        "vMerge",
        "gridSpan",
        "hyperlink",
    }
    for node in body.iter():
        if node.tag.removeprefix(W) in unsupported or "}oMath" in node.tag:
            limitations.add("document_complex_content_not_extracted")
    blocks = []
    cells = 0
    for element in body:
        if element.tag == W + "p":
            blocks.append({"kind": "paragraph", "text": _paragraph(element)})
        elif element.tag == W + "tbl":
            rows = []
            for row in element.findall(W + "tr"):
                data = []
                for cell in row.findall(W + "tc"):
                    cells += 1
                    if cells > MAX_CELLS:
                        raise MaterialError("document_content_too_large")
                    if cell.find(W + "tbl") is not None:
                        limitations.add("document_nested_table_not_extracted")
                    data.append("\n".join(_paragraph(p) for p in cell.iter(W + "p")))
                rows.append(data)
            blocks.append({"kind": "table", "rows": rows})
        elif element.tag != W + "sectPr":
            limitations.add("document_complex_content_not_extracted")
    # Numbered styles can supply numPr even when document.xml has none.
    with zipfile.ZipFile(path) as archive:
        if "word/styles.xml" in names:
            raw_styles = archive.read("word/styles.xml")
            if b"<!DOCTYPE" in raw_styles.upper() or b"<!ENTITY" in raw_styles.upper():
                raise MaterialError("document_archive_invalid")
            styles = ET.fromstring(raw_styles)
            numbered = {
                style.get(W + "styleId")
                for style in styles.findall(W + "style")
                if style.find(".//" + W + "numPr") is not None
            }
            if any(node.get(W + "val") in numbered for node in body.iter(W + "pStyle")):
                limitations.add("document_numbering_not_extracted")
    return blocks


def _xlsx(path: Path, limitations: set[str]) -> list[dict]:
    _archive(path)
    from openpyxl import load_workbook

    # data_only=False is essential: cached formula results may be absent/stale.
    with warnings.catch_warnings(record=True) as notices:
        warnings.simplefilter("always")
        workbook = load_workbook(path, read_only=True, data_only=False)
    if notices:
        limitations.add("document_workbook_features_not_extracted")
    blocks = []
    cells = 0
    try:
        with zipfile.ZipFile(path) as archive:
            if any(
                name.startswith(
                    (
                        "xl/charts/",
                        "xl/drawings/",
                        "xl/media/",
                        "xl/comments",
                        "xl/threadedComments/",
                        "xl/embeddings/",
                        "xl/pivotTables/",
                    )
                )
                for name in archive.namelist()
            ):
                limitations.add("document_workbook_features_not_extracted")
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    raw = archive.read(name)
                    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                        raise MaterialError("document_archive_invalid")
                    tree = ET.fromstring(raw)
                    if any(
                        node.tag.rsplit("}", 1)[-1] in {"mergeCells", "hyperlinks", "extLst"} for node in tree.iter()
                    ):
                        limitations.add("document_workbook_features_not_extracted")
        if len(workbook.worksheets) > 100:
            raise MaterialError("document_content_too_large")
        for sheet in workbook.worksheets:
            # Ignore potentially false stored dimensions; bound actual cells.
            sheet.reset_dimensions()
            rows = []
            for row in sheet.iter_rows():
                cells += len(row)
                if cells > MAX_CELLS or len(rows) >= 10000:
                    raise MaterialError("document_content_too_large")
                data = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        value = {"type": "formula", "expression": value}
                    elif isinstance(value, (datetime, date, time)):
                        value = {"type": "datetime", "value": value.isoformat()}
                    elif value is not None and type(value) not in (str, int, float, bool):
                        limitations.add("document_workbook_cell_not_extracted")
                        value = None
                    data.append(value)
                rows.append(data)
            blocks.append({"kind": "sheet", "name": sheet.title, "rows": rows})
    finally:
        workbook.close()
    return blocks


def _pdf(path: Path, limitations: set[str]) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(path, strict=True)
    try:
        if reader.is_encrypted:
            raise MaterialError("document_encrypted_unsupported")
        if len(reader.pages) > 200:
            raise MaterialError("document_content_too_large")
        blocks = []
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if page.images or not text.strip() or "\ufffd" in text or "\x00" in text:
                limitations.add("document_pdf_visual_content_not_extracted")
            if page.get("/Annots"):
                limitations.add("document_pdf_annotations_not_extracted")
            blocks.append({"kind": "page", "number": index, "text": text})
        return blocks
    finally:
        reader.close()


def read_document_material(path: Path) -> dict:
    """Return all extractable content, or fail. Never substitute a preview."""
    try:
        if path.stat().st_size > 50_000_000:
            raise MaterialError("document_source_too_large")
        fmt = path.suffix.lower().removeprefix(".")
        limitations: set[str] = set()
        if fmt in TEXT_FORMATS:
            content = _text(path)
            if fmt in {"csv", "tsv"}:
                rows = list(csv.reader(io.StringIO(content, newline=""), delimiter="\t" if fmt == "tsv" else ","))
                if sum(map(len, rows)) > MAX_CELLS:
                    raise MaterialError("document_content_too_large")
                blocks = [{"kind": "table", "rows": rows}]
            else:
                blocks = [{"kind": "text", "text": content}]
        elif fmt in {"docx", "xlsx", "pdf"}:
            blocks = {"docx": _docx, "xlsx": _xlsx, "pdf": _pdf}[fmt](path, limitations)
        else:
            raise MaterialError("document_format_unsupported")
        result = {
            "schema": "akane.document-material.v1",
            "source_format": fmt,
            "complete": not limitations,
            "limitations": sorted(limitations),
            "blocks": blocks,
        }
        encoded = json.dumps(result, ensure_ascii=False, allow_nan=False)
        if len(encoded) > MAX_CHARACTERS or len(encoded.encode("utf-8")) > MAX_JSON_BYTES:
            raise MaterialError("document_content_too_large")
        return result
    except MaterialError:
        raise
    except ImportError:
        raise MaterialError("document_dependency_missing") from None
    except (zipfile.BadZipFile, ET.ParseError):
        raise MaterialError("document_archive_invalid") from None
    except Exception:
        raise MaterialError("document_read_failed") from None


def main() -> int:
    try:
        source, output = map(Path, sys.argv[1:])
        material = read_document_material(source)
        with output.open("x", encoding="utf-8", newline="") as stream:
            json.dump(material, stream, ensure_ascii=False, allow_nan=False)
        print(json.dumps({"ok": True}))
        return 0
    except Exception as error:
        reason = str(error) if isinstance(error, MaterialError) else "document_read_failed"
        print(json.dumps({"ok": False, "reason": reason}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
