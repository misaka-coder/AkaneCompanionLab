"""Bounded Word, Excel, PDF and archive text rendering."""

from __future__ import annotations

import importlib.util
from collections.abc import Iterable
from pathlib import Path
import zipfile

from .definitions import WorkspaceReaderUnavailable


def render_window(lines: Iterable[str], *, max_chars: int, capped_flag: bool) -> str | tuple[str, bool]:
    """Retain at most one character beyond the budget to prove truncation."""
    probe_limit = max(0, int(max_chars)) + 1
    parts: list[str] = []
    char_count = 0
    capped = False
    for line in lines:
        if not line:
            continue
        piece = ('\n' if parts else '') + line
        remaining = probe_limit - char_count
        parts.append(piece[:remaining])
        char_count += min(len(piece), remaining)
        if char_count >= probe_limit:
            capped = True
            break
    content = ''.join(parts)
    return (content, capped) if capped_flag else content


class WorkspaceDocumentReadersMixin:
    def _read_docx(self, path: Path, *, max_chars: int, capped_flag: bool = False) -> str:
        if importlib.util.find_spec("docx") is None:
            raise WorkspaceReaderUnavailable("python-docx is not installed")
        from docx import Document  # type: ignore

        document = Document(str(path))

        def lines():
            for paragraph in document.paragraphs:
                text = str(paragraph.text or '').strip()
                if text:
                    yield text
            for table in document.tables:
                for row in table.rows:
                    cells = [str(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        yield '| ' + ' | '.join(cells) + ' |'

        return render_window(lines(), max_chars=max_chars, capped_flag=capped_flag)

    def _read_xlsx(self, path: Path, *, max_chars: int, capped_flag: bool = False) -> str:
        if importlib.util.find_spec("openpyxl") is None:
            raise WorkspaceReaderUnavailable("openpyxl is not installed")
        from openpyxl import load_workbook  # type: ignore

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        try:
            def lines():
                for sheet in workbook.worksheets:
                    yield f'## Sheet: {sheet.title}'
                    for row in sheet.iter_rows(values_only=True):
                        cells = ['' if cell is None else str(cell) for cell in row]
                        if any(cell.strip() for cell in cells):
                            yield '| ' + ' | '.join(cells) + ' |'

            return render_window(lines(), max_chars=max_chars, capped_flag=capped_flag)
        finally:
            workbook.close()

    def _read_pdf(self, path: Path, *, max_chars: int, capped_flag: bool = False) -> str:
        if importlib.util.find_spec("pypdf") is None:
            raise WorkspaceReaderUnavailable("pypdf is not installed")
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))

        def lines():
            for index, page in enumerate(reader.pages, start=1):
                text = str(page.extract_text() or '').strip()
                if text:
                    yield f'## Page {index}'
                    yield text

        rendered = render_window(lines(), max_chars=max_chars, capped_flag=capped_flag)
        if not (rendered[0] if capped_flag else rendered):
            raise WorkspaceReaderUnavailable('PDF has no readable text layer')
        return rendered

    def _read_zip_manifest(self, path: Path, *, max_chars: int, capped_flag: bool = False) -> str:
        with zipfile.ZipFile(path, "r") as archive:
            def lines():
                yield '## Archive entries'
                for info in archive.infolist():
                    yield f'- {info.filename} ({info.file_size} bytes)'

            return render_window(lines(), max_chars=max_chars, capped_flag=capped_flag)
