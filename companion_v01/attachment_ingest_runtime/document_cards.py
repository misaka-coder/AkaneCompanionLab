"""Optional Word, Excel and PDF readers with explicit unavailable states."""

from __future__ import annotations

import importlib.util
from pathlib import Path
from typing import Any


class AttachmentDocumentCardsMixin:
    """Methods use the state owned by AttachmentIngestService."""

    def _build_document_card(
        self,
        *,
        source_path: Path,
        title: str,
        suffix: str,
        mime_type: str,
        file_size: int,
    ) -> dict[str, Any]:
        if suffix == ".docx":
            return self._build_docx_card(
                source_path=source_path,
                title=title,
                mime_type=mime_type,
                file_size=file_size,
            )
        if suffix == ".xlsx":
            return self._build_xlsx_card(
                source_path=source_path,
                title=title,
                mime_type=mime_type,
                file_size=file_size,
            )
        if suffix == ".pdf":
            return self._build_pdf_card(
                source_path=source_path,
                title=title,
                mime_type=mime_type,
                file_size=file_size,
            )
        return self._build_unsupported_document_card(
            title=title,
            suffix=suffix,
            mime_type=mime_type,
            file_size=file_size,
            reason="暂不支持解析该文档格式。",
        )

    def _build_docx_card(self, *, source_path: Path, title: str, mime_type: str, file_size: int) -> dict[str, Any]:
        if importlib.util.find_spec("docx") is None:
            return self._build_unsupported_document_card(
                title=title,
                suffix=".docx",
                mime_type=mime_type,
                file_size=file_size,
                reason="缺少 python-docx，暂时只能看到文件名，不能读取 Word 内容。",
            )
        try:
            from docx import Document  # type: ignore

            document = Document(str(source_path))
            paragraphs = [str(paragraph.text or "").strip() for paragraph in document.paragraphs]
            paragraphs = [text for text in paragraphs if text]
            headings: list[str] = []
            for paragraph in document.paragraphs:
                text = str(paragraph.text or "").strip()
                style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
                if text and style_name.lower().startswith("heading"):
                    headings.append(text[:80])
                if len(headings) >= 8:
                    break

            table_summaries: list[dict[str, Any]] = []
            for index, table in enumerate(document.tables[:5], start=1):
                rows = []
                for row in table.rows[:3]:
                    rows.append([str(cell.text or "").strip()[:80] for cell in row.cells[:8]])
                table_summaries.append(
                    {
                        "index": index,
                        "row_count": len(table.rows),
                        "column_count": len(table.columns),
                        "preview_rows": rows,
                    }
                )

            preview = "\n".join(paragraphs[:20])[:4000]
            detail = {
                "summary": f"Word 文档 {title}，约 {len(paragraphs)} 段，{len(document.tables)} 个表格。",
                "file_kind": "docx",
                "mime_type": mime_type,
                "file_size": file_size,
                "paragraph_count": len(paragraphs),
                "table_count": len(document.tables),
                "headings": headings,
                "tables": table_summaries,
                "text_preview": preview,
                "preview_is_truncated": len(paragraphs) > 20
                or len("\n".join(paragraphs[:20])) > 4000
                or len(document.tables) > 5,
            }
            hint_parts = [detail["summary"]]
            if headings:
                hint_parts.append("标题：" + "、".join(headings[:4]))
            elif paragraphs:
                hint_parts.append("开头：" + " ".join(paragraphs[:2])[:100])
            return {
                "summary_title": title,
                "short_hint": " ".join(hint_parts)[:240],
                "detail": detail,
            }
        except Exception as exc:
            return self._build_unsupported_document_card(
                title=title,
                suffix=".docx",
                mime_type=mime_type,
                file_size=file_size,
                reason=f"Word 文档解析失败：{str(exc)[:160]}",
            )

    def _build_xlsx_card(self, *, source_path: Path, title: str, mime_type: str, file_size: int) -> dict[str, Any]:
        if importlib.util.find_spec("openpyxl") is None:
            return self._build_unsupported_document_card(
                title=title,
                suffix=".xlsx",
                mime_type=mime_type,
                file_size=file_size,
                reason="缺少 openpyxl，暂时只能看到文件名，不能读取 Excel 内容。",
            )
        try:
            from openpyxl import load_workbook  # type: ignore

            workbook = load_workbook(filename=str(source_path), read_only=True, data_only=True)
            try:
                sheet_count = len(workbook.sheetnames)
                sheet_cards: list[dict[str, Any]] = []
                for sheet_name in workbook.sheetnames[:5]:
                    sheet = workbook[sheet_name]
                    preview_rows: list[list[str]] = []
                    for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
                        values = [self._cell_to_text(value)[:80] for value in list(row)[:10]]
                        if any(values):
                            preview_rows.append(values)
                    columns = [cell for cell in (preview_rows[0] if preview_rows else []) if cell]
                    sheet_cards.append(
                        {
                            "name": str(sheet_name),
                            "max_row": int(sheet.max_row or 0),
                            "max_column": int(sheet.max_column or 0),
                            "columns": columns[:12],
                            "preview_rows": preview_rows,
                        }
                    )
            finally:
                workbook.close()

            sheet_names = [item["name"] for item in sheet_cards]
            first_columns = []
            if sheet_cards:
                first_columns = list(sheet_cards[0].get("columns") or [])
            detail = {
                "summary": f"Excel 工作簿 {title}，包含 {sheet_count} 个工作表。",
                "file_kind": "xlsx",
                "mime_type": mime_type,
                "file_size": file_size,
                "sheet_count": sheet_count,
                "sheet_names": sheet_names,
                "sheets": sheet_cards,
                "columns": first_columns,
                "text_preview": self._render_xlsx_preview(sheet_cards),
                "preview_is_truncated": any(
                    int(sheet.get("max_row") or 0) > 5 or int(sheet.get("max_column") or 0) > 10
                    for sheet in sheet_cards
                )
                or sheet_count > len(sheet_cards),
            }
            hint_parts = [detail["summary"]]
            if sheet_names:
                hint_parts.append("工作表：" + "、".join(sheet_names[:5]))
            if first_columns:
                hint_parts.append("首表列：" + "、".join(first_columns[:6]))
            return {
                "summary_title": title,
                "short_hint": " ".join(hint_parts)[:240],
                "detail": detail,
            }
        except Exception as exc:
            return self._build_unsupported_document_card(
                title=title,
                suffix=".xlsx",
                mime_type=mime_type,
                file_size=file_size,
                reason=f"Excel 文件解析失败：{str(exc)[:160]}",
            )

    def _build_pdf_card(self, *, source_path: Path, title: str, mime_type: str, file_size: int) -> dict[str, Any]:
        if importlib.util.find_spec("pypdf") is None:
            return self._build_unsupported_document_card(
                title=title,
                suffix=".pdf",
                mime_type=mime_type,
                file_size=file_size,
                reason="缺少 pypdf，暂时只能看到 PDF 文件名；安装依赖后可读取页数和文本层预览。",
            )
        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(str(source_path))
            page_count = len(reader.pages)
            page_texts: list[str] = []
            for page in reader.pages[:3]:
                page_texts.append(str(page.extract_text() or "").strip())
            text_preview = "\n\n".join(text for text in page_texts if text)[:4000]
            has_text = bool(text_preview.strip())
            summary = f"PDF 文档 {title}，共 {page_count} 页。"
            if not has_text:
                summary += " 目前没有读到明显文本层，可能是扫描件或图片型 PDF。"
            detail = {
                "summary": summary,
                "file_kind": "pdf",
                "mime_type": mime_type,
                "file_size": file_size,
                "page_count": page_count,
                "has_text_layer": has_text,
                "text_preview": text_preview,
                "preview_is_truncated": page_count > 3 or len("\n\n".join(text for text in page_texts if text)) > 4000,
            }
            return {
                "summary_title": title,
                "short_hint": summary[:240],
                "detail": detail,
            }
        except Exception as exc:
            return self._build_unsupported_document_card(
                title=title,
                suffix=".pdf",
                mime_type=mime_type,
                file_size=file_size,
                reason=f"PDF 解析失败：{str(exc)[:160]}",
            )

    def _build_unsupported_document_card(
        self,
        *,
        title: str,
        suffix: str,
        mime_type: str,
        file_size: int,
        reason: str,
    ) -> dict[str, Any]:
        file_kind = suffix.lstrip(".") or "document"
        return {
            "summary_title": title,
            "short_hint": f"{file_kind.upper()} 文件，约 {file_size} bytes；{reason}",
            "detail": {
                "summary": reason,
                "file_kind": file_kind,
                "mime_type": mime_type,
                "file_size": file_size,
                "parser_status": "unsupported_or_unavailable",
                "text_preview": "",
            },
        }

    def _cell_to_text(self, value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _render_xlsx_preview(self, sheets: list[dict[str, Any]]) -> str:
        lines: list[str] = []
        for sheet in sheets[:3]:
            name = str(sheet.get("name") or "Sheet")
            lines.append(f"[{name}]")
            for row in list(sheet.get("preview_rows") or [])[:5]:
                if isinstance(row, list):
                    lines.append(" | ".join(str(cell) for cell in row if str(cell).strip()))
        return "\n".join(lines)[:4000]
